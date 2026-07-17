# SPDX-License-Identifier: Apache-2.0
# The bounded selector and style-merge designs are adapted from OfficeCLI's
# Excel handler. See LICENSE.officecli and NOTICE.officecli in this package.
"""Bounded XLSX inspection, validation, and transactional cell formatting."""

from __future__ import annotations

import copy
import datetime as dt
import hashlib
import io
import math
import posixpath
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Any

from lxml import etree
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Color, PatternFill
from openpyxl.styles.numbers import BUILTIN_FORMATS_REVERSE
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from .errors import OfficeOperationError, OfficePackageError
from .models import (
    XlsxCellFormatOperation,
    XlsxCellFormatting,
    XlsxCellSelector,
    XlsxEditOperation,
    _validate_xlsx_number_format,
)
from .opc import enforce_package_preservation
from .receipt import OfficeEditTrace, xlsx_cell_path

_CONTENT_TYPES_XML = "[Content_Types].xml"
_ROOT_RELATIONSHIPS_XML = "_rels/.rels"
_WORKBOOK_XML = "xl/workbook.xml"
_WORKBOOK_RELATIONSHIPS_XML = "xl/_rels/workbook.xml.rels"
_STYLES_XML = "xl/styles.xml"
_REQUIRED_ENTRIES = frozenset(
    {
        _CONTENT_TYPES_XML,
        _ROOT_RELATIONSHIPS_XML,
        _WORKBOOK_XML,
        _WORKBOOK_RELATIONSHIPS_XML,
    }
)
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_SPREADSHEET_NAMESPACES = frozenset(
    {
        "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "http://purl.oclc.org/ooxml/spreadsheetml/main",
    }
)
_MAIN_WORKBOOK_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
_OFFICE_DOCUMENT_RELATIONSHIP_TYPES = frozenset(
    {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
        "http://purl.oclc.org/ooxml/officeDocument/relationships/officeDocument",
    }
)
_WORKSHEET_RELATIONSHIP_TYPES = frozenset(
    {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
        "http://purl.oclc.org/ooxml/officeDocument/relationships/worksheet",
    }
)

_MAX_PACKAGE_BYTES = 50 * 1024 * 1024
_MAX_ENTRY_COUNT = 5_000
_MAX_ENTRY_BYTES = 50 * 1024 * 1024
_MAX_TOTAL_UNCOMPRESSED_BYTES = 150 * 1024 * 1024
_MAX_COMPRESSION_RATIO = 250
_MAX_CONTROL_XML_BYTES = 2 * 1024 * 1024
_MAX_WORKSHEET_XML_BYTES = 35 * 1024 * 1024
_MAX_DECLARED_CELLS_PER_SHEET = 100_000
_MAX_DECLARED_CELLS = 200_000
_MAX_INSPECT_CELLS = 5_000
_MAX_INSPECT_OUTPUT_CHARS = 200_000
_MAX_INSPECT_CELL_VALUE_CHARS = 4_000
_MAX_INSPECT_HYPERLINK_CHARS = 2_000
_MAX_EDIT_CELLS = 10_000
_MAX_NEW_STYLES = 500
_DEFAULT_INSPECT_ROWS = 100
_DEFAULT_INSPECT_COLUMNS = 26
_RENDER_BLOCKED_FEATURES = frozenset(
    {
        "digital signatures",
        "macros",
        "external workbook links",
        "data connections",
        "embedded or ActiveX objects",
    }
)


@dataclass(frozen=True, slots=True)
class _PackageInfo:
    entry_count: int
    declared_cell_count: int
    sheet_cell_counts: dict[str, int]
    risky_features: tuple[str, ...]


@dataclass(slots=True)
class _EditablePackage:
    entries: list[tuple[zipfile.ZipInfo, bytes]]
    payloads: dict[str, bytes]
    comment: bytes


def _safe_xml_root(payload: bytes, *, label: str, limit: int) -> Any:
    if len(payload) > limit:
        raise OfficePackageError(f"XLSX {label} exceeds the supported size limit")
    if b"<!DOCTYPE" in payload.upper():
        raise OfficePackageError(f"XLSX {label} must not contain a document type declaration")
    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        huge_tree=False,
        remove_blank_text=False,
    )
    try:
        return etree.fromstring(payload, parser=parser)
    except (etree.XMLSyntaxError, ValueError) as exc:
        raise OfficePackageError(f"XLSX {label} is malformed: {exc}") from exc


def _validate_content_types(payload: bytes) -> None:
    root = _safe_xml_root(payload, label=_CONTENT_TYPES_XML, limit=_MAX_CONTROL_XML_BYTES)
    if root.tag != f"{{{_CONTENT_TYPES_NS}}}Types":
        raise OfficePackageError("XLSX [Content_Types].xml has an unexpected root element")

    overrides = [element for element in root.findall(f"{{{_CONTENT_TYPES_NS}}}Override") if element.get("PartName") == "/xl/workbook.xml"]
    if overrides:
        valid = len(overrides) == 1 and overrides[0].get("ContentType") == _MAIN_WORKBOOK_CONTENT_TYPE
    else:
        defaults = [element for element in root.findall(f"{{{_CONTENT_TYPES_NS}}}Default") if (element.get("Extension") or "").lower() == "xml"]
        valid = len(defaults) == 1 and defaults[0].get("ContentType") == _MAIN_WORKBOOK_CONTENT_TYPE
    if not valid:
        raise OfficePackageError("XLSX content types do not resolve a standard macro-free workbook part")


def _validate_root_relationships(payload: bytes) -> None:
    root = _safe_xml_root(payload, label=_ROOT_RELATIONSHIPS_XML, limit=_MAX_CONTROL_XML_BYTES)
    if root.tag != f"{{{_RELATIONSHIPS_NS}}}Relationships":
        raise OfficePackageError("XLSX root relationships have an unexpected root element")
    matches = []
    for relationship in root.findall(f"{{{_RELATIONSHIPS_NS}}}Relationship"):
        if relationship.get("Type") not in _OFFICE_DOCUMENT_RELATIONSHIP_TYPES:
            continue
        if relationship.get("TargetMode", "Internal").lower() == "external":
            raise OfficePackageError("XLSX workbook relationship must be internal")
        target = (relationship.get("Target") or "").replace("\\", "/").lstrip("/")
        if target == _WORKBOOK_XML:
            matches.append(relationship)
    if len(matches) != 1:
        raise OfficePackageError("XLSX root relationships do not target one workbook part")


def _validate_workbook_xml(payload: bytes) -> None:
    root = _safe_xml_root(payload, label=_WORKBOOK_XML, limit=_MAX_CONTROL_XML_BYTES)
    qname = etree.QName(root)
    if qname.localname != "workbook" or qname.namespace not in _SPREADSHEET_NAMESPACES:
        raise OfficePackageError("XLSX workbook.xml has an unexpected root element")
    sheets = next(
        (child for child in root if isinstance(child.tag, str) and etree.QName(child).localname == "sheets"),
        None,
    )
    if sheets is None or not any(isinstance(child.tag, str) and etree.QName(child).localname == "sheet" for child in sheets):
        raise OfficePackageError("XLSX workbook.xml does not contain a worksheet")


def _validate_zip_name(name: str) -> None:
    normalized = name.replace("\\", "/")
    path = PurePosixPath(normalized)
    if name != normalized or normalized.startswith("/") or ".." in path.parts:
        raise OfficePackageError(f"XLSX contains an unsafe package entry: {name}")


def _count_declared_cells(payload: bytes, *, label: str) -> int:
    root = _safe_xml_root(payload, label=label, limit=_MAX_WORKSHEET_XML_BYTES)
    root_name = etree.QName(root)
    if root_name.localname != "worksheet" or root_name.namespace not in _SPREADSHEET_NAMESPACES:
        raise OfficePackageError(f"XLSX {label} has an unexpected worksheet root element")
    count = 0
    for element in root.iter():
        if not isinstance(element.tag, str) or etree.QName(element).localname != "c":
            continue
        count += 1
        if count > _MAX_DECLARED_CELLS_PER_SHEET:
            raise OfficePackageError(f"XLSX {label} declares more than {_MAX_DECLARED_CELLS_PER_SHEET:,} cells")
    return count


def _validate_number_formats(payload: bytes) -> None:
    root = _safe_xml_root(
        payload,
        label=_STYLES_XML,
        limit=_MAX_CONTROL_XML_BYTES * 4,
    )
    root_name = etree.QName(root)
    if root_name.localname != "styleSheet" or root_name.namespace not in _SPREADSHEET_NAMESPACES:
        raise OfficePackageError("XLSX styles.xml has an unexpected root element")
    for element in root.iter():
        if not isinstance(element.tag, str) or etree.QName(element).localname != "numFmt":
            continue
        format_code = element.get("formatCode")
        if format_code is None:
            raise OfficePackageError("XLSX custom number format is missing formatCode")
        try:
            _validate_xlsx_number_format(format_code)
        except ValueError as exc:
            raise OfficePackageError(f"XLSX contains an invalid custom number format: {exc}") from exc


def _risk_labels_for_token(token: str) -> set[str]:
    normalized = token.lower()
    labels: set[str] = set()
    if "digital-signature" in normalized or "xmlsignature" in normalized:
        labels.add("digital signatures")
    if "vbaproject" in normalized or "macroenabled" in normalized or "macrosheet" in normalized:
        labels.add("macros")
    if "externallink" in normalized:
        labels.add("external workbook links")
    if "connections" in normalized or "querytable" in normalized:
        labels.add("data connections")
    if "pivotcache" in normalized or "pivottable" in normalized:
        labels.add("pivot tables")
    if "slicer" in normalized:
        labels.add("slicers")
    if "oleobject" in normalized or "activex" in normalized or normalized.endswith("/control"):
        labels.add("embedded or ActiveX objects")
    if normalized.endswith("/model") or "datamodel" in normalized:
        labels.add("workbook data models")
    return labels


def _detect_risky_features(
    archive: zipfile.ZipFile,
    names: set[str],
) -> tuple[str, ...]:
    normalized_names = {name.lower() for name in names}
    checks = {
        "digital signatures": lambda name: name.startswith("_xmlsignatures/"),
        "macros": lambda name: name.endswith("vbaproject.bin") or name.startswith("xl/macrosheets/"),
        "external workbook links": lambda name: name.startswith("xl/externallinks/"),
        "data connections": lambda name: name == "xl/connections.xml" or name.startswith("xl/querytables/"),
        "pivot tables": lambda name: name.startswith("xl/pivotcache/") or name.startswith("xl/pivottables/"),
        "slicers": lambda name: name.startswith("xl/slicers/") or name.startswith("xl/slicercaches/"),
        "embedded or ActiveX objects": lambda name: name.startswith("xl/embeddings/") or name.startswith("xl/activex/"),
        "workbook data models": lambda name: name.startswith("xl/model/"),
    }
    labels = {label for label, predicate in checks.items() if any(predicate(name) for name in normalized_names)}

    content_types = _safe_xml_root(
        archive.read(_CONTENT_TYPES_XML),
        label=_CONTENT_TYPES_XML,
        limit=_MAX_CONTROL_XML_BYTES,
    )
    for element in content_types:
        content_type = element.get("ContentType")
        if content_type:
            labels.update(_risk_labels_for_token(content_type))

    for name in sorted(candidate for candidate in names if candidate.lower().endswith(".rels")):
        relationships = _safe_xml_root(
            archive.read(name),
            label=name,
            limit=_MAX_CONTROL_XML_BYTES,
        )
        if relationships.tag != f"{{{_RELATIONSHIPS_NS}}}Relationships":
            raise OfficePackageError(f"XLSX {name} has an unexpected relationships root element")
        for relationship in relationships:
            relationship_type = relationship.get("Type")
            if relationship_type:
                labels.update(_risk_labels_for_token(relationship_type))

    return tuple(label for label in checks if label in labels)


def _load_package_info(data: bytes, *, editable: bool = False) -> _PackageInfo:
    if not data:
        raise OfficePackageError("XLSX input is empty")
    if len(data) > _MAX_PACKAGE_BYTES:
        raise OfficePackageError("XLSX input exceeds the 50 MiB package limit")

    try:
        with zipfile.ZipFile(io.BytesIO(data), mode="r") as archive:
            infos = archive.infolist()
            if len(infos) > _MAX_ENTRY_COUNT:
                raise OfficePackageError("XLSX contains too many package entries")

            names: set[str] = set()
            total_uncompressed = 0
            for info in infos:
                _validate_zip_name(info.filename)
                if info.filename in names:
                    raise OfficePackageError(f"XLSX contains a duplicate package entry: {info.filename}")
                names.add(info.filename)
                if info.flag_bits & 0x1:
                    raise OfficePackageError("Encrypted XLSX package entries are not supported")
                if info.file_size > _MAX_ENTRY_BYTES:
                    raise OfficePackageError(f"XLSX package entry is too large: {info.filename}")
                total_uncompressed += info.file_size
                if total_uncompressed > _MAX_TOTAL_UNCOMPRESSED_BYTES:
                    raise OfficePackageError("XLSX uncompressed content exceeds the supported limit")
                if info.file_size and info.compress_size == 0:
                    raise OfficePackageError(f"XLSX package entry has an invalid compression ratio: {info.filename}")
                if info.compress_size and info.file_size / info.compress_size > _MAX_COMPRESSION_RATIO:
                    raise OfficePackageError(f"XLSX package entry is suspiciously compressed: {info.filename}")

            missing = _REQUIRED_ENTRIES - names
            if missing:
                raise OfficePackageError(f"XLSX is missing required package entries: {', '.join(sorted(missing))}")
            _validate_content_types(archive.read(_CONTENT_TYPES_XML))
            _validate_root_relationships(archive.read(_ROOT_RELATIONSHIPS_XML))
            _validate_workbook_xml(archive.read(_WORKBOOK_XML))
            if _STYLES_XML in names:
                _validate_number_formats(archive.read(_STYLES_XML))

            sheet_counts: dict[str, int] = {}
            declared_cells = 0
            worksheet_parts = _worksheet_parts(
                archive.read(_WORKBOOK_XML),
                archive.read(_WORKBOOK_RELATIONSHIPS_XML),
                names,
            )
            for name in sorted(set(worksheet_parts.values())):
                count = _count_declared_cells(archive.read(name), label=name)
                sheet_counts[name] = count
                declared_cells += count
                if declared_cells > _MAX_DECLARED_CELLS:
                    raise OfficePackageError(f"XLSX worksheets declare more than {_MAX_DECLARED_CELLS:,} cells in total")

            risky_features = _detect_risky_features(archive, names)
            if "macros" in risky_features:
                raise OfficePackageError("XLSX macro content is not supported")
            if editable and risky_features:
                joined = ", ".join(risky_features)
                raise OfficeOperationError(f"XLSX editing is disabled for workbooks containing {joined}")
            return _PackageInfo(
                entry_count=len(infos),
                declared_cell_count=declared_cells,
                sheet_cell_counts=sheet_counts,
                risky_features=risky_features,
            )
    except (OfficeOperationError, OfficePackageError):
        raise
    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
        raise OfficePackageError(f"Invalid XLSX package: {exc}") from exc


def _load_editable_package(data: bytes) -> _EditablePackage:
    _load_package_info(data, editable=True)
    try:
        with zipfile.ZipFile(io.BytesIO(data), mode="r") as archive:
            entries = [(info, archive.read(info)) for info in archive.infolist()]
            return _EditablePackage(
                entries=entries,
                payloads={info.filename: payload for info, payload in entries},
                comment=archive.comment,
            )
    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
        raise OfficePackageError(f"Invalid XLSX package: {exc}") from exc


def _relationship_target(base_part: str, target: str) -> str:
    normalized_target = target.replace("\\", "/")
    if normalized_target.startswith("/"):
        normalized = posixpath.normpath(normalized_target.lstrip("/"))
    else:
        normalized = posixpath.normpath(posixpath.join(posixpath.dirname(base_part), normalized_target))
    if normalized.startswith("../") or normalized in {"", ".", ".."}:
        raise OfficePackageError("XLSX workbook contains an unsafe worksheet relationship")
    return normalized


def _worksheet_parts(
    workbook_payload: bytes,
    relationships_payload: bytes,
    package_names: set[str],
) -> dict[str, str]:
    workbook = _safe_xml_root(
        workbook_payload,
        label=_WORKBOOK_XML,
        limit=_MAX_CONTROL_XML_BYTES,
    )
    relationships = _safe_xml_root(
        relationships_payload,
        label=_WORKBOOK_RELATIONSHIPS_XML,
        limit=_MAX_CONTROL_XML_BYTES,
    )
    if relationships.tag != f"{{{_RELATIONSHIPS_NS}}}Relationships":
        raise OfficePackageError("XLSX workbook relationships have an unexpected root element")

    by_id: dict[str, Any] = {}
    for relationship in relationships:
        if not isinstance(relationship.tag, str) or etree.QName(relationship).localname != "Relationship":
            continue
        relationship_id = relationship.get("Id")
        if not relationship_id or relationship_id in by_id:
            raise OfficePackageError("XLSX workbook relationships contain a missing or duplicate ID")
        by_id[relationship_id] = relationship

    sheet_parts: dict[str, str] = {}
    casefolded_names: set[str] = set()
    claimed_parts: set[str] = set()
    for element in workbook.iter():
        if not isinstance(element.tag, str) or etree.QName(element).localname != "sheet":
            continue
        sheet_name = element.get("name")
        relationship_id = next(
            (value for key, value in element.attrib.items() if etree.QName(key).localname == "id"),
            None,
        )
        if not sheet_name or not relationship_id or relationship_id not in by_id:
            raise OfficePackageError("XLSX workbook contains a sheet without one relationship")
        folded_name = sheet_name.casefold()
        if folded_name in casefolded_names:
            raise OfficePackageError("XLSX workbook contains duplicate worksheet names")
        casefolded_names.add(folded_name)

        relationship = by_id[relationship_id]
        if relationship.get("TargetMode", "Internal").lower() == "external":
            raise OfficePackageError("XLSX worksheet relationship must be internal")
        relationship_type = relationship.get("Type") or ""
        if "macrosheet" in relationship_type.lower():
            raise OfficePackageError("XLSX macro content is not supported")
        if relationship_type not in _WORKSHEET_RELATIONSHIP_TYPES:
            raise OfficePackageError("XLSX workbook contains an unsupported sheet relationship type")
        target = relationship.get("Target")
        if not target:
            raise OfficePackageError("XLSX worksheet relationship is missing a target")
        part_name = _relationship_target(_WORKBOOK_XML, target)
        if part_name not in package_names:
            raise OfficePackageError("XLSX worksheet relationship targets a missing package part")
        if part_name in claimed_parts:
            raise OfficePackageError("XLSX worksheets must not share one package part")
        claimed_parts.add(part_name)
        sheet_parts[sheet_name] = part_name
    return sheet_parts


def _worksheet_part_for_name(package: _EditablePackage, sheet_name: str) -> str:
    relationships_payload = package.payloads.get(_WORKBOOK_RELATIONSHIPS_XML)
    if relationships_payload is None:
        raise OfficePackageError("XLSX is missing workbook relationships")
    sheet_parts = _worksheet_parts(
        package.payloads[_WORKBOOK_XML],
        relationships_payload,
        set(package.payloads),
    )
    matches = [part_name for name, part_name in sheet_parts.items() if name.casefold() == sheet_name.casefold()]
    if len(matches) != 1:
        raise OfficeOperationError(f"Worksheet was not found: {sheet_name}")
    return matches[0]


def _open_workbook(
    data: bytes,
    *,
    data_only: bool = False,
    read_only: bool = False,
    editable: bool = False,
):
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            workbook = load_workbook(
                io.BytesIO(data),
                data_only=data_only,
                read_only=read_only,
                keep_links=False,
                rich_text=True,
            )
        removal_warnings = [str(item.message) for item in caught if "not supported" in str(item.message).lower() or "will be removed" in str(item.message).lower()]
        if editable and removal_warnings:
            workbook.close()
            raise OfficeOperationError("XLSX editing is disabled because the workbook contains unsupported extensions that would be removed")
        return workbook
    except OfficeOperationError:
        raise
    except (InvalidFileException, KeyError, TypeError, ValueError, etree.XMLSyntaxError) as exc:
        raise OfficePackageError(f"XLSX workbook could not be parsed safely: {exc}") from exc


def _find_sheet(workbook, sheet_name: str):
    matches = [sheet for sheet in workbook.worksheets if sheet.title.casefold() == sheet_name.casefold()]
    if len(matches) != 1:
        raise OfficeOperationError(f"Worksheet was not found: {sheet_name}")
    return matches[0]


def _color_snapshot(color: Any) -> dict[str, Any] | None:
    if color is None or color.type is None:
        return None
    if color.type == "rgb" and color.rgb:
        raw = str(color.rgb).upper()
        if len(raw) == 8:
            return {"type": "rgb", "value": f"#{raw[-6:]}", "alpha": raw[:2]}
        return {"type": "rgb", "value": f"#{raw[-6:]}"}
    if color.type == "theme" and color.theme is not None:
        return {"type": "theme", "value": int(color.theme), "tint": float(color.tint or 0)}
    if color.type == "indexed" and color.indexed is not None:
        return {"type": "indexed", "value": int(color.indexed)}
    if color.type == "auto":
        return {"type": "automatic"}
    return {"type": str(color.type)}


def _rgb_value(color: Any) -> str | None:
    snapshot = _color_snapshot(color)
    if snapshot is None or snapshot.get("type") != "rgb":
        return None
    return str(snapshot["value"]).upper()


def _bounded_inspect_text(
    value: str | None,
    *,
    budget: dict[str, int],
    item_limit: int,
) -> tuple[str | None, bool]:
    if value is None:
        return None, False
    available = min(
        item_limit,
        max(0, _MAX_INSPECT_OUTPUT_CHARS - budget["characters"]),
    )
    returned = value[:available]
    budget["characters"] += len(returned)
    return returned, len(returned) < len(value)


def _style_snapshot(
    cell: Any,
    *,
    budget: dict[str, int],
) -> tuple[dict[str, Any], bool]:
    truncated = False

    def bounded(value: str | None, limit: int = 256) -> str | None:
        nonlocal truncated
        returned, value_truncated = _bounded_inspect_text(
            value,
            budget=budget,
            item_limit=limit,
        )
        truncated = truncated or value_truncated
        return returned

    snapshot = {
        "font": {
            "name": bounded(cell.font.name),
            "size": float(cell.font.sz) if cell.font.sz is not None else None,
            "bold": cell.font.bold,
            "italic": cell.font.italic,
            "strike": cell.font.strike,
            "underline": bounded(cell.font.underline),
            "color": _color_snapshot(cell.font.color),
        },
        "fill": {
            "pattern": bounded(cell.fill.patternType),
            "foreground": _color_snapshot(cell.fill.fgColor),
            "background": _color_snapshot(cell.fill.bgColor),
        },
        "alignment": {
            "horizontal": bounded(cell.alignment.horizontal),
            "vertical": bounded(cell.alignment.vertical),
            "wrap_text": cell.alignment.wrap_text,
            "shrink_to_fit": cell.alignment.shrink_to_fit,
            "text_rotation": cell.alignment.text_rotation,
            "indent": cell.alignment.indent,
        },
        "number_format": bounded(cell.number_format, 1_024),
        "protection": {
            "locked": cell.protection.locked,
            "hidden": cell.protection.hidden,
        },
    }
    return snapshot, truncated


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (dt.datetime, dt.date, dt.time, dt.timedelta)):
        return value.isoformat() if hasattr(value, "isoformat") else str(value)
    return str(value)


def _bounded_json_value(
    value: Any,
    *,
    budget: dict[str, int],
    item_limit: int,
) -> tuple[Any, bool]:
    normalized = _json_value(value)
    if not isinstance(normalized, str):
        return normalized, False
    return _bounded_inspect_text(
        normalized,
        budget=budget,
        item_limit=item_limit,
    )


def _formula_details(cell: Any) -> tuple[str | None, str | None, str | None]:
    if cell.data_type != "f":
        return None, None, None
    value = cell.value
    if isinstance(value, ArrayFormula):
        return value.text, "array", value.ref
    if isinstance(value, DataTableFormula):
        return None, "data_table", value.ref
    return str(value) if value is not None else None, "standard", None


def _selector_cell_value(cell: Any) -> Any:
    formula, _, _ = _formula_details(cell)
    return formula if cell.data_type == "f" else cell.value


def _cell_data_type(cell: Any) -> str:
    if cell.data_type == "f":
        return "formula"
    if cell.data_type == "e":
        return "error"
    value = cell.value
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if getattr(cell, "is_date", False) or isinstance(value, (dt.datetime, dt.date, dt.time)):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    return "string"


def _range_bounds(cell_range: str) -> tuple[int, int, int, int]:
    if "!" in cell_range or "$" in cell_range or "," in cell_range or " " in cell_range:
        raise OfficeOperationError("cell_range must be one unqualified rectangular A1 range")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range.upper())
    except ValueError as exc:
        raise OfficeOperationError("cell_range must be an A1 cell or rectangular A1:B10 range") from exc
    if min_col < 1 or min_row < 1 or max_col > 16_384 or max_row > 1_048_576:
        raise OfficeOperationError("cell_range exceeds Excel worksheet bounds")
    cell_count = (max_row - min_row + 1) * (max_col - min_col + 1)
    if cell_count > _MAX_INSPECT_CELLS:
        raise OfficeOperationError(f"cell_range may inspect at most {_MAX_INSPECT_CELLS:,} cells")
    return min_col, min_row, max_col, max_row


def _default_range(worksheet: Any) -> tuple[int, int, int, int]:
    min_row = max(1, worksheet.min_row or 1)
    min_col = max(1, worksheet.min_column or 1)
    max_row = min(worksheet.max_row or min_row, min_row + _DEFAULT_INSPECT_ROWS - 1)
    max_col = min(worksheet.max_column or min_col, min_col + _DEFAULT_INSPECT_COLUMNS - 1)
    return min_col, min_row, max_col, max_row


def _sheet_summary(worksheet: Any, index: int) -> dict[str, Any]:
    freeze_panes = worksheet.freeze_panes
    return {
        "index": index,
        "name": worksheet.title,
        "state": worksheet.sheet_state,
        "dimension": worksheet.calculate_dimension(),
        "max_row": worksheet.max_row,
        "max_column": worksheet.max_column,
        "merged_range_count": len(worksheet.merged_cells.ranges),
        "freeze_panes": freeze_panes.coordinate if hasattr(freeze_panes, "coordinate") else freeze_panes,
    }


def _ranges_intersect(
    first: tuple[int, int, int, int],
    second: tuple[int, int, int, int],
) -> bool:
    first_min_col, first_min_row, first_max_col, first_max_row = first
    second_min_col, second_min_row, second_max_col, second_max_row = second
    return not (first_max_col < second_min_col or second_max_col < first_min_col or first_max_row < second_min_row or second_max_row < first_min_row)


def inspect_xlsx(
    data: bytes,
    *,
    sheet_name: str | None = None,
    cell_range: str | None = None,
    include_cell_styles: bool = False,
) -> dict[str, Any]:
    """Return workbook metadata or a bounded machine-readable worksheet range."""
    if cell_range is not None and sheet_name is None:
        raise OfficeOperationError("sheet_name is required when cell_range is provided")
    package_info = _load_package_info(data)
    workbook = _open_workbook(data)
    cached_workbook = None
    content_budget = {"characters": 0}
    content_truncated = False
    try:
        sheets = [_sheet_summary(sheet, index) for index, sheet in enumerate(workbook.worksheets, start=1)]
        result: dict[str, Any] = {
            "format": "xlsx",
            "sheet_count": len(sheets),
            "sheets": sheets,
            "declared_cell_count": package_info.declared_cell_count,
            "risky_features": list(package_info.risky_features),
            "selected_sheet": None,
            "content_characters_returned": 0,
            "content_truncated": False,
        }
        if sheet_name is None:
            return result

        worksheet = _find_sheet(workbook, sheet_name)
        if cell_range is None:
            bounds = _default_range(worksheet)
        else:
            bounds = _range_bounds(cell_range)
        min_col, min_row, max_col, max_row = bounds
        normalized_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"

        cached_workbook = _open_workbook(data, data_only=True)
        cached_sheet = _find_sheet(cached_workbook, worksheet.title)
        cells: list[dict[str, Any]] = []
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = worksheet.cell(row=row, column=column)
                if isinstance(cell, MergedCell):
                    continue
                formula, formula_type, formula_range = _formula_details(cell)
                is_formula = cell.data_type == "f"
                if cell.value is None and not cell.has_style and cell.comment is None and cell.hyperlink is None:
                    continue
                item = {
                    "coordinate": cell.coordinate,
                    "row": row,
                    "column": column,
                    "formula_type": formula_type,
                    "data_type": _cell_data_type(cell),
                    "style_id": cell.style_id,
                }
                bounded_fields = (
                    (
                        "value",
                        None if is_formula else cell.value,
                        _MAX_INSPECT_CELL_VALUE_CHARS,
                        True,
                    ),
                    (
                        "formula",
                        formula,
                        _MAX_INSPECT_CELL_VALUE_CHARS,
                        True,
                    ),
                    ("formula_range", formula_range, 128, False),
                    (
                        "cached_value",
                        (cached_sheet.cell(row=row, column=column).value if is_formula else None),
                        _MAX_INSPECT_CELL_VALUE_CHARS,
                        True,
                    ),
                    ("number_format", cell.number_format, 1_024, False),
                    (
                        "hyperlink",
                        (cell.hyperlink.target if cell.hyperlink is not None else None),
                        _MAX_INSPECT_HYPERLINK_CHARS,
                        False,
                    ),
                )
                for key, raw_value, limit, normalize_json in bounded_fields:
                    if normalize_json:
                        returned, field_truncated = _bounded_json_value(
                            raw_value,
                            budget=content_budget,
                            item_limit=limit,
                        )
                    else:
                        returned, field_truncated = _bounded_inspect_text(
                            raw_value,
                            budget=content_budget,
                            item_limit=limit,
                        )
                    item[key] = returned
                    if field_truncated:
                        item[f"{key}_truncated"] = True
                        content_truncated = True
                if include_cell_styles:
                    formatting, formatting_truncated = _style_snapshot(
                        cell,
                        budget=content_budget,
                    )
                    item["formatting"] = formatting
                    if formatting_truncated:
                        item["formatting_truncated"] = True
                        content_truncated = True
                cells.append(item)

        merged_ranges = []
        for merged in worksheet.merged_cells.ranges:
            merged_bounds = (merged.min_col, merged.min_row, merged.max_col, merged.max_row)
            if _ranges_intersect(bounds, merged_bounds):
                merged_ranges.append(str(merged))

        result.update(
            {
                "selected_sheet": worksheet.title,
                "range": normalized_range,
                "range_cell_count": (max_row - min_row + 1) * (max_col - min_col + 1),
                "returned": len(cells),
                "has_more_rows": max_row < worksheet.max_row,
                "has_more_columns": max_col < worksheet.max_column,
                "include_cell_styles": include_cell_styles,
                "content_characters_returned": content_budget["characters"],
                "content_truncated": content_truncated,
                "merged_ranges": merged_ranges,
                "cells": cells,
            }
        )
        return result
    finally:
        if cached_workbook is not None:
            cached_workbook.close()
        workbook.close()


def _child(parent: Any, local_name: str) -> Any | None:
    return next(
        (item for item in parent if isinstance(item.tag, str) and etree.QName(item).localname == local_name),
        None,
    )


def _xml_key(element: Any) -> bytes:
    return etree.tostring(element, method="c14n", with_comments=True)


class _StyleEditor:
    def __init__(self, payload: bytes):
        self.root = _safe_xml_root(payload, label=_STYLES_XML, limit=_MAX_CONTROL_XML_BYTES * 4)
        root_name = etree.QName(self.root)
        if root_name.localname != "styleSheet" or root_name.namespace not in _SPREADSHEET_NAMESPACES:
            raise OfficePackageError("XLSX styles.xml has an unexpected root element")
        self.namespace = root_name.namespace
        self.fonts = self._required_collection("fonts")
        self.fills = self._required_collection("fills")
        self.cell_xfs = self._required_collection("cellXfs")
        self.num_fmts = _child(self.root, "numFmts")
        if not list(self.fonts) or not list(self.fills) or not list(self.cell_xfs):
            raise OfficePackageError("XLSX styles.xml is missing required base styles")
        self.new_records = 0

    def _q(self, local_name: str) -> str:
        return f"{{{self.namespace}}}{local_name}"

    def _required_collection(self, local_name: str) -> Any:
        collection = _child(self.root, local_name)
        if collection is None:
            raise OfficePackageError(f"XLSX styles.xml is missing {local_name}")
        return collection

    def _record_append(self) -> None:
        self.new_records += 1
        if self.new_records > _MAX_NEW_STYLES:
            raise OfficeOperationError(f"XLSX edit would create more than {_MAX_NEW_STYLES} style records")

    def _deduplicate(self, collection: Any, candidate: Any) -> int:
        candidate_key = _xml_key(candidate)
        for index, existing in enumerate(collection):
            if _xml_key(existing) == candidate_key:
                return index
        self._record_append()
        collection.append(candidate)
        collection.set("count", str(len(collection)))
        return len(collection) - 1

    def _indexed_clone(self, collection: Any, index: int, *, label: str) -> Any:
        values = list(collection)
        if index < 0 or index >= len(values):
            raise OfficePackageError(f"XLSX cell style references an invalid {label} index")
        return copy.deepcopy(values[index])

    def _font_child(
        self,
        font: Any,
        local_name: str,
        *,
        enabled: bool | None = None,
        value: str | None = None,
        attributes: dict[str, str] | None = None,
    ) -> None:
        matches = [item for item in font if isinstance(item.tag, str) and etree.QName(item).localname == local_name]
        if enabled is False:
            for item in matches:
                font.remove(item)
            return
        if enabled is None and value is None and attributes is None:
            return
        element = matches[0] if matches else etree.Element(self._q(local_name))
        for duplicate in matches[1:]:
            font.remove(duplicate)
        element.attrib.clear()
        if value is not None:
            element.set("val", value)
        if attributes:
            for key, item_value in attributes.items():
                element.set(key, item_value)
        if not matches:
            extension = _child(font, "extLst")
            if extension is None:
                font.append(element)
            else:
                font.insert(font.index(extension), element)

    def _font_id(self, base_xf: Any, formatting: XlsxCellFormatting) -> int:
        font_id = int(base_xf.get("fontId", "0"))
        font = self._indexed_clone(self.fonts, font_id, label="font")
        if formatting.bold is not None:
            self._font_child(font, "b", enabled=formatting.bold)
        if formatting.italic is not None:
            self._font_child(font, "i", enabled=formatting.italic)
        if formatting.strike is not None:
            self._font_child(font, "strike", enabled=formatting.strike)
        if formatting.underline is not None:
            if formatting.underline == "none":
                self._font_child(font, "u", enabled=False)
            else:
                underline_value = None if formatting.underline == "single" else "double"
                self._font_child(font, "u", enabled=True, value=underline_value)
        if formatting.font_name is not None:
            self._font_child(font, "name", value=formatting.font_name)
        if formatting.font_size is not None:
            self._font_child(font, "sz", value=f"{formatting.font_size:g}")
        if formatting.font_color is not None:
            self._font_child(
                font,
                "color",
                attributes={"rgb": _argb(formatting.font_color)},
            )
        return self._deduplicate(self.fonts, font)

    def _fill_id(self, formatting: XlsxCellFormatting) -> int:
        fill = etree.Element(self._q("fill"))
        pattern = etree.SubElement(fill, self._q("patternFill"), patternType="solid")
        etree.SubElement(pattern, self._q("fgColor"), rgb=_argb(formatting.fill_color or "000000"))
        etree.SubElement(pattern, self._q("bgColor"), indexed="64")
        return self._deduplicate(self.fills, fill)

    def _ensure_num_fmts(self) -> Any:
        if self.num_fmts is not None:
            return self.num_fmts
        self.num_fmts = etree.Element(self._q("numFmts"), count="0")
        fonts_index = self.root.index(self.fonts)
        self.root.insert(fonts_index, self.num_fmts)
        return self.num_fmts

    def _number_format_id(self, code: str) -> int:
        builtin = BUILTIN_FORMATS_REVERSE.get(code)
        if builtin is not None:
            return int(builtin)
        num_fmts = self._ensure_num_fmts()
        used_ids: list[int] = []
        for item in num_fmts:
            if not isinstance(item.tag, str) or etree.QName(item).localname != "numFmt":
                continue
            try:
                num_fmt_id = int(item.get("numFmtId", ""))
            except ValueError as exc:
                raise OfficePackageError("XLSX styles.xml contains an invalid number format ID") from exc
            used_ids.append(num_fmt_id)
            if item.get("formatCode") == code:
                return num_fmt_id
        next_id = max([163, *used_ids]) + 1
        if next_id > 65_535:
            raise OfficeOperationError("XLSX has no available custom number format IDs")
        self._record_append()
        num_fmts.append(etree.Element(self._q("numFmt"), numFmtId=str(next_id), formatCode=code))
        num_fmts.set("count", str(len(num_fmts)))
        return next_id

    def _alignment(self, cell_xf: Any, formatting: XlsxCellFormatting) -> None:
        alignment = _child(cell_xf, "alignment")
        if alignment is None:
            alignment = etree.Element(self._q("alignment"))
            insertion_target = _child(cell_xf, "protection")
            if insertion_target is None:
                insertion_target = _child(cell_xf, "extLst")
            if insertion_target is None:
                cell_xf.append(alignment)
            else:
                cell_xf.insert(cell_xf.index(insertion_target), alignment)
        values: tuple[tuple[str, Any], ...] = (
            ("horizontal", formatting.horizontal_alignment),
            ("vertical", formatting.vertical_alignment),
            ("wrapText", formatting.wrap_text),
            ("shrinkToFit", formatting.shrink_to_fit),
            ("textRotation", formatting.text_rotation),
            ("indent", formatting.indent),
        )
        for attribute, value in values:
            if value is None:
                continue
            if isinstance(value, bool):
                alignment.set(attribute, "1" if value else "0")
            else:
                alignment.set(attribute, str(value))
        cell_xf.set("applyAlignment", "1")

    def apply(self, base_index: int, formatting: XlsxCellFormatting) -> int:
        cell_xf = self._indexed_clone(self.cell_xfs, base_index, label="cell format")
        if any(
            value is not None
            for value in (
                formatting.bold,
                formatting.italic,
                formatting.strike,
                formatting.underline,
                formatting.font_name,
                formatting.font_size,
                formatting.font_color,
            )
        ):
            cell_xf.set("fontId", str(self._font_id(cell_xf, formatting)))
            cell_xf.set("applyFont", "1")
        if formatting.fill_color is not None:
            cell_xf.set("fillId", str(self._fill_id(formatting)))
            cell_xf.set("applyFill", "1")
        if any(
            value is not None
            for value in (
                formatting.horizontal_alignment,
                formatting.vertical_alignment,
                formatting.wrap_text,
                formatting.shrink_to_fit,
                formatting.text_rotation,
                formatting.indent,
            )
        ):
            self._alignment(cell_xf, formatting)
        if formatting.number_format is not None:
            cell_xf.set("numFmtId", str(self._number_format_id(formatting.number_format)))
            cell_xf.set("applyNumberFormat", "1")
        return self._deduplicate(self.cell_xfs, cell_xf)

    def serialize(self) -> bytes:
        return etree.tostring(
            self.root,
            encoding="UTF-8",
            xml_declaration=True,
            pretty_print=False,
        )


def _worksheet_cell_map(root: Any) -> dict[str, Any]:
    cells: dict[str, Any] = {}
    for element in root.iter():
        if not isinstance(element.tag, str) or etree.QName(element).localname != "c":
            continue
        coordinate = (element.get("r") or "").upper()
        try:
            min_col, min_row, max_col, max_row = range_boundaries(coordinate)
        except ValueError as exc:
            raise OfficePackageError("XLSX worksheet contains an invalid cell reference") from exc
        if min_col != max_col or min_row != max_row or coordinate in cells:
            raise OfficePackageError("XLSX worksheet contains duplicate or invalid cell references")
        cells[coordinate] = element
    return cells


def _worksheet_non_style_fingerprint(root: Any) -> bytes:
    clone = copy.deepcopy(root)
    for element in clone.iter():
        if isinstance(element.tag, str) and etree.QName(element).localname == "c":
            element.attrib.pop("s", None)
    return hashlib.sha256(_xml_key(clone)).digest()


def _serialize_xml(root: Any) -> bytes:
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, pretty_print=False)


def _serialize_package(package: _EditablePackage, replacements: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, mode="w") as archive:
        archive.comment = package.comment
        for original_info, payload in package.entries:
            info = copy.copy(original_info)
            archive.writestr(info, replacements.get(info.filename, payload))
    result = output.getvalue()
    if len(result) > _MAX_PACKAGE_BYTES:
        raise OfficePackageError("Edited XLSX exceeds the 50 MiB package limit")
    return result


def _semantic_value(value: Any) -> str:
    if isinstance(value, (dt.datetime, dt.date, dt.time, dt.timedelta)):
        return f"{type(value).__name__}:{value!s}"
    return f"{type(value).__name__}:{value!r}"


def _semantic_fingerprint(workbook: Any) -> str:
    digest = hashlib.sha256()
    for worksheet in workbook.worksheets:
        digest.update(f"sheet\0{worksheet.title}\0{worksheet.sheet_state}\0".encode())
        for merged in sorted(str(item) for item in worksheet.merged_cells.ranges):
            digest.update(f"merge\0{merged}\0".encode())
        for coordinate, cell in sorted(worksheet._cells.items()):
            if isinstance(cell, MergedCell):
                continue
            if cell.value is None and cell.hyperlink is None and cell.comment is None:
                continue
            digest.update(f"cell\0{coordinate[0]}\0{coordinate[1]}\0{cell.data_type}\0".encode())
            digest.update(_semantic_value(cell.value).encode("utf-8", errors="surrogatepass"))
            digest.update(b"\0")
            if cell.hyperlink is not None:
                digest.update(f"link\0{cell.hyperlink.target}\0{cell.hyperlink.location}\0{cell.hyperlink.tooltip}\0".encode("utf-8", errors="surrogatepass"))
            if cell.comment is not None:
                digest.update(f"comment\0{cell.comment.author}\0{cell.comment.text}\0".encode("utf-8", errors="surrogatepass"))
    return digest.hexdigest()


def _selector_value_matches(actual: Any, expected: str | int | float | bool) -> bool:
    if isinstance(expected, bool):
        return isinstance(actual, bool) and actual is expected
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and not isinstance(actual, bool) and actual == expected
    return str(actual) == expected


def _matches_selector(cell: Any, selector: XlsxCellSelector) -> bool:
    selector_value = _selector_cell_value(cell)
    if selector.contains_text is not None and selector.contains_text not in str(selector_value or ""):
        return False
    if selector.value_equals is not None and not _selector_value_matches(selector_value, selector.value_equals):
        return False
    if selector.has_formula is not None and (cell.data_type == "f") is not selector.has_formula:
        return False
    if selector.is_blank is not None and (cell.value is None) is not selector.is_blank:
        return False
    if selector.data_type is not None and _cell_data_type(cell) != selector.data_type:
        return False
    if selector.bold is not None and bool(cell.font.bold) is not selector.bold:
        return False
    if selector.italic is not None and bool(cell.font.italic) is not selector.italic:
        return False
    if selector.font_name is not None and (cell.font.name or "").casefold() != selector.font_name.casefold():
        return False
    if selector.font_size is not None and cell.font.sz != selector.font_size:
        return False
    if selector.font_color is not None and _rgb_value(cell.font.color) != f"#{selector.font_color.lstrip('#').upper()}":
        return False
    if selector.fill_color is not None and _rgb_value(cell.fill.fgColor) != f"#{selector.fill_color.lstrip('#').upper()}":
        return False
    if selector.horizontal_alignment is not None:
        actual_alignment = cell.alignment.horizontal or "general"
        if actual_alignment != selector.horizontal_alignment:
            return False
    return selector.number_format is None or cell.number_format == selector.number_format


def _argb(value: str) -> str:
    return f"FF{value.lstrip('#').upper()}"


def _apply_formatting(cell: Any, formatting: XlsxCellFormatting) -> None:
    if any(
        value is not None
        for value in (
            formatting.bold,
            formatting.italic,
            formatting.strike,
            formatting.underline,
            formatting.font_name,
            formatting.font_size,
            formatting.font_color,
        )
    ):
        font = copy.copy(cell.font)
        if formatting.bold is not None:
            font.bold = formatting.bold
        if formatting.italic is not None:
            font.italic = formatting.italic
        if formatting.strike is not None:
            font.strike = formatting.strike
        if formatting.underline is not None:
            font.underline = None if formatting.underline == "none" else formatting.underline
        if formatting.font_name is not None:
            font.name = formatting.font_name
        if formatting.font_size is not None:
            font.sz = formatting.font_size
        if formatting.font_color is not None:
            font.color = Color(rgb=_argb(formatting.font_color))
        cell.font = font

    if formatting.fill_color is not None:
        cell.fill = PatternFill(fill_type="solid", fgColor=_argb(formatting.fill_color))

    if any(
        value is not None
        for value in (
            formatting.horizontal_alignment,
            formatting.vertical_alignment,
            formatting.wrap_text,
            formatting.shrink_to_fit,
            formatting.text_rotation,
            formatting.indent,
        )
    ):
        alignment = copy.copy(cell.alignment)
        if formatting.horizontal_alignment is not None:
            alignment.horizontal = formatting.horizontal_alignment
        if formatting.vertical_alignment is not None:
            alignment.vertical = formatting.vertical_alignment
        if formatting.wrap_text is not None:
            alignment.wrap_text = formatting.wrap_text
        if formatting.shrink_to_fit is not None:
            alignment.shrink_to_fit = formatting.shrink_to_fit
        if formatting.text_rotation is not None:
            alignment.text_rotation = formatting.text_rotation
        if formatting.indent is not None:
            alignment.indent = formatting.indent
        cell.alignment = alignment

    if formatting.number_format is not None:
        cell.number_format = formatting.number_format


def _selected_cells(
    worksheet: Any,
    selector: XlsxCellSelector,
    *,
    declared_coordinates: set[str],
) -> list[Any]:
    cells: list[Any] = []
    seen: set[tuple[int, int]] = set()
    for cell_range in selector.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                coordinate = (row, column)
                if coordinate in seen:
                    continue
                seen.add(coordinate)
                if len(seen) > _MAX_EDIT_CELLS:
                    raise OfficeOperationError(f"An XLSX operation may select at most {_MAX_EDIT_CELLS:,} cells")
                cell_coordinate = f"{get_column_letter(column)}{row}"
                cell = worksheet._cells.get(coordinate)
                if cell_coordinate in declared_coordinates and cell is not None and not isinstance(cell, MergedCell):
                    cells.append(cell)
    return cells


def edit_xlsx(
    data: bytes,
    operations: list[XlsxEditOperation],
    *,
    receipt_trace: OfficeEditTrace | None = None,
) -> tuple[bytes, list[dict[str, Any]]]:
    """Patch worksheet style references and styles.xml as one transaction."""
    if not operations:
        raise OfficeOperationError("At least one Office edit operation is required")
    package = _load_editable_package(data)
    styles_payload = package.payloads.get(_STYLES_XML)
    if styles_payload is None:
        raise OfficeOperationError("XLSX formatting requires an existing workbook styles part")
    style_editor = _StyleEditor(styles_payload)
    workbook = _open_workbook(data)
    worksheet_roots: dict[str, Any] = {}
    worksheet_cells: dict[str, dict[str, Any]] = {}
    worksheet_fingerprints: dict[str, bytes] = {}
    total_match_count = 0
    try:
        reports: list[dict[str, Any]] = []
        for operation_index, operation in enumerate(operations, start=1):
            if not isinstance(operation, XlsxCellFormatOperation):
                raise OfficeOperationError(f"Unsupported Office operation at position {operation_index}")
            if receipt_trace is not None:
                receipt_trace.start_operation(
                    operation_index,
                    operation.type,
                )
            worksheet = _find_sheet(workbook, operation.cells.sheet_name)
            sheet_part = _worksheet_part_for_name(package, worksheet.title)
            if sheet_part not in worksheet_roots:
                worksheet_root = _safe_xml_root(
                    package.payloads[sheet_part],
                    label=sheet_part,
                    limit=_MAX_WORKSHEET_XML_BYTES,
                )
                worksheet_roots[sheet_part] = worksheet_root
                worksheet_cells[sheet_part] = _worksheet_cell_map(worksheet_root)
                worksheet_fingerprints[sheet_part] = _worksheet_non_style_fingerprint(worksheet_root)
            cell_elements = worksheet_cells[sheet_part]
            matched: list[str] = []
            for cell in _selected_cells(
                worksheet,
                operation.cells,
                declared_coordinates=set(cell_elements),
            ):
                if not _matches_selector(cell, operation.cells):
                    continue
                cell_element = cell_elements[cell.coordinate]
                try:
                    base_style_index = int(cell_element.get("s", "0"))
                except ValueError as exc:
                    raise OfficePackageError("XLSX cell references an invalid style index") from exc
                new_style_index = style_editor.apply(base_style_index, operation.formatting)
                if new_style_index == 0:
                    cell_element.attrib.pop("s", None)
                else:
                    cell_element.set("s", str(new_style_index))
                _apply_formatting(cell, operation.formatting)
                matched.append(cell.coordinate)
                if operation.occurrence == "first":
                    break
            if operation.require_match and not matched:
                raise OfficeOperationError(f"Operation {operation_index} matched no targets; no document changes were written")
            total_match_count += len(matched)
            if receipt_trace is not None:
                for coordinate in matched:
                    receipt_trace.add_target(
                        operation_index,
                        path=xlsx_cell_path(worksheet.title, coordinate),
                        kind="xlsx_cell_format",
                        locator=(worksheet.title, coordinate),
                    )
                receipt_trace.finish_operation(
                    operation_index,
                    match_count=len(matched),
                )
            reports.append(
                {
                    "operation": operation_index,
                    "type": operation.type,
                    "occurrence": operation.occurrence,
                    "sheet_name": worksheet.title,
                    "match_count": len(matched),
                    "sample_coordinates": matched[:20],
                    "sample_truncated": len(matched) > 20,
                }
            )
    finally:
        workbook.close()

    if total_match_count == 0:
        return data, reports
    replacements = {_STYLES_XML: style_editor.serialize()}
    for sheet_part, worksheet_root in worksheet_roots.items():
        if _worksheet_non_style_fingerprint(worksheet_root) != worksheet_fingerprints[sheet_part]:
            raise OfficeOperationError("XLSX formatting changed worksheet content; no output was written")
        replacements[sheet_part] = _serialize_xml(worksheet_root)
    result = _serialize_package(package, replacements)
    enforce_package_preservation(
        data,
        result,
        changed_parts=replacements,
        label="XLSX",
    )
    validate_xlsx(result)
    return result, reports


def validate_xlsx(data: bytes) -> dict[str, Any]:
    """Validate the supported XLSX package and parser invariants."""
    package_info = _load_package_info(data)
    workbook = _open_workbook(data, read_only=True)
    try:
        return {
            "valid": True,
            "format": "xlsx",
            "entry_count": package_info.entry_count,
            "sheet_count": len(workbook.sheetnames),
            "declared_cell_count": package_info.declared_cell_count,
            "risky_features": list(package_info.risky_features),
        }
    finally:
        workbook.close()


def validate_xlsx_renderable(data: bytes) -> dict[str, Any]:
    """Validate XLSX and reject active content before Office conversion."""
    validation = validate_xlsx(data)
    blocked = [feature for feature in validation["risky_features"] if feature in _RENDER_BLOCKED_FEATURES]
    if blocked:
        raise OfficeOperationError(f"XLSX rendering is disabled for workbooks containing {', '.join(blocked)}")
    return validation

"""Versioned semantic intent and deterministic native PPTX generation."""

from __future__ import annotations

import hashlib
import json
import math
import re
import zipfile
from collections import Counter
from collections.abc import Mapping, Sequence
from datetime import datetime
from io import BytesIO
from typing import Annotated, Any, Literal

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from pydantic import BaseModel, ConfigDict, Field, StringConstraints, field_validator, model_validator

from .errors import OfficeOperationError, OfficePackageError
from .pptx import inspect_pptx

PRESENTATION_INTENT_SCHEMA = "vassilflow.office.presentation_intent.v1"
PRESENTATION_GENERATION_RECEIPT_SCHEMA = "vassilflow.office.presentation_generation_receipt.v1"
PRESENTATION_COMPILER_ID = "vassilflow-native-pptx"
PRESENTATION_COMPILER_VERSION = 1

_MAX_SLIDES = 20
_MAX_ELEMENTS_PER_SLIDE = 12
_MAX_TOTAL_TEXT_CHARS = 48_000
_MAX_IMAGE_BYTES = 20 * 1024 * 1024
_MAX_IMAGE_DIMENSION = 12_000
_MAX_IMAGE_PIXELS = 40_000_000
_FIXED_PACKAGE_TIMESTAMP = (2000, 1, 1, 0, 0, 0)
_FIXED_CORE_TIMESTAMP = datetime(2000, 1, 1)
_HEX_RE = re.compile(r"^[0-9A-F]{6}$")

StableIntentId = Annotated[
    str,
    StringConstraints(
        strip_whitespace=True,
        min_length=1,
        max_length=48,
        pattern=r"^[a-z][a-z0-9_-]*$",
    ),
]
PresentationLayout = Literal[
    "title",
    "title_content",
    "two_column",
    "picture_caption",
    "closing",
]
ThemeColorToken = Literal[
    "background",
    "surface",
    "text",
    "muted",
    "accent",
    "accent_alt",
]
FontFamily = Literal[
    "Aptos",
    "Arial",
    "Calibri",
    "Georgia",
    "Liberation Sans",
    "Noto Sans",
]


def _canonical_json_bytes(value: object) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _normalized_hex(value: str) -> str:
    normalized = value.strip().removeprefix("#").upper()
    if _HEX_RE.fullmatch(normalized) is None:
        raise ValueError("colors must use six hexadecimal RGB digits")
    return normalized


class PresentationTheme(BaseModel):
    """Bounded tokens consumed by the compiler, not direct object formatting."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    background: str = "F7F8FA"
    surface: str = "FFFFFF"
    text: str = "17212B"
    muted: str = "56616F"
    accent: str = "1F6FEB"
    accent_alt: str = "0F766E"
    heading_font: FontFamily = "Arial"
    body_font: FontFamily = "Arial"

    @field_validator(
        "background",
        "surface",
        "text",
        "muted",
        "accent",
        "accent_alt",
        mode="before",
    )
    @classmethod
    def normalize_color(cls, value: object) -> str:
        if not isinstance(value, str):
            raise ValueError("theme colors must be strings")
        return _normalized_hex(value)


class PresentationTextElement(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    kind: Literal["text"] = "text"
    id: StableIntentId
    role: Literal["title", "subtitle", "content", "left", "right", "caption"]
    text: str | None = Field(default=None, max_length=4_000)
    bullets: list[str] = Field(default_factory=list, max_length=12)
    emphasis: Literal["regular", "strong"] = "regular"

    @field_validator("text")
    @classmethod
    def validate_text(cls, value: str | None) -> str | None:
        if value is None:
            return None
        value = value.strip()
        if not value:
            raise ValueError("text cannot be blank")
        if any(character in value for character in ("\x00", "\x0b", "\x0c")):
            raise ValueError("text contains unsupported control characters")
        return value

    @field_validator("bullets")
    @classmethod
    def validate_bullets(cls, values: list[str]) -> list[str]:
        normalized: list[str] = []
        for value in values:
            if not isinstance(value, str):
                raise ValueError("bullet values must be strings")
            item = value.strip()
            if not item or len(item) > 500:
                raise ValueError("each bullet must contain 1 through 500 characters")
            if any(character in item for character in ("\x00", "\x0b", "\x0c")):
                raise ValueError("bullet text contains unsupported control characters")
            normalized.append(item)
        return normalized

    @model_validator(mode="after")
    def validate_content(self) -> PresentationTextElement:
        if (self.text is None) == (not self.bullets):
            raise ValueError("text elements require exactly one of text or bullets")
        if self.role in {"title", "subtitle", "caption"} and self.bullets:
            raise ValueError(f"{self.role} text cannot use bullets")
        return self

    def paragraphs(self) -> tuple[str, ...]:
        if self.bullets:
            return tuple(self.bullets)
        return tuple((self.text or "").splitlines()) or (self.text or "",)


class PresentationImageElement(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    kind: Literal["image"] = "image"
    id: StableIntentId
    role: Literal["content", "left", "right", "media"]
    source_path: str = Field(
        min_length=1,
        max_length=1_024,
        pattern=r"^/mnt/user-data/",
    )
    alt_text: str = Field(min_length=1, max_length=512)
    fit: Literal["contain", "cover"] = "cover"

    @field_validator("alt_text")
    @classmethod
    def normalize_alt_text(cls, value: str) -> str:
        normalized = value.strip()
        if not normalized:
            raise ValueError("image alt text cannot be blank")
        return normalized


class PresentationShapeElement(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    kind: Literal["shape"] = "shape"
    id: StableIntentId
    placement: Literal["accent_bar", "top_right", "bottom_left"]
    shape: Literal["rectangle", "rounded_rectangle", "ellipse"] = "rectangle"
    fill: ThemeColorToken = "accent"


PresentationElement = Annotated[
    PresentationTextElement | PresentationImageElement | PresentationShapeElement,
    Field(discriminator="kind"),
]


_LAYOUT_REQUIRED_ROLES: dict[PresentationLayout, frozenset[str]] = {
    "title": frozenset({"title"}),
    "title_content": frozenset({"title", "content"}),
    "two_column": frozenset({"title", "left", "right"}),
    "picture_caption": frozenset({"title", "media", "caption"}),
    "closing": frozenset({"title"}),
}
_LAYOUT_ALLOWED_ROLES: dict[PresentationLayout, frozenset[str]] = {
    "title": frozenset({"title", "subtitle"}),
    "title_content": frozenset({"title", "content"}),
    "two_column": frozenset({"title", "left", "right"}),
    "picture_caption": frozenset({"title", "media", "caption"}),
    "closing": frozenset({"title", "subtitle"}),
}


class PresentationSlideIntent(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    id: StableIntentId
    purpose: str = Field(min_length=1, max_length=160)
    layout: PresentationLayout
    elements: list[PresentationElement] = Field(
        min_length=1,
        max_length=_MAX_ELEMENTS_PER_SLIDE,
    )

    @field_validator("purpose")
    @classmethod
    def normalize_purpose(cls, value: str) -> str:
        return value.strip()

    @model_validator(mode="after")
    def validate_layout_contract(self) -> PresentationSlideIntent:
        ids = [element.id for element in self.elements]
        if len(set(ids)) != len(ids):
            raise ValueError("element IDs must be unique within a slide")

        content_elements = [element for element in self.elements if not isinstance(element, PresentationShapeElement)]
        roles = [element.role for element in content_elements]
        if len(set(roles)) != len(roles):
            raise ValueError("semantic content roles must be unique within a slide")
        required = _LAYOUT_REQUIRED_ROLES[self.layout]
        allowed = _LAYOUT_ALLOWED_ROLES[self.layout]
        if not required.issubset(roles) or not set(roles).issubset(allowed):
            raise ValueError(f"layout {self.layout} requires {sorted(required)} and allows only {sorted(allowed)}")

        title = next(
            (element for element in content_elements if element.role == "title"),
            None,
        )
        if not isinstance(title, PresentationTextElement) or title.text is None:
            raise ValueError("every slide requires one non-bulleted text title")
        if self.layout == "picture_caption":
            media = next(element for element in content_elements if element.role == "media")
            if not isinstance(media, PresentationImageElement):
                raise ValueError("picture_caption requires an image in the media role")

        placements = [element.placement for element in self.elements if isinstance(element, PresentationShapeElement)]
        if len(set(placements)) != len(placements):
            raise ValueError("shape placements must be unique within a slide")
        return self


class PresentationIntent(BaseModel):
    """Model-authored semantic input; geometry remains compiler-owned."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    schema_version: Literal[PRESENTATION_INTENT_SCHEMA] = Field(
        default=PRESENTATION_INTENT_SCHEMA,
        alias="schema",
    )
    title: str = Field(min_length=1, max_length=256)
    aspect_ratio: Literal["16:9", "4:3"] = "16:9"
    theme: PresentationTheme = Field(default_factory=PresentationTheme)
    slides: list[PresentationSlideIntent] = Field(
        min_length=1,
        max_length=_MAX_SLIDES,
    )

    @field_validator("title")
    @classmethod
    def normalize_title(cls, value: str) -> str:
        return value.strip()

    @model_validator(mode="after")
    def validate_deck_contract(self) -> PresentationIntent:
        slide_ids = [slide.id for slide in self.slides]
        if len(set(slide_ids)) != len(slide_ids):
            raise ValueError("slide IDs must be unique")
        element_ids = [element.id for slide in self.slides for element in slide.elements]
        if len(set(element_ids)) != len(element_ids):
            raise ValueError("element IDs must be unique across the presentation")
        total_text = sum(len(element.text or "") + sum(len(item) for item in element.bullets) for slide in self.slides for element in slide.elements if isinstance(element, PresentationTextElement))
        if total_text > _MAX_TOTAL_TEXT_CHARS:
            raise ValueError("presentation text exceeds the generation limit")
        return self

    def canonical_bytes(self) -> bytes:
        return _canonical_json_bytes(
            self.model_dump(
                by_alias=True,
                exclude_none=True,
                mode="json",
            )
        )

    def image_paths(self) -> tuple[str, ...]:
        return tuple(dict.fromkeys(element.source_path for slide in self.slides for element in slide.elements if isinstance(element, PresentationImageElement)))


class _Rect(BaseModel):
    model_config = ConfigDict(frozen=True)

    left: float
    top: float
    width: float
    height: float


_LAYOUT_RECTS: dict[PresentationLayout, dict[str, _Rect]] = {
    "title": {
        "title": _Rect(left=0.10, top=0.31, width=0.80, height=0.20),
        "subtitle": _Rect(left=0.16, top=0.55, width=0.68, height=0.12),
    },
    "title_content": {
        "title": _Rect(left=0.06, top=0.06, width=0.88, height=0.14),
        "content": _Rect(left=0.08, top=0.25, width=0.84, height=0.62),
    },
    "two_column": {
        "title": _Rect(left=0.06, top=0.06, width=0.88, height=0.14),
        "left": _Rect(left=0.06, top=0.25, width=0.42, height=0.62),
        "right": _Rect(left=0.52, top=0.25, width=0.42, height=0.62),
    },
    "picture_caption": {
        "title": _Rect(left=0.06, top=0.06, width=0.88, height=0.14),
        "media": _Rect(left=0.06, top=0.24, width=0.58, height=0.62),
        "caption": _Rect(left=0.68, top=0.24, width=0.26, height=0.62),
    },
    "closing": {
        "title": _Rect(left=0.10, top=0.34, width=0.80, height=0.18),
        "subtitle": _Rect(left=0.18, top=0.56, width=0.64, height=0.11),
    },
}

_TEXT_SIZE_BOUNDS: dict[str, tuple[int, int]] = {
    "title": (48, 24),
    "subtitle": (25, 16),
    "content": (24, 14),
    "left": (22, 14),
    "right": (22, 14),
    "caption": (19, 12),
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _relative_luminance(value: str) -> float:
    channels = [int(value[index : index + 2], 16) / 255 for index in (0, 2, 4)]
    linear = [channel / 12.92 if channel <= 0.04045 else ((channel + 0.055) / 1.055) ** 2.4 for channel in channels]
    return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]


def _contrast_ratio(first: str, second: str) -> float:
    first_luminance = _relative_luminance(first)
    second_luminance = _relative_luminance(second)
    lighter = max(first_luminance, second_luminance)
    darker = min(first_luminance, second_luminance)
    return (lighter + 0.05) / (darker + 0.05)


def _readable_color(preferred: str, background: str) -> str:
    candidates = (preferred, "FFFFFF", "111111")
    return max(candidates, key=lambda candidate: _contrast_ratio(candidate, background))


def _rect_to_emu(
    rect: _Rect,
    *,
    slide_width: int,
    slide_height: int,
) -> tuple[int, int, int, int]:
    return (
        round(slide_width * rect.left),
        round(slide_height * rect.top),
        round(slide_width * rect.width),
        round(slide_height * rect.height),
    )


def _estimate_line_count(text: str, *, font_size: int, width_points: float) -> int:
    capacity = max(1, math.floor(width_points / (font_size * 0.52)))
    return sum(max(1, math.ceil(len(segment) / capacity)) for segment in text.splitlines() or [text])


def _fit_font_size(
    element: PresentationTextElement,
    *,
    width_emu: int,
    height_emu: int,
) -> int:
    maximum, minimum = _TEXT_SIZE_BOUNDS[element.role]
    width_points = width_emu / 12_700
    height_points = height_emu / 12_700
    paragraphs = element.paragraphs()
    for font_size in range(maximum, minimum - 1, -1):
        content_width = width_points - (font_size * 1.1 if element.bullets else 8)
        lines = sum(
            _estimate_line_count(
                paragraph,
                font_size=font_size,
                width_points=max(1, content_width),
            )
            for paragraph in paragraphs
        )
        estimated_height = lines * font_size * 1.18
        estimated_height += max(0, len(paragraphs) - 1) * min(10, font_size * 0.45)
        if estimated_height <= height_points - 8:
            return font_size
    raise OfficeOperationError(f"Presentation text element {element.id} does not fit its semantic layout above the minimum readable size")


def _set_bullet(paragraph: Any) -> None:
    properties = paragraph._p.get_or_add_pPr()  # noqa: SLF001 - generated OOXML only
    for child in list(properties):
        if child.tag.rsplit("}", 1)[-1] in {"buNone", "buChar", "buAutoNum"}:
            properties.remove(child)
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "\u2022")
    properties.append(bullet)


def _authored_name(slide_id: str, element_id: str) -> str:
    return f"VFGEN:{slide_id}:{element_id}"


def _set_shape_name(shape: Any, name: str) -> None:
    shape.name = name


def _set_picture_alt_text(picture: Any, alt_text: str) -> None:
    properties = picture._element.nvPicPr.cNvPr  # noqa: SLF001 - generated OOXML only
    properties.set("descr", alt_text)
    properties.set("title", alt_text)


def _send_shape_to_back(slide: Any, shape: Any) -> None:
    tree = slide.shapes._spTree  # noqa: SLF001 - generated OOXML only
    element = shape._element  # noqa: SLF001 - generated OOXML only
    tree.remove(element)
    tree.insert(2, element)


def _add_text(
    slide: Any,
    element: PresentationTextElement,
    *,
    rect: tuple[int, int, int, int],
    theme: PresentationTheme,
    slide_id: str,
    title_shape: Any,
    centered: bool,
) -> Any:
    left, top, width, height = rect
    shape = title_shape if element.role == "title" else slide.shapes.add_textbox(left, top, width, height)
    shape.left = left
    shape.top = top
    shape.width = width
    shape.height = height
    _set_shape_name(shape, _authored_name(slide_id, element.id))

    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if element.role in {"title", "subtitle"} else MSO_ANCHOR.TOP
    font_size = _fit_font_size(
        element,
        width_emu=width,
        height_emu=height,
    )
    preferred = theme.muted if element.role in {"subtitle", "caption"} else theme.text
    font_color = _readable_color(preferred, theme.background)
    font_family = theme.heading_font if element.role in {"title", "subtitle"} else theme.body_font

    paragraphs = element.paragraphs()
    for index, value in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
        paragraph.line_spacing = 1.12
        paragraph.space_after = Pt(min(10, font_size * 0.45))
        if element.bullets:
            _set_bullet(paragraph)
            paragraph.margin_left = Pt(font_size * 1.05)
            paragraph.indent = Pt(-font_size * 0.45)
        run = paragraph.add_run()
        run.text = value
        run.font.name = font_family
        run.font.size = Pt(font_size)
        run.font.bold = element.role == "title" or element.emphasis == "strong"
        run.font.color.rgb = _rgb(font_color)
    return shape


def _validated_image(data: bytes, *, element_id: str) -> tuple[str, int, int]:
    if not data or len(data) > _MAX_IMAGE_BYTES:
        raise OfficeOperationError(f"Presentation image {element_id} must contain 1 through {_MAX_IMAGE_BYTES} bytes")
    try:
        with Image.open(BytesIO(data)) as image:
            image.load()
            image_format = (image.format or "").upper()
            width, height = image.size
            if image_format not in {"PNG", "JPEG"}:
                raise OfficeOperationError(f"Presentation image {element_id} must be PNG or baseline JPEG")
            if image_format == "JPEG" and (image.info.get("progressive") or image.info.get("progression")):
                raise OfficeOperationError(f"Presentation image {element_id} must use baseline JPEG encoding")
            if width < 1 or height < 1 or width > _MAX_IMAGE_DIMENSION or height > _MAX_IMAGE_DIMENSION or width * height > _MAX_IMAGE_PIXELS:
                raise OfficeOperationError(f"Presentation image {element_id} dimensions exceed the generation limit")
            return image_format.lower(), width, height
    except OfficeOperationError:
        raise
    except (OSError, UnidentifiedImageError) as exc:
        raise OfficeOperationError(f"Presentation image {element_id} is invalid") from exc


def _add_image(
    slide: Any,
    element: PresentationImageElement,
    *,
    rect: tuple[int, int, int, int],
    data: bytes,
    dimensions: tuple[int, int],
    slide_id: str,
) -> Any:
    left, top, width, height = rect
    image_width, image_height = dimensions
    frame_aspect = width / height
    image_aspect = image_width / image_height
    if element.fit == "contain":
        if image_aspect >= frame_aspect:
            rendered_width = width
            rendered_height = round(width / image_aspect)
        else:
            rendered_height = height
            rendered_width = round(height * image_aspect)
        picture = slide.shapes.add_picture(
            BytesIO(data),
            left + round((width - rendered_width) / 2),
            top + round((height - rendered_height) / 2),
            rendered_width,
            rendered_height,
        )
    else:
        picture = slide.shapes.add_picture(
            BytesIO(data),
            left,
            top,
            width,
            height,
        )
        if image_aspect > frame_aspect:
            crop = (1 - frame_aspect / image_aspect) / 2
            picture.crop_left = crop
            picture.crop_right = crop
        elif image_aspect < frame_aspect:
            crop = (1 - image_aspect / frame_aspect) / 2
            picture.crop_top = crop
            picture.crop_bottom = crop
    _set_shape_name(picture, _authored_name(slide_id, element.id))
    _set_picture_alt_text(picture, element.alt_text)
    return picture


def _add_decoration(
    slide: Any,
    element: PresentationShapeElement,
    *,
    slide_width: int,
    slide_height: int,
    theme: PresentationTheme,
    slide_id: str,
) -> Any:
    placements = {
        "accent_bar": _Rect(left=0.07, top=0.90, width=0.25, height=0.012),
        "top_right": _Rect(left=0.88, top=0.04, width=0.08, height=0.10),
        "bottom_left": _Rect(left=0.02, top=0.82, width=0.11, height=0.15),
    }
    shape_types = {
        "rectangle": MSO_SHAPE.RECTANGLE,
        "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
        "ellipse": MSO_SHAPE.OVAL,
    }
    left, top, width, height = _rect_to_emu(
        placements[element.placement],
        slide_width=slide_width,
        slide_height=slide_height,
    )
    shape = slide.shapes.add_shape(
        shape_types[element.shape],
        left,
        top,
        width,
        height,
    )
    _set_shape_name(shape, _authored_name(slide_id, element.id))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(getattr(theme, element.fill))
    shape.line.fill.background()
    _send_shape_to_back(slide, shape)
    return shape


def _canonicalize_pptx_package(data: bytes) -> bytes:
    output = BytesIO()
    try:
        with zipfile.ZipFile(BytesIO(data), "r") as source:
            names = source.namelist()
            if len(set(names)) != len(names):
                raise OfficePackageError("Generated PPTX contains duplicate package entries")
            with zipfile.ZipFile(
                output,
                "w",
                compression=zipfile.ZIP_DEFLATED,
                compresslevel=9,
                strict_timestamps=True,
            ) as destination:
                for name in sorted(names):
                    payload = source.read(name)
                    info = zipfile.ZipInfo(name, date_time=_FIXED_PACKAGE_TIMESTAMP)
                    info.compress_type = zipfile.ZIP_DEFLATED
                    info.create_system = 3
                    info.external_attr = 0o600 << 16
                    destination.writestr(
                        info,
                        payload,
                        compress_type=zipfile.ZIP_DEFLATED,
                        compresslevel=9,
                    )
    except (OSError, zipfile.BadZipFile) as exc:
        raise OfficePackageError("Generated PPTX package could not be canonicalized") from exc
    return output.getvalue()


def _object_records(
    document: bytes,
    *,
    intent: PresentationIntent,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    inspection = inspect_pptx(
        document,
        start_slide=1,
        max_slides=len(intent.slides),
        include_formatting=False,
    )
    inspected_slides = inspection.get("slides")
    if not isinstance(inspected_slides, list) or len(inspected_slides) != len(intent.slides):
        raise OfficePackageError("Generated PPTX inspection did not return every slide")

    all_objects: list[dict[str, Any]] = []
    slide_records: list[dict[str, Any]] = []
    for index, (slide_intent, inspected_slide) in enumerate(
        zip(intent.slides, inspected_slides, strict=True),
        start=1,
    ):
        objects = inspected_slide.get("objects")
        if not isinstance(objects, list):
            raise OfficePackageError("Generated PPTX inspection omitted object identities")
        by_name = {item.get("name"): item for item in objects if isinstance(item, dict) and isinstance(item.get("name"), str)}
        mapped: list[dict[str, Any]] = []
        for element in slide_intent.elements:
            authored_name = _authored_name(slide_intent.id, element.id)
            inspected = by_name.get(authored_name)
            if not isinstance(inspected, dict):
                raise OfficePackageError(f"Generated PPTX object {element.id} could not be resolved after compilation")
            path = inspected.get("path")
            kind = inspected.get("kind")
            identity_source = inspected.get("identity_source")
            if not isinstance(path, str) or not isinstance(kind, str) or identity_source != "cNvPr.id":
                raise OfficePackageError(f"Generated PPTX object {element.id} lacks a stable authored identity")
            record = {
                "slide_id": slide_intent.id,
                "element_id": element.id,
                "element_kind": element.kind,
                "role": (element.placement if isinstance(element, PresentationShapeElement) else element.role),
                "object_path": path,
                "object_kind": kind,
                "authored_name": authored_name,
            }
            mapped.append(record)
            all_objects.append(record)
        slide_records.append(
            {
                "slide_id": slide_intent.id,
                "slide_index": index,
                "slide_path": inspected_slide.get("path"),
                "layout": slide_intent.layout,
                "purpose": slide_intent.purpose,
                "object_count": len(mapped),
                "objects": mapped,
            }
        )
    return slide_records, all_objects


def _embedded_asset_records(
    document: bytes,
    *,
    intent: PresentationIntent,
) -> list[dict[str, Any]]:
    """Rebuild bounded image evidence from the media embedded in the package."""

    try:
        presentation = Presentation(BytesIO(document))
    except Exception as exc:
        raise OfficePackageError("Generated PPTX media could not be inspected") from exc

    records: dict[str, dict[str, Any]] = {}
    for slide_intent, slide in zip(intent.slides, presentation.slides, strict=True):
        by_name = {shape.name: shape for shape in slide.shapes}
        for element in slide_intent.elements:
            if not isinstance(element, PresentationImageElement):
                continue
            shape = by_name.get(_authored_name(slide_intent.id, element.id))
            image = getattr(shape, "image", None)
            if image is None:
                raise OfficePackageError(f"Generated PPTX image {element.id} could not be resolved")
            payload = image.blob
            image_format, width, height = _validated_image(
                payload,
                element_id=element.id,
            )
            properties = shape._element.nvPicPr.cNvPr  # noqa: SLF001 - read-only OOXML evidence
            if properties.get("descr") != element.alt_text or properties.get("title") != element.alt_text:
                raise OfficePackageError(f"Generated PPTX image {element.id} alt text is inconsistent")
            records[element.id] = {
                "element_id": element.id,
                "sha256": _sha256(payload),
                "size_bytes": len(payload),
                "format": image_format,
                "width_px": width,
                "height_px": height,
                "alt_text_sha256": _sha256(element.alt_text.encode("utf-8")),
            }
    return [records[key] for key in sorted(records)]


def _presentation_dimensions(aspect_ratio: str) -> tuple[int, int]:
    if aspect_ratio == "4:3":
        return int(Inches(10)), int(Inches(7.5))
    return int(Inches(13.333333)), int(Inches(7.5))


def compile_presentation(
    intent: PresentationIntent,
    *,
    image_assets: Mapping[str, bytes] | None = None,
) -> tuple[bytes, dict[str, Any]]:
    """Compile semantic intent to deterministic editable PPTX bytes and a receipt."""

    assets = dict(image_assets or {})
    expected_paths = set(intent.image_paths())
    if set(assets) != expected_paths:
        missing = sorted(expected_paths - set(assets))
        unexpected = sorted(set(assets) - expected_paths)
        raise OfficeOperationError(f"Presentation image assets do not match intent; missing={missing}, unexpected={unexpected}")

    image_evidence: dict[str, dict[str, Any]] = {}
    for slide in intent.slides:
        for element in slide.elements:
            if not isinstance(element, PresentationImageElement):
                continue
            payload = assets[element.source_path]
            image_format, width, height = _validated_image(
                payload,
                element_id=element.id,
            )
            image_evidence[element.id] = {
                "element_id": element.id,
                "sha256": _sha256(payload),
                "size_bytes": len(payload),
                "format": image_format,
                "width_px": width,
                "height_px": height,
                "alt_text_sha256": _sha256(element.alt_text.encode("utf-8")),
            }

    presentation = Presentation()
    slide_width, slide_height = _presentation_dimensions(intent.aspect_ratio)
    presentation.slide_width = slide_width
    presentation.slide_height = slide_height
    properties = presentation.core_properties
    properties.title = intent.title
    properties.subject = "Editable presentation generated by VassilFlow"
    properties.author = "VassilFlow"
    properties.last_modified_by = "VassilFlow"
    properties.comments = f"{PRESENTATION_COMPILER_ID} v{PRESENTATION_COMPILER_VERSION}"
    properties.created = _FIXED_CORE_TIMESTAMP
    properties.modified = _FIXED_CORE_TIMESTAMP
    properties.revision = 1

    title_layout = presentation.slide_layouts[5]
    for slide_intent in intent.slides:
        slide = presentation.slides.add_slide(title_layout)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = _rgb(intent.theme.background)
        title_shape = slide.shapes.title
        if title_shape is None:
            raise OfficePackageError("PPTX compiler title layout has no title placeholder")

        for element in slide_intent.elements:
            if isinstance(element, PresentationShapeElement):
                _add_decoration(
                    slide,
                    element,
                    slide_width=slide_width,
                    slide_height=slide_height,
                    theme=intent.theme,
                    slide_id=slide_intent.id,
                )

        for element in slide_intent.elements:
            if isinstance(element, PresentationShapeElement):
                continue
            semantic_rect = _LAYOUT_RECTS[slide_intent.layout][element.role]
            rect = _rect_to_emu(
                semantic_rect,
                slide_width=slide_width,
                slide_height=slide_height,
            )
            if isinstance(element, PresentationTextElement):
                _add_text(
                    slide,
                    element,
                    rect=rect,
                    theme=intent.theme,
                    slide_id=slide_intent.id,
                    title_shape=title_shape,
                    centered=slide_intent.layout in {"title", "closing"},
                )
            else:
                evidence = image_evidence[element.id]
                _add_image(
                    slide,
                    element,
                    rect=rect,
                    data=assets[element.source_path],
                    dimensions=(evidence["width_px"], evidence["height_px"]),
                    slide_id=slide_intent.id,
                )

    stream = BytesIO()
    presentation.save(stream)
    document = _canonicalize_pptx_package(stream.getvalue())
    intent_bytes = intent.canonical_bytes()
    slide_records, object_records = _object_records(document, intent=intent)
    mapping_sha256 = _sha256(_canonical_json_bytes(slide_records))
    receipt = {
        "schema": PRESENTATION_GENERATION_RECEIPT_SCHEMA,
        "compiler": {
            "id": PRESENTATION_COMPILER_ID,
            "version": PRESENTATION_COMPILER_VERSION,
        },
        "intent": {
            "schema": PRESENTATION_INTENT_SCHEMA,
            "sha256": _sha256(intent_bytes),
            "size_bytes": len(intent_bytes),
        },
        "output": {
            "format": "pptx",
            "sha256": _sha256(document),
            "size_bytes": len(document),
        },
        "presentation": {
            "title": intent.title,
            "aspect_ratio": intent.aspect_ratio,
            "slide_width_emu": slide_width,
            "slide_height_emu": slide_height,
            "slide_count": len(intent.slides),
            "object_count": len(object_records),
            "mapping_sha256": mapping_sha256,
        },
        "slides": slide_records,
        "assets": [image_evidence[key] for key in sorted(image_evidence)],
        "bounds": {
            "slides_returned": len(slide_records),
            "slides_truncated": False,
            "objects_returned": len(object_records),
            "objects_truncated": False,
            "assets_returned": len(image_evidence),
            "assets_truncated": False,
        },
    }
    return document, receipt


def validate_generation_evidence(
    *,
    artifact: bytes,
    intent: Mapping[str, Any],
    receipt: Mapping[str, Any],
    preflight: Mapping[str, Any],
) -> None:
    """Verify generated revision evidence before durable publication or readback."""

    try:
        parsed_intent = PresentationIntent.model_validate(intent)
    except Exception as exc:
        raise OfficePackageError("Presentation generation intent is invalid") from exc
    intent_bytes = parsed_intent.canonical_bytes()
    artifact_identity = {
        "format": "pptx",
        "sha256": _sha256(artifact),
        "size_bytes": len(artifact),
    }
    if (
        set(receipt)
        != {
            "schema",
            "compiler",
            "intent",
            "output",
            "presentation",
            "slides",
            "assets",
            "bounds",
        }
        or receipt.get("schema") != PRESENTATION_GENERATION_RECEIPT_SCHEMA
        or receipt.get("compiler")
        != {
            "id": PRESENTATION_COMPILER_ID,
            "version": PRESENTATION_COMPILER_VERSION,
        }
        or receipt.get("intent")
        != {
            "schema": PRESENTATION_INTENT_SCHEMA,
            "sha256": _sha256(intent_bytes),
            "size_bytes": len(intent_bytes),
        }
        or receipt.get("output") != artifact_identity
    ):
        raise OfficePackageError("Presentation generation receipt identity is invalid")
    presentation = receipt.get("presentation")
    bounds = receipt.get("bounds")
    slides = receipt.get("slides")
    assets = receipt.get("assets")
    expected_slides, expected_objects = _object_records(
        artifact,
        intent=parsed_intent,
    )
    expected_assets = _embedded_asset_records(
        artifact,
        intent=parsed_intent,
    )
    slide_width, slide_height = _presentation_dimensions(parsed_intent.aspect_ratio)
    expected_presentation = {
        "title": parsed_intent.title,
        "aspect_ratio": parsed_intent.aspect_ratio,
        "slide_width_emu": slide_width,
        "slide_height_emu": slide_height,
        "slide_count": len(parsed_intent.slides),
        "object_count": len(expected_objects),
        "mapping_sha256": _sha256(_canonical_json_bytes(expected_slides)),
    }
    expected_bounds = {
        "slides_returned": len(expected_slides),
        "slides_truncated": False,
        "objects_returned": len(expected_objects),
        "objects_truncated": False,
        "assets_returned": len(expected_assets),
        "assets_truncated": False,
    }
    if (
        not isinstance(presentation, Mapping)
        or dict(presentation) != expected_presentation
        or not isinstance(slides, Sequence)
        or isinstance(slides, (str, bytes))
        or list(slides) != expected_slides
        or not isinstance(assets, Sequence)
        or isinstance(assets, (str, bytes))
        or list(assets) != expected_assets
        or not isinstance(bounds, Mapping)
        or dict(bounds) != expected_bounds
    ):
        raise OfficePackageError("Presentation generation receipt bounds are invalid")
    summary = preflight.get("summary")
    findings = preflight.get("findings")
    severity_counts = summary.get("findings_by_severity") if isinstance(summary, Mapping) else None
    code_counts = summary.get("findings_by_code") if isinstance(summary, Mapping) else None
    image_quality = preflight.get("image_quality")
    image_assessments = image_quality.get("assessments") if isinstance(image_quality, Mapping) else None
    expected_severity_counts: Counter[str] = Counter()
    expected_code_counts: Counter[str] = Counter()
    findings_valid = isinstance(findings, Sequence) and not isinstance(findings, (str, bytes))
    if findings_valid:
        for finding in findings:
            if not isinstance(finding, Mapping):
                findings_valid = False
                break
            severity = finding.get("severity")
            code = finding.get("code")
            if severity not in {"info", "warning", "error"} or not isinstance(code, str) or not code:
                findings_valid = False
                break
            expected_severity_counts[severity] += 1
            expected_code_counts[code] += 1
    if (
        preflight.get("schema") != "vassilflow.office.pptx.quality_preflight.v1"
        or preflight.get("source_sha256") != artifact_identity["sha256"]
        or preflight.get("source_size_bytes") != artifact_identity["size_bytes"]
        or not isinstance(summary, Mapping)
        or summary.get("slide_count") != len(parsed_intent.slides)
        or summary.get("object_count") != len(expected_objects)
        or summary.get("picture_count") != len(expected_assets)
        or summary.get("findings_truncated") is not False
        or not findings_valid
        or summary.get("findings_returned") != len(findings)
        or summary.get("finding_count") != len(findings)
        or not isinstance(severity_counts, Mapping)
        or dict(severity_counts) != dict(sorted(expected_severity_counts.items()))
        or expected_severity_counts.get("error", 0) != 0
        or not isinstance(code_counts, Mapping)
        or dict(code_counts) != dict(sorted(expected_code_counts.items()))
        or not isinstance(image_quality, Mapping)
        or image_quality.get("assessment_count") != len(expected_assets)
        or image_quality.get("assessments_returned") != len(expected_assets)
        or image_quality.get("assessments_truncated") is not False
        or not isinstance(image_assessments, Sequence)
        or isinstance(image_assessments, (str, bytes))
        or len(image_assessments) != len(expected_assets)
        or any(not isinstance(item, Mapping) for item in image_assessments)
    ):
        raise OfficePackageError("Presentation generation preflight evidence is invalid")

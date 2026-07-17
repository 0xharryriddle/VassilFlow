from __future__ import annotations

import copy
import errno
import hashlib
import io
import json
from pathlib import Path

import pytest
from langchain.tools import ToolRuntime
from PIL import Image
from pptx import Presentation
from pydantic import ValidationError

from vassilflow.community.office.engine import office_engine
from vassilflow.community.office.errors import (
    OfficeOperationError,
    OfficePackageError,
    OfficeRevisionIntegrityError,
)
from vassilflow.community.office.generation import (
    PRESENTATION_GENERATION_RECEIPT_SCHEMA,
    PRESENTATION_INTENT_SCHEMA,
    PresentationIntent,
    compile_presentation,
    validate_generation_evidence,
)
from vassilflow.community.office.models import PptxTextReplacement
from vassilflow.community.office.revisions import OfficeRevisionStore
from vassilflow.community.office.tools import office_generate_tool

_IMAGE_PATH = "/mnt/user-data/uploads/hero.png"


def _png(*, width: int = 1600, height: int = 900) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (width, height), (42, 112, 164)).save(
        output,
        format="PNG",
        optimize=False,
    )
    return output.getvalue()


def _jpeg(*, progressive: bool) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (1200, 800), (42, 112, 164)).save(
        output,
        format="JPEG",
        quality=90,
        progressive=progressive,
        optimize=False,
    )
    return output.getvalue()


def _intent_payload(*, aspect_ratio: str = "16:9") -> dict:
    return {
        "schema": PRESENTATION_INTENT_SCHEMA,
        "title": "Quarterly Update",
        "aspect_ratio": aspect_ratio,
        "theme": {
            "background": "F7F8FA",
            "surface": "FFFFFF",
            "text": "17212B",
            "muted": "56616F",
            "accent": "1F6FEB",
            "accent_alt": "0F766E",
            "heading_font": "Arial",
            "body_font": "Georgia",
        },
        "slides": [
            {
                "id": "opening",
                "purpose": "Open the review",
                "layout": "title",
                "elements": [
                    {
                        "kind": "shape",
                        "id": "opening-accent",
                        "placement": "accent_bar",
                        "shape": "rectangle",
                        "fill": "accent",
                    },
                    {
                        "kind": "text",
                        "id": "opening-title",
                        "role": "title",
                        "text": "Quarterly Update",
                    },
                    {
                        "kind": "text",
                        "id": "opening-subtitle",
                        "role": "subtitle",
                        "text": "Decisions, evidence, and next actions",
                    },
                ],
            },
            {
                "id": "summary",
                "purpose": "Summarize outcomes",
                "layout": "title_content",
                "elements": [
                    {
                        "kind": "text",
                        "id": "summary-title",
                        "role": "title",
                        "text": "Executive summary",
                    },
                    {
                        "kind": "text",
                        "id": "summary-content",
                        "role": "content",
                        "bullets": [
                            "Customer retention improved across core accounts",
                            "Delivery reliability remained above target",
                            "Two decisions are required for the next quarter",
                        ],
                    },
                    {
                        "kind": "shape",
                        "id": "summary-mark",
                        "placement": "top_right",
                        "shape": "ellipse",
                        "fill": "accent_alt",
                    },
                ],
            },
            {
                "id": "evidence",
                "purpose": "Pair visual evidence with interpretation",
                "layout": "two_column",
                "elements": [
                    {
                        "kind": "text",
                        "id": "evidence-title",
                        "role": "title",
                        "text": "Evidence and interpretation",
                    },
                    {
                        "kind": "image",
                        "id": "evidence-image",
                        "role": "left",
                        "source_path": _IMAGE_PATH,
                        "alt_text": "Blue evidence panel used in the quarterly review",
                        "fit": "cover",
                    },
                    {
                        "kind": "text",
                        "id": "evidence-notes",
                        "role": "right",
                        "bullets": [
                            "The signal is sustained across three reporting periods",
                            "No material variance appears in the validation sample",
                        ],
                    },
                ],
            },
            {
                "id": "detail",
                "purpose": "Explain one visual",
                "layout": "picture_caption",
                "elements": [
                    {
                        "kind": "text",
                        "id": "detail-title",
                        "role": "title",
                        "text": "What changed",
                    },
                    {
                        "kind": "image",
                        "id": "detail-image",
                        "role": "media",
                        "source_path": _IMAGE_PATH,
                        "alt_text": "Blue detail panel supporting the change explanation",
                        "fit": "contain",
                    },
                    {
                        "kind": "text",
                        "id": "detail-caption",
                        "role": "caption",
                        "text": "The operating model now protects the highest-value path.",
                    },
                ],
            },
            {
                "id": "closing",
                "purpose": "Close with a decision",
                "layout": "closing",
                "elements": [
                    {
                        "kind": "shape",
                        "id": "closing-mark",
                        "placement": "bottom_left",
                        "shape": "rounded_rectangle",
                        "fill": "accent_alt",
                    },
                    {
                        "kind": "text",
                        "id": "closing-title",
                        "role": "title",
                        "text": "Decision requested",
                    },
                    {
                        "kind": "text",
                        "id": "closing-subtitle",
                        "role": "subtitle",
                        "text": "Approve the next-quarter delivery plan",
                    },
                ],
            },
        ],
    }


def _compiled(
    *,
    aspect_ratio: str = "16:9",
) -> tuple[PresentationIntent, bytes, dict, dict]:
    intent = PresentationIntent.model_validate(_intent_payload(aspect_ratio=aspect_ratio))
    artifact, receipt = compile_presentation(
        intent,
        image_assets={_IMAGE_PATH: _png()},
    )
    preflight = office_engine.preflight(artifact, suffix=".pptx")
    return intent, artifact, receipt, preflight


def _stored_intent(intent: PresentationIntent) -> dict:
    return intent.model_dump(by_alias=True, exclude_none=True, mode="json")


def test_intent_schema_is_typed_and_rejects_unowned_geometry() -> None:
    intent = PresentationIntent.model_validate(_intent_payload())

    assert intent.schema_version == PRESENTATION_INTENT_SCHEMA
    assert intent.image_paths() == (_IMAGE_PATH,)

    raw_geometry = _intent_payload()
    raw_geometry["slides"][0]["elements"][1]["left"] = 120
    with pytest.raises(ValidationError, match="Extra inputs are not permitted"):
        PresentationIntent.model_validate(raw_geometry)

    duplicate_id = _intent_payload()
    duplicate_id["slides"][1]["elements"][0]["id"] = "opening-title"
    with pytest.raises(ValidationError, match="unique across"):
        PresentationIntent.model_validate(duplicate_id)

    missing_accessible_title = _intent_payload()
    missing_accessible_title["slides"][0]["elements"].pop(1)
    with pytest.raises(ValidationError, match="requires"):
        PresentationIntent.model_validate(missing_accessible_title)


def test_compile_is_deterministic_native_editable_and_fully_preflighted() -> None:
    intent, first, first_receipt, preflight = _compiled()
    second, second_receipt = compile_presentation(
        intent,
        image_assets={_IMAGE_PATH: _png()},
    )

    assert first == second
    assert first_receipt == second_receipt
    assert first_receipt["schema"] == PRESENTATION_GENERATION_RECEIPT_SCHEMA
    assert first_receipt["output"] == {
        "format": "pptx",
        "sha256": hashlib.sha256(first).hexdigest(),
        "size_bytes": len(first),
    }
    assert first_receipt["presentation"]["slide_count"] == 5
    assert first_receipt["presentation"]["object_count"] == 15
    assert first_receipt["bounds"] == {
        "slides_returned": 5,
        "slides_truncated": False,
        "objects_returned": 15,
        "objects_truncated": False,
        "assets_returned": 2,
        "assets_truncated": False,
    }
    assert preflight["summary"]["slide_count"] == 5
    assert preflight["summary"]["findings_truncated"] is False
    assert preflight["summary"]["findings_by_severity"].get("error", 0) == 0
    assert office_engine.validate(first, suffix=".pptx")["valid"] is True

    validate_generation_evidence(
        artifact=first,
        intent=_stored_intent(intent),
        receipt=first_receipt,
        preflight=preflight,
    )

    presentation = Presentation(io.BytesIO(first))
    assert len(presentation.slides) == 5
    assert presentation.slides[0].shapes.title is not None
    assert presentation.slides[0].shapes.title.is_placeholder
    authored_names = {shape.name for slide in presentation.slides for shape in slide.shapes}
    assert "VFGEN:opening:opening-title" in authored_names
    assert "VFGEN:evidence:evidence-image" in authored_names
    assert "VFGEN:closing:closing-mark" in authored_names

    evidence_picture = next(shape for shape in presentation.slides[2].shapes if shape.name == "VFGEN:evidence:evidence-image")
    detail_picture = next(shape for shape in presentation.slides[3].shapes if shape.name == "VFGEN:detail:detail-image")
    assert evidence_picture.crop_left > 0
    assert evidence_picture.crop_right > 0
    assert detail_picture.crop_left == detail_picture.crop_right == 0
    assert detail_picture._element.nvPicPr.cNvPr.get("descr") == ("Blue detail panel supporting the change explanation")


def test_generation_supports_four_by_three_and_rejects_unsafe_assets_or_overflow() -> None:
    intent, artifact, receipt, _preflight = _compiled(aspect_ratio="4:3")
    presentation = Presentation(io.BytesIO(artifact))

    assert presentation.slide_width == 9_144_000
    assert presentation.slide_height == 6_858_000
    assert receipt["presentation"]["aspect_ratio"] == "4:3"

    with pytest.raises(OfficeOperationError, match="missing=.*hero.png"):
        compile_presentation(intent)
    with pytest.raises(OfficeOperationError, match="unexpected"):
        compile_presentation(
            intent,
            image_assets={
                _IMAGE_PATH: _png(),
                "/mnt/user-data/uploads/extra.png": _png(),
            },
        )
    with pytest.raises(OfficeOperationError, match="baseline JPEG"):
        compile_presentation(
            intent,
            image_assets={_IMAGE_PATH: _jpeg(progressive=True)},
        )

    overflow = _intent_payload()
    overflow["slides"][0]["elements"][1]["text"] = "W" * 4_000
    with pytest.raises(OfficeOperationError, match="does not fit"):
        compile_presentation(
            PresentationIntent.model_validate(overflow),
            image_assets={_IMAGE_PATH: _png()},
        )


@pytest.mark.parametrize(
    "mutation",
    [
        lambda receipt: receipt["slides"][0]["objects"][0].__setitem__("object_path", "/slide[1]/shape[@id=999]"),
        lambda receipt: receipt["presentation"].__setitem__("mapping_sha256", "0" * 64),
        lambda receipt: receipt["assets"][0].__setitem__("sha256", "0" * 64),
        lambda receipt: receipt.__setitem__("unexpected", True),
    ],
)
def test_generation_evidence_rejects_tampered_receipt(mutation) -> None:
    intent, artifact, receipt, preflight = _compiled()
    tampered = copy.deepcopy(receipt)
    mutation(tampered)

    with pytest.raises(OfficePackageError):
        validate_generation_evidence(
            artifact=artifact,
            intent=_stored_intent(intent),
            receipt=tampered,
            preflight=preflight,
        )


def test_generation_evidence_rejects_tampered_preflight_and_intent() -> None:
    intent, artifact, receipt, preflight = _compiled()
    tampered_preflight = copy.deepcopy(preflight)
    tampered_preflight["summary"]["findings_by_severity"] = {"error": 1}
    with pytest.raises(OfficePackageError, match="preflight"):
        validate_generation_evidence(
            artifact=artifact,
            intent=_stored_intent(intent),
            receipt=receipt,
            preflight=tampered_preflight,
        )

    malformed_counts = copy.deepcopy(preflight)
    malformed_counts["summary"]["findings_by_severity"] = {"warning": "1"}
    with pytest.raises(OfficePackageError, match="preflight"):
        validate_generation_evidence(
            artifact=artifact,
            intent=_stored_intent(intent),
            receipt=receipt,
            preflight=malformed_counts,
        )

    tampered_intent = _stored_intent(intent)
    tampered_intent["title"] = "Changed title"
    with pytest.raises(OfficePackageError, match="identity"):
        validate_generation_evidence(
            artifact=artifact,
            intent=tampered_intent,
            receipt=receipt,
            preflight=preflight,
        )


def test_generated_revision_roundtrip_can_append_a_native_edit(tmp_path: Path) -> None:
    intent, artifact, receipt, preflight = _compiled()
    store = OfficeRevisionStore(tmp_path / "office")
    generated = store.commit_generation(
        artifact=artifact,
        output_path="/mnt/user-data/outputs/generated.pptx",
        thread_id="thread-generation",
        intent=_stored_intent(intent),
        receipt=receipt,
        preflight=preflight,
        validation=office_engine.validate(artifact, suffix=".pptx"),
    )

    assert generated.parent_revision_id is None
    assert generated.sequence == 1
    stored = store.load_revision(generated.project_id, generated.revision_id)
    assert stored["kind"] == "generation"
    assert stored["generation"]["receipt"] == receipt
    _revision, stored_artifact = store.read_revision_artifact(
        generated.project_id,
        generated.revision_id,
    )
    assert stored_artifact == artifact

    title_path = next(item["object_path"] for item in receipt["slides"][0]["objects"] if item["element_id"] == "opening-title")
    edited, _reports, edit_receipt = office_engine.edit_with_receipt(
        artifact,
        suffix=".pptx",
        operations=[
            PptxTextReplacement(
                paths=[title_path],
                find="Quarterly Update",
                replace="Quarterly Review",
            )
        ],
    )
    appended = store.commit_edit(
        source=artifact,
        result=edited,
        suffix=".pptx",
        source_path="/mnt/user-data/outputs/generated.pptx",
        output_path="/mnt/user-data/outputs/generated-v2.pptx",
        thread_id="thread-generation",
        receipt=edit_receipt,
        source_validation=office_engine.validate(artifact, suffix=".pptx"),
        result_validation=office_engine.validate(edited, suffix=".pptx"),
        project_id=generated.project_id,
        parent_revision_id=generated.revision_id,
    )

    assert appended.parent_revision_id == generated.revision_id
    assert appended.sequence == 2
    assert [item["kind"] for item in store.list_revisions(generated.project_id)] == [
        "edit",
        "generation",
    ]

    edit_metadata_path = tmp_path / "office" / "projects" / generated.project_id / "revisions" / appended.revision_id / "revision.json"
    edit_metadata = json.loads(edit_metadata_path.read_text(encoding="utf-8"))
    edit_metadata["generation"] = stored["generation"]
    edit_metadata_path.write_text(json.dumps(edit_metadata), encoding="utf-8")
    with pytest.raises(OfficeRevisionIntegrityError, match="change revision evidence"):
        store.load_revision(generated.project_id, appended.revision_id)

    metadata_path = tmp_path / "office" / "projects" / generated.project_id / "revisions" / generated.revision_id / "revision.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    metadata["generation"]["receipt"]["slides"][0]["slide_id"] = "tampered"
    metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
    with pytest.raises(
        OfficeRevisionIntegrityError,
        match="generation evidence",
    ):
        store.load_revision(generated.project_id, generated.revision_id)


class _Sandbox:
    id = "office-generation-test"

    def __init__(self, files: dict[str, bytes]) -> None:
        self.files = dict(files)

    def download_file(self, path: str) -> bytes:
        if path not in self.files:
            raise FileNotFoundError(errno.ENOENT, "not found", path)
        return self.files[path]

    def update_file(self, path: str, content: bytes) -> None:
        self.files[path] = content

    def replace_file(self, source_path: str, destination_path: str) -> None:
        self.files[destination_path] = self.files.pop(source_path)


def _runtime() -> ToolRuntime:
    return ToolRuntime(
        state={"sandbox": {"sandbox_id": "office-generation-test"}},
        context={"thread_id": "office-generation-thread"},
        config={"configurable": {"thread_id": "office-generation-thread"}},
        stream_writer=lambda _: None,
        tools=[],
        tool_call_id="office-generate-call",
        store=None,
    )


def test_office_generate_tool_persists_exact_trusted_artifact(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    sandbox = _Sandbox({_IMAGE_PATH: _png()})
    store = OfficeRevisionStore(tmp_path / "office")
    monkeypatch.setattr(
        "vassilflow.community.office.tools.ensure_sandbox_initialized",
        lambda _runtime: sandbox,
    )
    monkeypatch.setattr(
        "vassilflow.community.office.tools.ensure_thread_directories_exist",
        lambda _runtime: None,
    )
    monkeypatch.setattr(
        "vassilflow.community.office.tools._office_revision_store",
        lambda _runtime: store,
    )
    monkeypatch.setattr(
        "vassilflow.community.office.tools._mirror_binary_to_gateway_if_needed",
        lambda _runtime, _path, _content: True,
    )
    intent = PresentationIntent.model_validate(_intent_payload())
    output_path = "/mnt/user-data/outputs/generated.pptx"

    response = json.loads(
        office_generate_tool.func(
            _runtime(),
            intent,
            output_path,
        )
    )

    assert response["ok"] is True
    assert response["commit_status"] == "committed"
    assert response["output_path"] == output_path
    assert response["generation_receipt"]["schema"] == (PRESENTATION_GENERATION_RECEIPT_SCHEMA)
    assert "source_path" not in response
    project_id = response["project"]["project_id"]
    revision_id = response["revision"]["revision_id"]
    _revision, stored = store.read_revision_artifact(
        project_id,
        revision_id,
    )
    assert sandbox.files[output_path] == stored
    assert hashlib.sha256(stored).hexdigest() == response["revision"]["artifact"]["sha256"]

    duplicate = json.loads(
        office_generate_tool.func(
            _runtime(),
            intent,
            output_path,
        )
    )
    assert duplicate["ok"] is False
    assert "already exists" in duplicate["error"]
    assert len(store.list_projects()) == 1


def test_office_generate_tool_schema_exposes_only_semantic_intent() -> None:
    assert office_generate_tool.name == "office_generate"
    schema = office_generate_tool.tool_call_schema.model_json_schema()
    intent_schema = schema["$defs"]["PresentationIntent"]

    assert intent_schema["additionalProperties"] is False
    assert set(intent_schema["properties"]) == {
        "schema",
        "title",
        "aspect_ratio",
        "theme",
        "slides",
    }
    serialized = json.dumps(schema)
    assert "left_emu" not in serialized
    assert "raw_ooxml" not in serialized

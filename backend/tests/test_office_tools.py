from __future__ import annotations

import asyncio
import base64
import errno
import hashlib
import io
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from langchain.tools import ToolRuntime
from openpyxl import Workbook, load_workbook
from PIL import Image
from pydantic import ValidationError

from vassilflow.community.office.engine import office_engine
from vassilflow.community.office.models import (
    DocxRunFormatOperation,
    DocxTextReplacement,
    PptxLineFormatOperation,
    PptxLineSelector,
    PptxObjectSelector,
    PptxParagraphFormatOperation,
    PptxParagraphSelector,
    PptxPictureSourceReplacementOperation,
    PptxRunFormatOperation,
    PptxRunSelector,
    PptxShapeFormatOperation,
    PptxShapeSelector,
    PptxTextReplacement,
    XlsxCellFormatOperation,
)
from vassilflow.community.office.render import OfficeRenderResult, RenderedOfficePage
from vassilflow.community.office.revisions import OfficeRevisionCommit, OfficeRevisionStore
from vassilflow.community.office.selection import (
    build_pptx_selection_surface,
    resolve_pptx_selection,
)
from vassilflow.community.office.tools import (
    _mirror_binary_to_gateway_if_needed,
    _office_inspect_async,
    _visual_review_contract,
    office_edit_tool,
    office_inspect_tool,
    office_render_tool,
)

_PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=")


def _jpeg() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (2, 2), (68, 114, 196)).save(
        output,
        format="JPEG",
        quality=90,
        progressive=False,
        optimize=False,
    )
    return output.getvalue()


_JPEG = _jpeg()

_FAKE_PROJECT_ID = f"ofp_{'1' * 32}"
_FAKE_BASELINE_REVISION_ID = f"ofr_{'2' * 32}"
_FAKE_REVISION_ID = f"ofr_{'3' * 32}"
_PPTX_CORPUS = Path(__file__).parent / "fixtures" / "office" / "pptx" / "roundtrip" / "seed-v1.pptx"


def _docx(text: str) -> bytes:
    document = (f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p><w:sectPr/></w:body></w:document>').encode()
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>',
        )
        archive.writestr("word/document.xml", document)
    return output.getvalue()


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Revenue"
    sheet["B1"] = 42
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pptx(text: str, *, picture_source: bytes | None = None) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        image_content_types = '<Default Extension="png" ContentType="image/png"/><Default Extension="jpg" ContentType="image/jpeg"/>' if picture_source is not None else ""
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            f"{image_content_types}"
            '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
            '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "ppt/presentation.xml",
            '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst>'
            '<p:sldSz cx="12192000" cy="6858000"/></p:presentation>',
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "ppt/slides/slide1.xml",
            '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            "<p:cSld><p:spTree>"
            '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
            "<p:grpSpPr/><p:sp>"
            '<p:nvSpPr><p:cNvPr id="2" name="Text"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f"<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody>"
            "</p:sp>"
            + (
                '<p:pic><p:nvPicPr><p:cNvPr id="3" name="Picture" descr="Preserved alt text"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
                '<p:blipFill dpi="144" rotWithShape="0"><a:blip r:embed="rIdImage"><a:grayscl/></a:blip>'
                '<a:srcRect l="5000" t="10000" r="0" b="0"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
                '<p:spPr><a:xfrm><a:off x="10" y="20"/><a:ext cx="30" cy="40"/></a:xfrm></p:spPr></p:pic>'
                if picture_source is not None
                else ""
            )
            + "</p:spTree></p:cSld></p:sld>",
        )
        if picture_source is not None:
            archive.writestr(
                "ppt/slides/_rels/slide1.xml.rels",
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rIdImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/>'
                "</Relationships>",
            )
            archive.writestr("ppt/media/image1.png", picture_source)
    return output.getvalue()


class _Sandbox:
    id = "office-test"

    def __init__(self, files: dict[str, bytes]) -> None:
        self.files = dict(files)

    def download_file(self, path: str) -> bytes:
        if path not in self.files:
            raise FileNotFoundError(errno.ENOENT, "not found", path)
        return self.files[path]

    def update_file(self, path: str, content: bytes) -> None:
        self.files[path] = content

    def replace_file(self, source_path: str, destination_path: str) -> None:
        self.files[destination_path] = self.files.pop(source_path)

    def list_dir(self, path: str, max_depth: int = 2) -> list[str]:
        del max_depth
        prefix = f"{path.rstrip('/')}/"
        matching = sorted(file_path for file_path in self.files if file_path.startswith(prefix))
        return [path, *matching] if matching else []


class _FakeRevisionStore:
    def commit_edit(self, *, result: bytes, **_kwargs) -> OfficeRevisionCommit:
        return OfficeRevisionCommit(
            project_id=_FAKE_PROJECT_ID,
            revision_id=_FAKE_REVISION_ID,
            parent_revision_id=_FAKE_BASELINE_REVISION_ID,
            sequence=2,
            project_created=True,
            revision_count=2,
            artifact_sha256=hashlib.sha256(result).hexdigest(),
            artifact_size_bytes=len(result),
            baseline_revision_id=_FAKE_BASELINE_REVISION_ID,
        )


def _patch_sandbox(monkeypatch, sandbox: _Sandbox) -> None:
    monkeypatch.setattr(
        "vassilflow.community.office.tools.ensure_sandbox_initialized",
        lambda _runtime: sandbox,
    )
    monkeypatch.setattr(
        "vassilflow.community.office.tools.ensure_thread_directories_exist",
        lambda _runtime: None,
    )
    monkeypatch.setattr(
        "vassilflow.community.office.tools._office_revision_store",
        lambda _runtime: _FakeRevisionStore(),
    )


def _runtime() -> ToolRuntime:
    return ToolRuntime(
        state={"sandbox": {"sandbox_id": "office-test"}, "thread_data": {}},
        context={"thread_id": "office-thread"},
        config={"configurable": {"thread_id": "office-thread"}},
        stream_writer=lambda _: None,
        tools=[],
        tool_call_id="office-call",
        store=None,
    )


def _seed_selected_pptx(
    store: OfficeRevisionStore,
) -> tuple[ToolRuntime, str, str, bytes]:
    source = _PPTX_CORPUS.read_bytes()
    result, _reports, receipt = office_engine.edit_with_receipt(
        source,
        suffix=".pptx",
        operations=[
            PptxTextReplacement(
                paths=["/slide[1]/shape[@id=5]"],
                find="Final",
                replace="Reviewed",
            )
        ],
    )
    commit = store.commit_edit(
        source=source,
        result=result,
        suffix=".pptx",
        source_path="/mnt/user-data/uploads/source.pptx",
        output_path="/mnt/user-data/outputs/result.pptx",
        thread_id="office-thread",
        receipt=receipt,
        source_validation=office_engine.validate(source, suffix=".pptx"),
        result_validation=office_engine.validate(result, suffix=".pptx"),
    )
    surface = build_pptx_selection_surface(result, slide_index=1)
    selected = next(item for item in surface["objects"] if item["path"] == "/slide[1]/shape[@id=3]")
    resolved = resolve_pptx_selection(
        result,
        slide_index=1,
        source_sha256=surface["source_sha256"],
        object_path=selected["path"],
        object_fingerprint=selected["object_fingerprint"],
    )
    runtime = _runtime()
    runtime.context["office_selection"] = {
        **resolved,
        "project_id": commit.project_id,
        "revision_id": commit.revision_id,
        "artifact_size_bytes": len(result),
    }
    return runtime, commit.project_id, commit.revision_id, result


def test_office_tool_names_and_nested_operation_schema() -> None:
    assert office_inspect_tool.name == "office_inspect"
    assert office_edit_tool.name == "office_edit"
    assert office_render_tool.name == "office_render"

    schema = office_edit_tool.tool_call_schema.model_json_schema()
    assert "source_path" in schema["required"]
    assert {item.get("type") for item in schema["properties"]["source_path"]["anyOf"]} == {
        "string",
        "null",
    }
    assert schema["properties"]["project_id"]["default"] is None
    assert schema["properties"]["parent_revision_id"]["default"] is None
    for selector_name in (
        "PptxRunSelector",
        "PptxParagraphSelector",
        "PptxShapeSelector",
        "PptxLineSelector",
        "PptxSlideSelector",
        "PptxPictureSelector",
    ):
        selector_schema = schema["$defs"][selector_name]
        assert selector_schema["additionalProperties"] is False
        assert set(selector_schema["properties"]) == {"targets"}
        assert "accepts only ``targets``" in selector_schema["description"]
    for operation_name, field_name in (
        ("PptxRunFormatOperation", "runs"),
        ("PptxParagraphFormatOperation", "paragraphs"),
        ("PptxShapeFormatOperation", "shapes"),
        ("PptxLineFormatOperation", "lines"),
        ("PptxSlideBackgroundFormatOperation", "slides"),
        ("PptxPictureSourceReplacementOperation", "pictures"),
    ):
        operation_schema = schema["$defs"][operation_name]
        field_schema = operation_schema["properties"][field_name]
        assert operation_schema["additionalProperties"] is False
        assert "with only ``type``" in operation_schema["description"]
        assert "occurrence" not in operation_schema["properties"]
        assert "require_match" not in operation_schema["properties"]
        assert "containing exactly targets" in field_schema["description"]
        assert "occurrence" in field_schema["description"]

    shape_target_schema = schema["$defs"]["PptxShapeTarget"]
    assert "expected_name" in shape_target_schema["required"]
    assert {item.get("type") for item in shape_target_schema["properties"]["expected_name"]["anyOf"]} == {
        "string",
        "null",
    }
    line_schema = schema["$defs"]["PptxShapeLineFormatting"]
    assert {
        "fill",
        "width_points",
        "cap",
        "dash",
        "join",
        "miter_limit_percent",
    } == set(line_schema["properties"])
    direct_line_schema = schema["$defs"]["PptxLineFormatting"]
    assert {
        "fill",
        "width_points",
        "cap",
        "dash",
        "join",
        "miter_limit_percent",
        "compound",
        "alignment",
        "head_end",
        "tail_end",
    } == set(direct_line_schema["properties"])
    line_target_schema = schema["$defs"]["PptxLineTarget"]
    assert "expected_name" in line_target_schema["required"]
    picture_target_schema = schema["$defs"]["PptxPictureTarget"]
    assert picture_target_schema["required"] == [
        "path",
        "expected_name",
        "expected_source_sha256",
    ]
    assert picture_target_schema["properties"]["expected_source_sha256"]["pattern"] == "^[0-9a-f]{64}$"
    image_fill_schema = schema["$defs"]["PptxShapeFillFormatting"]
    assert {"mode", "crop", "tile"} <= set(image_fill_schema["properties"])
    assert set(schema["$defs"]["PptxImageTileFormatting"]["required"]) == {
        "offset_x_emu",
        "offset_y_emu",
        "scale_x_percent",
        "scale_y_percent",
        "alignment",
        "flip",
    }

    assert "selector contains only targets" in office_edit_tool.description
    assert "explicit tile scale/offset" in office_edit_tool.description
    assert "source SHA-256 guards" in office_edit_tool.description
    assert "preserves crop, framing, effects, geometry, alt text, and interactions" in office_edit_tool.description
    assert "occurrence and require_match are not accepted" in office_edit_tool.description
    assert "semantic change receipt built before commit" in office_edit_tool.description
    assert "package part/relationship changes" in office_edit_tool.description
    inspect_schema = office_inspect_tool.tool_call_schema.model_json_schema()
    analysis_mode_schema = inspect_schema["properties"]["analysis_mode"]
    assert analysis_mode_schema["enum"] == [
        "inspect",
        "pptx_quality_preflight",
    ]
    assert "does not mutate, render, or claim visual QA" in (analysis_mode_schema["description"])
    media_gc_description = inspect_schema["properties"]["include_pptx_media_gc_plan"]["description"]
    assert "non-destructive dry-run" in media_gc_description
    assert "never deletes data" in media_gc_description
    render_schema = office_render_tool.tool_call_schema.model_json_schema()
    assert render_schema["properties"]["project_id"]["default"] is None
    assert render_schema["properties"]["revision_id"]["default"] is None
    assert "persist its manifest and preview PNGs" in office_render_tool.description

    parsed = office_edit_tool.tool_call_schema.model_validate(
        {
            "source_path": "/mnt/user-data/uploads/in.docx",
            "output_path": "/mnt/user-data/workspace/out.docx",
            "operations": [{"type": "replace_text", "find": "old", "replace": "new"}],
        }
    )
    assert isinstance(parsed.operations[0], DocxTextReplacement)

    formatted = office_edit_tool.tool_call_schema.model_validate(
        {
            "source_path": "/mnt/user-data/uploads/in.docx",
            "output_path": "/mnt/user-data/workspace/formatted.docx",
            "operations": [
                {
                    "type": "format_runs",
                    "paragraphs": {"paragraph_indices": [1]},
                    "runs": {"contains_text": "old"},
                    "formatting": {"bold": True, "color": "#FF0000"},
                }
            ],
        }
    )
    assert isinstance(formatted.operations[0], DocxRunFormatOperation)

    xlsx_formatted = office_edit_tool.tool_call_schema.model_validate(
        {
            "source_path": "/mnt/user-data/uploads/in.xlsx",
            "output_path": "/mnt/user-data/workspace/formatted.xlsx",
            "operations": [
                {
                    "type": "format_cells",
                    "cells": {"sheet_name": "Data", "ranges": ["B1"]},
                    "formatting": {"bold": True, "number_format": "#,##0"},
                }
            ],
        }
    )
    assert isinstance(xlsx_formatted.operations[0], XlsxCellFormatOperation)

    pptx_replaced = office_edit_tool.tool_call_schema.model_validate(
        {
            "source_path": "/mnt/user-data/uploads/in.pptx",
            "output_path": "/mnt/user-data/workspace/replaced.pptx",
            "operations": [
                {
                    "type": "replace_pptx_text",
                    "paths": ["/slide[1]/shape[@id=2]"],
                    "find": "old",
                    "replace": "new",
                }
            ],
        }
    )
    assert isinstance(pptx_replaced.operations[0], PptxTextReplacement)

    pptx_formatted = office_edit_tool.tool_call_schema.model_validate(
        {
            "source_path": "/mnt/user-data/uploads/in.pptx",
            "output_path": "/mnt/user-data/workspace/formatted.pptx",
            "operations": [
                {
                    "type": "format_pptx_runs",
                    "runs": {
                        "targets": [
                            {
                                "path": "/slide[1]/shape[@id=2]/paragraph[1]/run[1]",
                                "expected_text": "target",
                            }
                        ]
                    },
                    "formatting": {"bold": True, "color": "#FF0000"},
                },
                {
                    "type": "format_pptx_paragraphs",
                    "paragraphs": {
                        "targets": [
                            {
                                "path": "/slide[1]/shape[@id=2]/paragraph[1]",
                                "expected_text": "target",
                            }
                        ]
                    },
                    "formatting": {"alignment": "center"},
                },
                {
                    "type": "format_pptx_shapes",
                    "shapes": {
                        "targets": [
                            {
                                "path": "/slide[1]/shape[@id=2]",
                                "expected_name": "Text",
                            }
                        ]
                    },
                    "formatting": {
                        "fill": {"type": "solid", "color": "#336699"},
                        "line": {
                            "cap": "square",
                            "dash": "large_dash_dot_dot",
                            "join": "miter",
                            "miter_limit_percent": 800,
                        },
                        "text_box": {"vertical_anchor": "middle"},
                    },
                },
                {
                    "type": "format_pptx_lines",
                    "lines": {
                        "targets": [
                            {
                                "path": "/slide[1]/connector[@id=9]",
                                "expected_name": "Connector",
                            }
                        ]
                    },
                    "formatting": {
                        "compound": "triple",
                        "alignment": "inset",
                        "head_end": {
                            "type": "stealth",
                            "width": "large",
                            "length": "small",
                        },
                    },
                },
            ],
        }
    )
    assert isinstance(pptx_formatted.operations[0], PptxRunFormatOperation)
    assert isinstance(
        pptx_formatted.operations[1],
        PptxParagraphFormatOperation,
    )
    assert isinstance(pptx_formatted.operations[2], PptxShapeFormatOperation)
    assert isinstance(pptx_formatted.operations[3], PptxLineFormatOperation)

    pptx_picture_replaced = office_edit_tool.tool_call_schema.model_validate(
        {
            "source_path": "/mnt/user-data/uploads/in.pptx",
            "output_path": "/mnt/user-data/workspace/picture-replaced.pptx",
            "operations": [
                {
                    "type": "replace_pptx_picture_sources",
                    "pictures": {
                        "targets": [
                            {
                                "path": "/slide[1]/picture[@id=3]",
                                "expected_name": "Picture",
                                "expected_source_sha256": "0" * 64,
                            }
                        ]
                    },
                    "image_path": "/mnt/user-data/uploads/replacement.jpg",
                }
            ],
        }
    )
    assert isinstance(
        pptx_picture_replaced.operations[0],
        PptxPictureSourceReplacementOperation,
    )

    rendered = office_render_tool.tool_call_schema.model_validate(
        {
            "path": "/mnt/user-data/uploads/in.docx",
            "output_dir": "/mnt/user-data/workspace/visual-qa",
        }
    )
    assert rendered.max_pages == 1
    assert rendered.dpi == 120

    pptx_inspected = office_inspect_tool.tool_call_schema.model_validate(
        {
            "path": "/mnt/user-data/uploads/in.pptx",
            "start_slide": 2,
            "max_slides": 3,
            "include_pptx_formatting": True,
            "include_pptx_annotations": True,
            "include_pptx_dynamics": True,
            "include_pptx_media_gc_plan": True,
            "pptx_selector": {
                "paths": ["/slide[2]/shape[@id=7]"],
                "kinds": ["text_box"],
            },
        }
    )
    assert pptx_inspected.start_slide == 2
    assert pptx_inspected.max_slides == 3
    assert pptx_inspected.include_pptx_formatting is True
    assert pptx_inspected.include_pptx_annotations is True
    assert pptx_inspected.include_pptx_dynamics is True
    assert pptx_inspected.include_pptx_media_gc_plan is True
    assert isinstance(pptx_inspected.pptx_selector, PptxObjectSelector)


def test_office_edit_normalizes_only_inert_pptx_selector_union_fields() -> None:
    run_target = {
        "path": "/slide[1]/shape[@id=2]/paragraph[1]/run[1]",
        "expected_text": "Alpha",
    }
    paragraph_target = {
        "path": "/slide[1]/shape[@id=2]/paragraph[1]",
        "expected_text": "Alpha",
    }
    shape_target = {
        "path": "/slide[1]/shape[@id=2]",
        "expected_name": "Text",
    }
    line_target = {
        "path": "/slide[1]/connector[@id=9]",
        "expected_name": "Connector",
    }
    parsed = office_edit_tool.tool_call_schema.model_validate(
        {
            "source_path": "/mnt/user-data/uploads/in.pptx",
            "output_path": "/mnt/user-data/workspace/formatted.pptx",
            "operations": [
                {
                    "type": "format_pptx_runs",
                    "runs": {
                        "targets": [run_target],
                        "contains_text": None,
                        "occurrence": "all",
                        "require_match": True,
                    },
                    "formatting": {"italic": True},
                },
                {
                    "type": "format_pptx_paragraphs",
                    "paragraphs": {
                        "targets": [paragraph_target],
                        "contains_text": None,
                        "occurrence": None,
                        "require_match": None,
                    },
                    "formatting": {"alignment": "justify"},
                },
                {
                    "type": "format_pptx_shapes",
                    "shapes": {
                        "targets": [shape_target],
                        "contains_text": None,
                        "occurrence": "all",
                        "require_match": True,
                    },
                    "formatting": {
                        "fill": {"type": "none"},
                    },
                },
                {
                    "type": "format_pptx_lines",
                    "lines": {
                        "targets": [line_target],
                        "contains_text": None,
                        "occurrence": "all",
                        "require_match": True,
                    },
                    "formatting": {"compound": "double"},
                },
            ],
        }
    )

    assert parsed.operations[0].runs.model_dump() == {"targets": [run_target]}
    assert parsed.operations[1].paragraphs.model_dump() == {"targets": [paragraph_target]}
    assert parsed.operations[2].shapes.model_dump() == {"targets": [shape_target]}
    assert parsed.operations[3].lines.model_dump() == {"targets": [line_target]}

    invalid_selector_fields = (
        {"contains_text": "Alpha"},
        {"occurrence": "first"},
        {"require_match": False},
        {"unexpected": None},
    )
    for extra in invalid_selector_fields:
        with pytest.raises(ValidationError):
            office_edit_tool.tool_call_schema.model_validate(
                {
                    "source_path": "/mnt/user-data/uploads/in.pptx",
                    "output_path": "/mnt/user-data/workspace/formatted.pptx",
                    "operations": [
                        {
                            "type": "format_pptx_runs",
                            "runs": {"targets": [run_target], **extra},
                            "formatting": {"italic": True},
                        }
                    ],
                }
            )

    with pytest.raises(ValidationError):
        PptxRunSelector.model_validate({"targets": [run_target], "contains_text": None})
    with pytest.raises(ValidationError):
        PptxParagraphSelector.model_validate({"targets": [paragraph_target], "require_match": True})
    with pytest.raises(ValidationError):
        PptxShapeSelector.model_validate({"targets": [shape_target], "contains_text": None})
    with pytest.raises(ValidationError):
        PptxLineSelector.model_validate({"targets": [line_target], "contains_text": None})
    with pytest.raises(ValidationError):
        office_edit_tool.tool_call_schema.model_validate(
            {
                "source_path": "/mnt/user-data/uploads/in.pptx",
                "output_path": "/mnt/user-data/workspace/formatted.pptx",
                "operations": [
                    {
                        "type": "format_pptx_runs",
                        "runs": {"targets": [run_target]},
                        "formatting": {"italic": True},
                        "occurrence": "all",
                    }
                ],
            }
        )


def test_office_inspect_reads_binary_through_sandbox(monkeypatch) -> None:
    path = "/mnt/user-data/uploads/in.docx"
    sandbox = _Sandbox({path: _docx("hello")})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(office_inspect_tool.func(SimpleNamespace(), path))

    assert result["ok"] is True
    assert result["paragraphs"][0]["text"] == "hello"


def test_office_inspect_lists_and_reads_bounded_xlsx_range(monkeypatch) -> None:
    path = "/mnt/user-data/uploads/in.xlsx"
    sandbox = _Sandbox({path: _xlsx()})
    _patch_sandbox(monkeypatch, sandbox)

    summary = json.loads(office_inspect_tool.func(SimpleNamespace(), path))
    inspected = json.loads(
        office_inspect_tool.func(
            SimpleNamespace(),
            path,
            sheet_name="Data",
            cell_range="A1:B1",
            include_cell_styles=True,
        )
    )

    assert summary["ok"] is True
    assert summary["sheets"][0]["name"] == "Data"
    assert inspected["ok"] is True
    assert [cell["coordinate"] for cell in inspected["cells"]] == ["A1", "B1"]


def test_office_inspect_passes_bounded_pptx_slide_window(monkeypatch) -> None:
    path = "/mnt/user-data/uploads/in.pptx"
    sandbox = _Sandbox({path: b"pptx"})
    _patch_sandbox(monkeypatch, sandbox)
    captured: dict[str, object] = {}

    def fake_inspect(document: bytes, **kwargs):
        captured.update(document=document, **kwargs)
        return {"format": "pptx", "slide_count": 4, "slides": []}

    monkeypatch.setattr("vassilflow.community.office.tools.office_engine.inspect", fake_inspect)
    selector = PptxObjectSelector(paths=["/slide[2]/shape[@id=7]"])

    result = json.loads(
        office_inspect_tool.func(
            SimpleNamespace(),
            path,
            start_slide=2,
            max_slides=2,
            include_pptx_formatting=True,
            include_pptx_annotations=True,
            include_pptx_dynamics=True,
            include_pptx_media_gc_plan=True,
            pptx_selector=selector,
        )
    )

    assert result["ok"] is True
    assert captured["document"] == b"pptx"
    assert captured["suffix"] == ".pptx"
    assert captured["start_slide"] == 2
    assert captured["max_slides"] == 2
    assert captured["include_pptx_formatting"] is True
    assert captured["include_pptx_annotations"] is True
    assert captured["include_pptx_dynamics"] is True
    assert captured["include_pptx_media_gc_plan"] is True
    assert captured["pptx_selector"] is selector


def test_office_inspect_returns_opt_in_pptx_media_gc_plan(monkeypatch) -> None:
    path = "/mnt/user-data/uploads/in.pptx"
    sandbox = _Sandbox({path: _pptx("hello", picture_source=_PNG)})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_inspect_tool.func(
            SimpleNamespace(),
            path,
            include_pptx_media_gc_plan=True,
        )
    )

    assert result["ok"] is True
    assert result["media_gc_plan_included"] is True
    assert result["media_gc_plan"]["contract_version"] == 2
    assert result["media_gc_plan"]["mode"] == "dry_run"
    assert result["media_gc_plan"]["destructive"] is False
    assert result["media_gc_plan"]["analysis_complete"] is True
    assert result["media_gc_plan"]["package_xml_integrity_complete"] is True
    assert result["media_gc_plan"]["root_reachable_scoped_image_part_count"] == 1
    assert result["media_gc_plan"]["live_relationship_count"] == 1
    assert result["media_gc_plan"]["relationship_candidate_count"] == 0
    assert result["media_gc_plan"]["part_candidate_count"] == 0


def test_office_inspect_runs_explicit_pptx_quality_preflight(monkeypatch) -> None:
    path = "/mnt/user-data/uploads/in.pptx"
    document = _pptx("hello")
    sandbox = _Sandbox({path: document})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_inspect_tool.func(
            SimpleNamespace(),
            path,
            analysis_mode="pptx_quality_preflight",
        )
    )

    assert result["ok"] is True
    assert result["path"] == path
    assert result["analysis"] == "quality_preflight"
    assert result["source_sha256"] == hashlib.sha256(document).hexdigest()
    assert result["visual_review_status"] == "not_performed"


def test_office_inspect_rejects_invalid_quality_preflight_combinations(
    monkeypatch,
) -> None:
    pptx_path = "/mnt/user-data/uploads/in.pptx"
    docx_path = "/mnt/user-data/uploads/in.docx"
    sandbox = _Sandbox(
        {
            pptx_path: _pptx("hello"),
            docx_path: _docx("hello"),
        }
    )
    _patch_sandbox(monkeypatch, sandbox)

    with_window = json.loads(
        office_inspect_tool.func(
            SimpleNamespace(),
            pptx_path,
            start_slide=2,
            analysis_mode="pptx_quality_preflight",
        )
    )
    wrong_format = json.loads(
        office_inspect_tool.func(
            SimpleNamespace(),
            docx_path,
            analysis_mode="pptx_quality_preflight",
        )
    )

    assert with_window == {
        "ok": False,
        "error": ("PPTX quality preflight cannot be combined with inspect windows, selectors, or detail flags"),
    }
    assert wrong_format == {
        "ok": False,
        "error": "PPTX quality preflight requires a .pptx document",
    }


def test_office_inspect_rejects_unknown_analysis_mode_for_direct_callers(
    monkeypatch,
) -> None:
    path = "/mnt/user-data/uploads/in.pptx"
    sandbox = _Sandbox({path: _pptx("hello")})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_inspect_tool.func(
            SimpleNamespace(),
            path,
            analysis_mode="unexpected",
        )
    )

    assert result == {
        "ok": False,
        "error": "Unsupported Office inspection analysis mode: unexpected",
    }


def test_office_inspect_async_forwards_pptx_options_and_selector(monkeypatch) -> None:
    captured: dict[str, object] = {}

    async def fake_initialize(_runtime) -> None:
        return None

    def fake_inspect(*args, **kwargs) -> str:
        captured["args"] = args
        captured["kwargs"] = kwargs
        return "inspected"

    monkeypatch.setattr(
        "vassilflow.community.office.tools.ensure_sandbox_initialized_async",
        fake_initialize,
    )
    monkeypatch.setattr(office_inspect_tool, "func", fake_inspect)
    selector = PptxObjectSelector(paths=["/slide[2]/shape[@id=7]"])

    result = asyncio.run(
        _office_inspect_async(
            SimpleNamespace(),
            "/mnt/user-data/uploads/in.pptx",
            start_slide=2,
            max_slides=3,
            include_pptx_formatting=True,
            include_pptx_annotations=True,
            include_pptx_dynamics=True,
            include_pptx_media_gc_plan=True,
            pptx_selector=selector,
        )
    )

    assert result == "inspected"
    assert captured["args"] == (
        SimpleNamespace(),
        "/mnt/user-data/uploads/in.pptx",
    )
    assert captured["kwargs"] == {
        "start_paragraph": 1,
        "max_paragraphs": 100,
        "include_runs": False,
        "sheet_name": None,
        "cell_range": None,
        "include_cell_styles": False,
        "start_slide": 2,
        "max_slides": 3,
        "include_pptx_formatting": True,
        "pptx_selector": selector,
        "include_pptx_annotations": True,
        "include_pptx_dynamics": True,
        "include_pptx_media_gc_plan": True,
        "analysis_mode": "inspect",
    }


def test_office_edit_keeps_upload_immutable_and_writes_valid_output(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output = "/mnt/user-data/workspace/out.docx"
    original = _docx("old value")
    sandbox = _Sandbox({source: original})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.func(
            SimpleNamespace(),
            source,
            output,
            [DocxTextReplacement(find="old", replace="new")],
        )
    )

    assert result["ok"] is True
    assert result["commit_status"] == "committed"
    assert result["gateway_mirror_status"] == "available"
    assert result["replacement_count"] == 1
    assert result["validation"]["valid"] is True
    assert result["project"] == {
        "schema": "vassilflow.office.project.v1",
        "project_id": _FAKE_PROJECT_ID,
        "created": True,
        "current_revision_id": _FAKE_REVISION_ID,
        "revision_count": 2,
        "storage_scope": "trusted_user",
    }
    assert result["revision"]["revision_id"] == _FAKE_REVISION_ID
    assert result["revision"]["parent_revision_id"] == _FAKE_BASELINE_REVISION_ID
    assert sandbox.files[source] == original
    receipt = result["receipt"]
    assert receipt["source"]["sha256"] == hashlib.sha256(original).hexdigest()
    assert receipt["result"]["sha256"] == hashlib.sha256(sandbox.files[output]).hexdigest()
    assert receipt["applied_operation_ids"] == [result["operations"][0]["operation_id"]]
    assert receipt["semantic_changes"]["changed_target_paths"] == ["/document/paragraph[1]"]
    inspected = json.loads(office_inspect_tool.func(SimpleNamespace(), output))
    assert inspected["paragraphs"][0]["text"] == "new value"


def test_office_edit_does_not_commit_when_receipt_construction_fails(
    monkeypatch,
) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output = "/mnt/user-data/workspace/out.docx"
    original = _docx("old value")
    sandbox = _Sandbox({source: original})
    _patch_sandbox(monkeypatch, sandbox)

    def fail_receipt(*_args, **_kwargs):
        raise RuntimeError("receipt construction failed")

    monkeypatch.setattr(
        "vassilflow.community.office.engine.build_semantic_change_receipt",
        fail_receipt,
    )

    result = json.loads(
        office_edit_tool.func(
            SimpleNamespace(),
            source,
            output,
            [DocxTextReplacement(find="old", replace="new")],
        )
    )

    assert result["ok"] is False
    assert sandbox.files[source] == original
    assert output not in sandbox.files


def test_office_edit_writes_bounded_pptx_text_replacement(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.pptx"
    output = "/mnt/user-data/workspace/out.pptx"
    original = _pptx("old value")
    sandbox = _Sandbox({source: original})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.func(
            SimpleNamespace(),
            source,
            output,
            [
                PptxTextReplacement(
                    paths=["/slide[1]/shape[@id=2]"],
                    find="old",
                    replace="new",
                )
            ],
        )
    )

    assert result["ok"] is True
    assert result["commit_status"] == "committed"
    assert result["replacement_count"] == 1
    assert result["validation"]["valid"] is True
    assert sandbox.files[source] == original
    with zipfile.ZipFile(io.BytesIO(sandbox.files[output])) as archive:
        slide = archive.read("ppt/slides/slide1.xml")
    assert b"new value" in slide
    assert b"old value" not in slide


def test_office_edit_invoke_applies_typed_pptx_formatting_payload(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.pptx"
    output = "/mnt/user-data/workspace/formatted.pptx"
    original = _pptx("target")
    sandbox = _Sandbox({source: original})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.invoke(
            {
                "runtime": _runtime(),
                "source_path": source,
                "output_path": output,
                "operations": [
                    {
                        "type": "format_pptx_runs",
                        "runs": {
                            "targets": [
                                {
                                    "path": "/slide[1]/shape[@id=2]/paragraph[1]/run[1]",
                                    "expected_text": "target",
                                }
                            ]
                        },
                        "formatting": {
                            "bold": True,
                            "color": "#FF0000",
                        },
                    },
                    {
                        "type": "format_pptx_paragraphs",
                        "paragraphs": {
                            "targets": [
                                {
                                    "path": "/slide[1]/shape[@id=2]/paragraph[1]",
                                    "expected_text": "target",
                                }
                            ]
                        },
                        "formatting": {
                            "alignment": "center",
                            "space_after": 6.5,
                        },
                    },
                    {
                        "type": "format_pptx_shapes",
                        "shapes": {
                            "targets": [
                                {
                                    "path": "/slide[1]/shape[@id=2]",
                                    "expected_name": "Text",
                                }
                            ]
                        },
                        "formatting": {
                            "fill": {"type": "solid", "color": "#336699"},
                            "line": {
                                "fill": {"type": "solid", "color": "#778899"},
                                "width_points": 2.5,
                                "cap": "square",
                                "dash": "large_dash",
                                "join": "miter",
                                "miter_limit_percent": 800,
                            },
                            "text_box": {
                                "margin_left": 8,
                                "vertical_anchor": "middle",
                            },
                        },
                    },
                ],
            }
        )
    )

    assert result["ok"] is True
    assert result["change_count"] == 3
    assert result["replacement_count"] == 0
    assert sandbox.files[source] == original
    inspected = json.loads(
        office_inspect_tool.func(
            SimpleNamespace(),
            output,
            include_pptx_formatting=True,
        )
    )
    paragraph = inspected["slides"][0]["objects"][0]["text_body"]["paragraphs"][0]
    assert paragraph["formatting"]["alignment"] == "center"
    assert paragraph["formatting"]["space_after"] == {
        "unit": "points",
        "value": 6.5,
    }
    assert paragraph["segments"][0]["formatting"]["bold"] is True
    assert paragraph["segments"][0]["formatting"]["fill"] == {
        "type": "solid",
        "color": {"type": "rgb", "value": "#FF0000"},
    }
    style = inspected["slides"][0]["objects"][0]["style"]
    assert style["fill"]["color"]["value"] == "#336699"
    assert style["line"] == {
        "width_emu": 31_750,
        "cap": "square",
        "fill": {
            "type": "solid",
            "color": {"type": "rgb", "value": "#778899"},
        },
        "dash": "large_dash",
        "join": "miter",
        "miter_limit_percent": 800.0,
    }
    assert style["text_box"]["margins"]["left_emu"] == 101_600
    assert style["text_box"]["vertical_anchor"] == "middle"


@pytest.mark.parametrize(
    ("image_name", "image_payload", "media_part"),
    [
        ("fill.png", _PNG, "ppt/media/image1.png"),
        ("fill.jpg", _JPEG, "ppt/media/image1.jpg"),
    ],
)
def test_office_edit_embeds_locked_image_asset_without_mutating_source(
    monkeypatch,
    image_name: str,
    image_payload: bytes,
    media_part: str,
) -> None:
    source = "/mnt/user-data/uploads/in.pptx"
    image_path = f"/mnt/user-data/uploads/{image_name}"
    output = "/mnt/user-data/workspace/image-filled.pptx"
    original = _pptx("target")

    class RecordingSandbox(_Sandbox):
        def __init__(self, files: dict[str, bytes]) -> None:
            super().__init__(files)
            self.downloads: list[str] = []

        def download_file(self, path: str) -> bytes:
            self.downloads.append(path)
            return super().download_file(path)

    sandbox = RecordingSandbox({source: original, image_path: image_payload})
    _patch_sandbox(monkeypatch, sandbox)
    locked_paths: list[str] = []

    class Lock:
        def __init__(self, path: str) -> None:
            self.path = path

        def __enter__(self):
            locked_paths.append(self.path)

        def __exit__(self, *_args) -> None:
            return None

    monkeypatch.setattr(
        "vassilflow.community.office.tools.get_file_operation_lock",
        lambda _sandbox, path: Lock(path),
    )

    result = json.loads(
        office_edit_tool.invoke(
            {
                "runtime": _runtime(),
                "source_path": source,
                "output_path": output,
                "operations": [
                    {
                        "type": "format_pptx_shapes",
                        "shapes": {
                            "targets": [
                                {
                                    "path": "/slide[1]/shape[@id=2]",
                                    "expected_name": "Text",
                                }
                            ]
                        },
                        "formatting": {
                            "fill": {
                                "type": "image",
                                "image_path": image_path,
                                "rotate_with_shape": False,
                                "crop": {
                                    "left_percent": 5,
                                    "top_percent": 10,
                                    "right_percent": 15,
                                    "bottom_percent": 20,
                                },
                            }
                        },
                    }
                ],
            }
        )
    )

    assert result["ok"] is True
    assert locked_paths == sorted({source, image_path, output})
    assert image_path in sandbox.downloads
    assert sandbox.files[source] == original
    assert sandbox.files[image_path] == image_payload
    with zipfile.ZipFile(io.BytesIO(sandbox.files[output])) as archive:
        assert archive.read(media_part) == image_payload


def test_office_edit_embeds_locked_slide_background_asset(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.pptx"
    image_path = "/mnt/user-data/uploads/background.png"
    output = "/mnt/user-data/workspace/background-filled.pptx"
    original = _pptx("target")

    class RecordingSandbox(_Sandbox):
        def __init__(self, files: dict[str, bytes]) -> None:
            super().__init__(files)
            self.downloads: list[str] = []

        def download_file(self, path: str) -> bytes:
            self.downloads.append(path)
            return super().download_file(path)

    sandbox = RecordingSandbox({source: original, image_path: _PNG})
    _patch_sandbox(monkeypatch, sandbox)
    locked_paths: list[str] = []

    class Lock:
        def __init__(self, path: str) -> None:
            self.path = path

        def __enter__(self):
            locked_paths.append(self.path)

        def __exit__(self, *_args) -> None:
            return None

    monkeypatch.setattr(
        "vassilflow.community.office.tools.get_file_operation_lock",
        lambda _sandbox, path: Lock(path),
    )

    result = json.loads(
        office_edit_tool.invoke(
            {
                "runtime": _runtime(),
                "source_path": source,
                "output_path": output,
                "operations": [
                    {
                        "type": "format_pptx_slide_backgrounds",
                        "slides": {
                            "targets": [
                                {
                                    "path": "/slide[1]",
                                    "expected_part_name": "ppt/slides/slide1.xml",
                                }
                            ]
                        },
                        "formatting": {
                            "type": "image",
                            "image_path": image_path,
                            "mode": "center",
                        },
                    }
                ],
            }
        )
    )

    assert result["ok"] is True
    assert locked_paths == sorted({source, image_path, output})
    assert image_path in sandbox.downloads
    assert sandbox.files[source] == original
    assert sandbox.files[image_path] == _PNG
    inspected = office_engine.inspect(
        sandbox.files[output],
        suffix=".pptx",
        max_slides=1,
    )
    background = inspected["slides"][0]["background"]
    assert {key: value for key, value in background.items() if key != "relationship_ids"} == {
        "scope": "direct",
        "type": "image",
        "fill_mode": "center",
    }
    with zipfile.ZipFile(io.BytesIO(sandbox.files[output])) as archive:
        assert archive.read("ppt/media/image1.png") == _PNG


def test_office_edit_replaces_locked_picture_source_without_mutating_inputs(
    monkeypatch,
) -> None:
    source = "/mnt/user-data/uploads/in.pptx"
    image_path = "/mnt/user-data/uploads/replacement.jpg"
    output = "/mnt/user-data/workspace/picture-replaced.pptx"
    original = _pptx("target", picture_source=_PNG)

    class RecordingSandbox(_Sandbox):
        def __init__(self, files: dict[str, bytes]) -> None:
            super().__init__(files)
            self.downloads: list[str] = []

        def download_file(self, path: str) -> bytes:
            self.downloads.append(path)
            return super().download_file(path)

    sandbox = RecordingSandbox(
        {
            source: original,
            image_path: _JPEG,
        }
    )
    _patch_sandbox(monkeypatch, sandbox)
    locked_paths: list[str] = []

    class Lock:
        def __init__(self, path: str) -> None:
            self.path = path

        def __enter__(self):
            locked_paths.append(self.path)

        def __exit__(self, *_args) -> None:
            return None

    monkeypatch.setattr(
        "vassilflow.community.office.tools.get_file_operation_lock",
        lambda _sandbox, path: Lock(path),
    )

    result = json.loads(
        office_edit_tool.invoke(
            {
                "runtime": _runtime(),
                "source_path": source,
                "output_path": output,
                "operations": [
                    {
                        "type": "replace_pptx_picture_sources",
                        "pictures": {
                            "targets": [
                                {
                                    "path": "/slide[1]/picture[@id=3]",
                                    "expected_name": "Picture",
                                    "expected_source_sha256": hashlib.sha256(_PNG).hexdigest(),
                                }
                            ]
                        },
                        "image_path": image_path,
                    }
                ],
            }
        )
    )

    assert result["ok"] is True
    assert result["change_count"] == 1
    assert result["replacement_count"] == 1
    assert locked_paths == sorted({source, image_path, output})
    assert source in sandbox.downloads
    assert image_path in sandbox.downloads
    assert sandbox.files[source] == original
    assert sandbox.files[image_path] == _JPEG
    inspected = office_engine.inspect(
        sandbox.files[output],
        suffix=".pptx",
        max_slides=1,
        include_pptx_formatting=True,
    )
    picture = next(item for item in inspected["slides"][0]["objects"] if item["kind"] == "picture")
    assert picture["alt_text"] == "Preserved alt text"
    assert picture["geometry"] == {
        "x_emu": 10,
        "y_emu": 20,
        "width_emu": 30,
        "height_emu": 40,
    }
    assert picture["style"]["picture"]["crop"] == {
        "left_percent": 5.0,
        "top_percent": 10.0,
        "right_percent": 0.0,
        "bottom_percent": 0.0,
    }
    assert picture["style"]["picture"]["effects"] == {"grayscale": True}
    assert picture["source"]["part_name"] == "ppt/media/image2.jpg"
    assert picture["source"]["content_type"] == "image/jpeg"
    assert picture["source"]["sha256"] == hashlib.sha256(_JPEG).hexdigest()
    with zipfile.ZipFile(io.BytesIO(sandbox.files[output])) as archive:
        assert archive.read("ppt/media/image1.png") == _PNG
        assert archive.read("ppt/media/image2.jpg") == _JPEG


def test_office_edit_rejects_non_workspace_image_source_before_writing(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.pptx"
    output = "/mnt/user-data/workspace/image-filled.pptx"
    sandbox = _Sandbox({source: _pptx("target")})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.invoke(
            {
                "runtime": _runtime(),
                "source_path": source,
                "output_path": output,
                "operations": [
                    {
                        "type": "format_pptx_shapes",
                        "shapes": {
                            "targets": [
                                {
                                    "path": "/slide[1]/shape[@id=2]",
                                    "expected_name": "Text",
                                }
                            ]
                        },
                        "formatting": {
                            "fill": {
                                "type": "image",
                                "image_path": "https://example.test/fill.png",
                                "rotate_with_shape": False,
                            }
                        },
                    }
                ],
            }
        )
    )

    assert result["ok"] is False
    assert "absolute /mnt/user-data path" in result["error"]
    assert output not in sandbox.files


def test_office_edit_invoke_validates_nested_operation_payload(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output = "/mnt/user-data/workspace/out.docx"
    sandbox = _Sandbox({source: _docx("old value")})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.invoke(
            {
                "runtime": _runtime(),
                "source_path": source,
                "output_path": output,
                "operations": [{"type": "replace_text", "find": "old", "replace": "new"}],
            }
        )
    )

    assert result["ok"] is True
    assert result["replacement_count"] == 1


def test_office_edit_invoke_applies_typed_formatting_payload(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output = "/mnt/user-data/workspace/formatted.docx"
    sandbox = _Sandbox({source: _docx("target")})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.invoke(
            {
                "runtime": _runtime(),
                "source_path": source,
                "output_path": output,
                "operations": [
                    {
                        "type": "format_runs",
                        "paragraphs": {"paragraph_indices": [1]},
                        "runs": {"contains_text": "target"},
                        "formatting": {"bold": True, "color": "#FF0000"},
                    }
                ],
            }
        )
    )

    assert result["ok"] is True
    assert result["change_count"] == 1
    assert result["replacement_count"] == 0
    inspected = json.loads(office_inspect_tool.func(SimpleNamespace(), output, include_runs=True))
    assert inspected["paragraphs"][0]["runs"][0]["formatting"]["bold"] is True
    assert inspected["paragraphs"][0]["runs"][0]["formatting"]["color"] == "#FF0000"


def test_office_edit_invoke_applies_typed_xlsx_formatting_payload(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.xlsx"
    output = "/mnt/user-data/workspace/formatted.xlsx"
    original = _xlsx()
    sandbox = _Sandbox({source: original})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.invoke(
            {
                "runtime": _runtime(),
                "source_path": source,
                "output_path": output,
                "operations": [
                    {
                        "type": "format_cells",
                        "cells": {"sheet_name": "Data", "ranges": ["B1"]},
                        "formatting": {"bold": True, "fill_color": "#FFF2CC"},
                    }
                ],
            }
        )
    )

    assert result["ok"] is True
    assert result["change_count"] == 1
    assert result["replacement_count"] == 0
    assert sandbox.files[source] == original
    workbook = load_workbook(io.BytesIO(sandbox.files[output]))
    try:
        assert workbook["Data"]["B1"].font.bold is True
        assert workbook["Data"]["B1"].fill.fgColor.rgb == "FFFFF2CC"
    finally:
        workbook.close()


def test_office_edit_refuses_to_write_into_uploads(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    sandbox = _Sandbox({source: _docx("old")})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.func(
            SimpleNamespace(),
            source,
            "/mnt/user-data/uploads/out.docx",
            [DocxTextReplacement(find="old", replace="new")],
        )
    )

    assert result["ok"] is False
    assert "/mnt/user-data/uploads/out.docx" not in sandbox.files


def test_office_edit_does_not_overwrite_a_different_existing_output_by_default(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output = "/mnt/user-data/outputs/result.docx"
    existing = _docx("keep me")
    sandbox = _Sandbox({source: _docx("old"), output: existing})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.func(
            SimpleNamespace(),
            source,
            output,
            [DocxTextReplacement(find="old", replace="new")],
        )
    )

    assert result["ok"] is False
    assert sandbox.files[output] == existing


def test_office_edit_rejects_in_place_write_and_preserves_source(monkeypatch) -> None:
    source = "/mnt/user-data/workspace/in-place.docx"
    original = _docx("old")
    sandbox = _Sandbox({source: original})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.func(
            SimpleNamespace(),
            source,
            source,
            [DocxTextReplacement(find="old", replace="new")],
            True,
        )
    )

    assert result["ok"] is False
    assert "In-place" in result["error"]
    assert sandbox.files[source] == original


def test_office_edit_atomic_replace_failure_preserves_existing_output(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output = "/mnt/user-data/workspace/existing.docx"
    existing = _docx("keep")

    class ReplaceFailingSandbox(_Sandbox):
        def replace_file(self, source_path: str, destination_path: str) -> None:
            del source_path, destination_path
            raise OSError(errno.EIO, "replace failed")

    sandbox = ReplaceFailingSandbox({source: _docx("old"), output: existing})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_edit_tool.func(
            SimpleNamespace(),
            source,
            output,
            [DocxTextReplacement(find="old", replace="new")],
            True,
        )
    )

    assert result["ok"] is False
    assert result["commit_status"] == "revision_committed_output_failed"
    assert result["project"]["current_revision_id"] == _FAKE_REVISION_ID
    assert result["revision"]["revision_id"] == _FAKE_REVISION_ID
    assert sandbox.files[source] != sandbox.files[output]
    assert sandbox.files[output] == existing


def test_office_edit_persists_and_continues_a_real_revision_chain(
    monkeypatch,
    tmp_path,
) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    first_output = "/mnt/user-data/workspace/v1.docx"
    second_output = "/mnt/user-data/workspace/v2.docx"
    rejected_output = "/mnt/user-data/workspace/rejected.docx"
    stale_output = "/mnt/user-data/workspace/stale.docx"
    original = _docx("first value")
    sandbox = _Sandbox({source: original})
    _patch_sandbox(monkeypatch, sandbox)
    store = OfficeRevisionStore(tmp_path / "office")
    monkeypatch.setattr(
        "vassilflow.community.office.tools._office_revision_store",
        lambda _runtime: store,
    )

    first = json.loads(
        office_edit_tool.func(
            _runtime(),
            source,
            first_output,
            [DocxTextReplacement(find="first", replace="second")],
        )
    )
    rejected = json.loads(
        office_edit_tool.func(
            _runtime(),
            first_output,
            rejected_output,
            [DocxTextReplacement(find="second", replace="rejected")],
            False,
            first["project"]["project_id"],
        )
    )
    second = json.loads(
        office_edit_tool.func(
            _runtime(),
            first_output,
            second_output,
            [DocxTextReplacement(find="second", replace="third")],
            False,
            first["project"]["project_id"],
            first["revision"]["revision_id"],
        )
    )
    stale = json.loads(
        office_edit_tool.func(
            _runtime(),
            first_output,
            stale_output,
            [DocxTextReplacement(find="second", replace="stale")],
            False,
            first["project"]["project_id"],
            first["revision"]["revision_id"],
        )
    )

    assert first["ok"] is True
    assert first["project"]["created"] is True
    assert rejected["ok"] is False
    assert "provided together" in rejected["error"]
    assert rejected_output not in sandbox.files
    assert second["ok"] is True
    assert second["project"]["created"] is False
    assert second["project"]["project_id"] == first["project"]["project_id"]
    assert second["revision"]["parent_revision_id"] == first["revision"]["revision_id"]
    assert second["revision"]["sequence"] == 3
    assert stale["ok"] is False
    assert "newer current revision" in stale["error"]
    assert stale_output not in sandbox.files
    project = store.load_project(first["project"]["project_id"])
    assert project["primary_thread_id"] == "office-thread"
    assert project["current_revision_id"] == second["revision"]["revision_id"]
    assert project["revision_count"] == 3


def test_office_edit_reads_verified_selection_from_canonical_project(
    monkeypatch,
    tmp_path,
) -> None:
    output = "/mnt/user-data/workspace/selected-object.pptx"
    sandbox = _Sandbox({})
    _patch_sandbox(monkeypatch, sandbox)
    store = OfficeRevisionStore(tmp_path / "office")
    runtime, project_id, revision_id, source = _seed_selected_pptx(store)
    monkeypatch.setattr(
        "vassilflow.community.office.tools._office_revision_store",
        lambda _runtime: store,
    )

    result = json.loads(
        office_edit_tool.func(
            runtime,
            None,
            output,
            [
                PptxTextReplacement(
                    paths=["/slide[1]/shape[@id=3]"],
                    find="Alpha",
                    replace="Gamma",
                )
            ],
        )
    )

    assert result["ok"] is True
    assert result["source_path"] is None
    assert result["project"]["project_id"] == project_id
    assert result["revision"]["parent_revision_id"] == revision_id
    assert sandbox.files[output] != source
    with zipfile.ZipFile(io.BytesIO(sandbox.files[output])) as archive:
        slide = archive.read("ppt/slides/slide1.xml")
    assert b"Gamma" in slide
    assert b"Alpha" not in slide
    current = store.load_project(project_id)
    assert current["revision_count"] == 3
    assert current["current_revision_id"] == result["revision"]["revision_id"]


def test_office_edit_selection_rejects_source_path_and_sibling_target(
    monkeypatch,
    tmp_path,
) -> None:
    output = "/mnt/user-data/workspace/rejected-selection.pptx"
    sandbox = _Sandbox({})
    _patch_sandbox(monkeypatch, sandbox)
    store = OfficeRevisionStore(tmp_path / "office")
    runtime, project_id, _revision_id, _source = _seed_selected_pptx(store)
    monkeypatch.setattr(
        "vassilflow.community.office.tools._office_revision_store",
        lambda _runtime: store,
    )

    with_source = json.loads(
        office_edit_tool.func(
            runtime,
            "/mnt/user-data/uploads/spoofed.pptx",
            output,
            [
                PptxTextReplacement(
                    paths=["/slide[1]/shape[@id=3]"],
                    find="Alpha",
                    replace="Gamma",
                )
            ],
        )
    )
    sibling = json.loads(
        office_edit_tool.func(
            runtime,
            None,
            output,
            [
                PptxTextReplacement(
                    paths=["/slide[1]/shape[@id=5]"],
                    find="Reviewed",
                    replace="Escaped",
                )
            ],
        )
    )

    assert with_source == {
        "ok": False,
        "error": "source_path must be null for a verified Office project selection",
    }
    assert sibling == {
        "ok": False,
        "error": "Office edit operation escapes the selected object scope",
    }
    assert output not in sandbox.files
    assert store.load_project(project_id)["revision_count"] == 2


def test_office_edit_selection_rejects_out_of_scope_semantic_receipt(
    monkeypatch,
    tmp_path,
) -> None:
    output = "/mnt/user-data/workspace/rejected-receipt.pptx"
    sandbox = _Sandbox({})
    _patch_sandbox(monkeypatch, sandbox)
    store = OfficeRevisionStore(tmp_path / "office")
    runtime, project_id, _revision_id, _source = _seed_selected_pptx(store)
    monkeypatch.setattr(
        "vassilflow.community.office.tools._office_revision_store",
        lambda _runtime: store,
    )
    original_edit = office_engine.edit_with_receipt

    def tamper_receipt(*args, **kwargs):
        edited, reports, receipt = original_edit(*args, **kwargs)
        semantic = receipt["semantic_changes"]
        semantic["changed_target_paths"] = ["/slide[1]/shape[@id=5]"]
        semantic["semantic_deltas"][0]["path"] = "/slide[1]/shape[@id=5]"
        return edited, reports, receipt

    monkeypatch.setattr(office_engine, "edit_with_receipt", tamper_receipt)

    result = json.loads(
        office_edit_tool.func(
            runtime,
            None,
            output,
            [
                PptxTextReplacement(
                    paths=["/slide[1]/shape[@id=3]"],
                    find="Alpha",
                    replace="Gamma",
                )
            ],
        )
    )

    assert result == {
        "ok": False,
        "error": "Office selection semantic changes escape the selected object scope",
    }
    assert output not in sandbox.files
    assert store.load_project(project_id)["revision_count"] == 2


def test_remote_office_artifact_is_mirrored_for_gateway_file_tools(monkeypatch, tmp_path) -> None:
    workspace = tmp_path / "workspace"
    uploads = tmp_path / "uploads"
    outputs = tmp_path / "outputs"
    for directory in (workspace, uploads, outputs):
        directory.mkdir()
    runtime = SimpleNamespace(
        state={
            "thread_data": {
                "workspace_path": str(workspace),
                "uploads_path": str(uploads),
                "outputs_path": str(outputs),
            }
        }
    )
    monkeypatch.setattr(
        "vassilflow.sandbox.sandbox_provider.get_sandbox_provider",
        lambda: SimpleNamespace(uses_thread_data_mounts=False),
    )

    _mirror_binary_to_gateway_if_needed(
        runtime,
        "/mnt/user-data/workspace/qa/page-001.png",
        _PNG,
    )

    assert (workspace / "qa" / "page-001.png").read_bytes() == _PNG


def test_remote_office_mirror_failure_is_reported_without_raising(monkeypatch) -> None:
    monkeypatch.setattr(
        "vassilflow.sandbox.sandbox_provider.get_sandbox_provider",
        lambda: SimpleNamespace(uses_thread_data_mounts=False),
    )

    mirrored = _mirror_binary_to_gateway_if_needed(
        SimpleNamespace(state={}),
        "/mnt/user-data/workspace/qa/page-001.png",
        _PNG,
    )

    assert mirrored is False


def test_office_edit_reports_committed_output_when_gateway_mirror_is_unavailable(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output = "/mnt/user-data/workspace/out.docx"
    sandbox = _Sandbox({source: _docx("old")})
    _patch_sandbox(monkeypatch, sandbox)
    monkeypatch.setattr(
        "vassilflow.community.office.tools._mirror_binary_to_gateway_if_needed",
        lambda *_args, **_kwargs: False,
    )

    result = json.loads(
        office_edit_tool.func(
            SimpleNamespace(state={}),
            source,
            output,
            [DocxTextReplacement(find="old", replace="new")],
        )
    )

    assert result["ok"] is True
    assert result["commit_status"] == "committed"
    assert result["gateway_mirror_status"] == "unavailable"
    assert "committed in the sandbox" in result["warning"]
    assert output in sandbox.files


def test_visual_review_requires_external_client_when_gateway_image_is_unavailable() -> None:
    runtime = SimpleNamespace(tools=[SimpleNamespace(name="view_image")])

    status, _ = _visual_review_contract(runtime, gateway_images_available=False)

    assert status == "external_review_required"


def test_office_render_writes_pages_then_pending_review_manifest(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.pptx"
    output_dir = "/mnt/user-data/workspace/visual-qa"
    sandbox = _Sandbox({source: b"pptx"})
    _patch_sandbox(monkeypatch, sandbox)
    monkeypatch.setattr(
        "vassilflow.community.office.tools.office_engine.render",
        lambda *_args, **_kwargs: OfficeRenderResult(
            format="pptx",
            renderer="test-renderer",
            renderer_version="1.0",
            pdfium_version="5.11.0",
            source_sha256="a" * 64,
            pipeline_fingerprint="b" * 64,
            page_count=1,
            start_page=1,
            dpi=120,
            pages=(
                RenderedOfficePage(
                    page=1,
                    filename="page-001.png",
                    width=1,
                    height=1,
                    sha256="digest",
                    data=_PNG,
                    source_slide=1,
                ),
            ),
        ),
    )

    result = json.loads(
        office_render_tool.func(
            SimpleNamespace(tools=[SimpleNamespace(name="view_image")]),
            source,
            output_dir,
        )
    )

    assert result["ok"] is True
    assert result["commit_status"] == "committed"
    assert result["gateway_mirror_status"] == "available"
    assert result["page_paths"] == [f"{output_dir}/page-001.png"]
    assert result["visual_review_status"] == "pending"
    assert sandbox.files[f"{output_dir}/page-001.png"] == _PNG
    manifest = json.loads(sandbox.files[f"{output_dir}/render-manifest.json"])
    assert manifest["complete"] is True
    assert manifest["source_sha256"] == "a" * 64
    assert manifest["pipeline_fingerprint"] == "b" * 64
    assert manifest["pages"][0]["source_slide"] == 1
    assert manifest["visual_review_status"] == "pending"
    assert manifest["gateway_page_access"] == "available"


def test_office_render_persists_preview_evidence_for_exact_revision(
    monkeypatch,
    tmp_path,
) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    edited_path = "/mnt/user-data/workspace/edited.docx"
    output_dir = "/mnt/user-data/workspace/rendered"
    sandbox = _Sandbox({source: _docx("old")})
    _patch_sandbox(monkeypatch, sandbox)
    store = OfficeRevisionStore(tmp_path / "office")
    monkeypatch.setattr(
        "vassilflow.community.office.tools._office_revision_store",
        lambda _runtime: store,
    )
    edited = json.loads(
        office_edit_tool.func(
            _runtime(),
            source,
            edited_path,
            [DocxTextReplacement(find="old", replace="new")],
        )
    )
    edited_bytes = sandbox.files[edited_path]
    page_sha256 = hashlib.sha256(_PNG).hexdigest()
    monkeypatch.setattr(
        "vassilflow.community.office.tools.office_engine.render",
        lambda *_args, **_kwargs: OfficeRenderResult(
            format="docx",
            renderer="test-renderer",
            renderer_version="1.0",
            pdfium_version="5.11.0",
            source_sha256=hashlib.sha256(edited_bytes).hexdigest(),
            pipeline_fingerprint="b" * 64,
            page_count=1,
            start_page=1,
            dpi=120,
            pages=(
                RenderedOfficePage(
                    page=1,
                    filename="page-001.png",
                    width=1,
                    height=1,
                    sha256=page_sha256,
                    data=_PNG,
                ),
            ),
        ),
    )

    rendered = json.loads(
        office_render_tool.func(
            SimpleNamespace(tools=[SimpleNamespace(name="view_image")]),
            edited_path,
            output_dir,
            1,
            1,
            120,
            edited["project"]["project_id"],
            edited["revision"]["revision_id"],
        )
    )

    assert rendered["ok"] is True
    evidence = rendered["render_evidence"]
    assert evidence["schema"] == "vassilflow.office.render_evidence.v1"
    assert evidence["project_id"] == edited["project"]["project_id"]
    assert evidence["revision_id"] == edited["revision"]["revision_id"]
    assert evidence["source_sha256"] == hashlib.sha256(edited_bytes).hexdigest()
    manifest = json.loads(sandbox.files[f"{output_dir}/render-manifest.json"])
    assert manifest["schema"] == "vassilflow.office.render_manifest.v1"
    assert manifest["project_id"] == evidence["project_id"]
    assert manifest["revision_id"] == evidence["revision_id"]
    persisted = store.load_render_evidence(evidence["project_id"], evidence["evidence_id"])
    assert persisted["render_manifest"] == manifest
    assert persisted["preview_pages"][0]["sha256"] == page_sha256


def test_office_render_rejects_partial_revision_binding_before_render(
    monkeypatch,
) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output_dir = "/mnt/user-data/workspace/rendered"
    sandbox = _Sandbox({source: _docx("hello")})
    _patch_sandbox(monkeypatch, sandbox)
    render_called = False

    def render_should_not_run(*_args, **_kwargs):
        nonlocal render_called
        render_called = True
        raise AssertionError("render should not run")

    monkeypatch.setattr(
        "vassilflow.community.office.tools.office_engine.render",
        render_should_not_run,
    )

    result = json.loads(
        office_render_tool.func(
            SimpleNamespace(),
            source,
            output_dir,
            project_id=_FAKE_PROJECT_ID,
        )
    )

    assert result["ok"] is False
    assert "provided together" in result["error"]
    assert render_called is False
    assert not any(path.startswith(f"{output_dir}/") for path in sandbox.files)


def test_office_render_reports_committed_evidence_when_manifest_write_fails(
    monkeypatch,
    tmp_path,
) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    edited_path = "/mnt/user-data/workspace/edited.docx"
    output_dir = "/mnt/user-data/workspace/rendered"

    class ManifestFailingSandbox(_Sandbox):
        def update_file(self, path: str, content: bytes) -> None:
            if path.endswith("render-manifest.json"):
                raise OSError(errno.EIO, "manifest write failed", path)
            super().update_file(path, content)

    sandbox = ManifestFailingSandbox({source: _docx("old")})
    _patch_sandbox(monkeypatch, sandbox)
    store = OfficeRevisionStore(tmp_path / "office")
    monkeypatch.setattr(
        "vassilflow.community.office.tools._office_revision_store",
        lambda _runtime: store,
    )
    edited = json.loads(
        office_edit_tool.func(
            _runtime(),
            source,
            edited_path,
            [DocxTextReplacement(find="old", replace="new")],
        )
    )
    edited_bytes = sandbox.files[edited_path]
    monkeypatch.setattr(
        "vassilflow.community.office.tools.office_engine.render",
        lambda *_args, **_kwargs: OfficeRenderResult(
            format="docx",
            renderer="test-renderer",
            renderer_version="1.0",
            pdfium_version="5.11.0",
            source_sha256=hashlib.sha256(edited_bytes).hexdigest(),
            pipeline_fingerprint="b" * 64,
            page_count=1,
            start_page=1,
            dpi=120,
            pages=(
                RenderedOfficePage(
                    page=1,
                    filename="page-001.png",
                    width=1,
                    height=1,
                    sha256=hashlib.sha256(_PNG).hexdigest(),
                    data=_PNG,
                ),
            ),
        ),
    )

    result = json.loads(
        office_render_tool.func(
            SimpleNamespace(tools=[SimpleNamespace(name="view_image")]),
            edited_path,
            output_dir,
            1,
            1,
            120,
            edited["project"]["project_id"],
            edited["revision"]["revision_id"],
        )
    )

    assert result["ok"] is False
    assert result["commit_status"] == "render_evidence_committed_output_failed"
    evidence = result["render_evidence"]
    assert f"{output_dir}/page-001.png" in sandbox.files
    assert f"{output_dir}/render-manifest.json" not in sandbox.files
    persisted = store.load_render_evidence(evidence["project_id"], evidence["evidence_id"])
    assert persisted["revision_id"] == edited["revision"]["revision_id"]


def test_office_render_commits_with_external_review_when_gateway_mirror_fails(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output_dir = "/mnt/user-data/workspace/remote-visual-qa"
    sandbox = _Sandbox({source: _docx("hello")})
    _patch_sandbox(monkeypatch, sandbox)
    monkeypatch.setattr(
        "vassilflow.community.office.tools._mirror_binary_to_gateway_if_needed",
        lambda *_args, **_kwargs: False,
    )
    monkeypatch.setattr(
        "vassilflow.community.office.tools.office_engine.render",
        lambda *_args, **_kwargs: OfficeRenderResult(
            format="docx",
            renderer="test-renderer",
            renderer_version="1.0",
            pdfium_version="5.11.0",
            source_sha256="a" * 64,
            pipeline_fingerprint="b" * 64,
            page_count=1,
            start_page=1,
            dpi=120,
            pages=(
                RenderedOfficePage(
                    page=1,
                    filename="page-001.png",
                    width=1,
                    height=1,
                    sha256="digest",
                    data=_PNG,
                ),
            ),
        ),
    )

    result = json.loads(
        office_render_tool.func(
            SimpleNamespace(state={}, tools=[SimpleNamespace(name="view_image")]),
            source,
            output_dir,
        )
    )

    assert result["ok"] is True
    assert result["commit_status"] == "committed"
    assert result["gateway_mirror_status"] == "unavailable"
    assert result["visual_review_status"] == "external_review_required"
    assert "committed in the sandbox" in result["warning"]
    manifest = json.loads(sandbox.files[f"{output_dir}/render-manifest.json"])
    assert manifest["complete"] is True
    assert manifest["gateway_page_access"] == "unavailable"


def test_office_render_requires_external_review_without_view_image(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output_dir = "/mnt/user-data/workspace/external-visual-qa"
    sandbox = _Sandbox({source: _docx("hello")})
    _patch_sandbox(monkeypatch, sandbox)
    monkeypatch.setattr(
        "vassilflow.community.office.tools.office_engine.render",
        lambda *_args, **_kwargs: OfficeRenderResult(
            format="docx",
            renderer="test-renderer",
            renderer_version="1.0",
            pdfium_version="5.11.0",
            source_sha256="a" * 64,
            pipeline_fingerprint="b" * 64,
            page_count=1,
            start_page=1,
            dpi=120,
            pages=(
                RenderedOfficePage(
                    page=1,
                    filename="page-001.png",
                    width=1,
                    height=1,
                    sha256="digest",
                    data=_PNG,
                ),
            ),
        ),
    )

    result = json.loads(office_render_tool.func(SimpleNamespace(tools=[]), source, output_dir))

    assert result["ok"] is True
    assert result["visual_review_status"] == "external_review_required"
    manifest = json.loads(sandbox.files[f"{output_dir}/render-manifest.json"])
    assert manifest["visual_review_status"] == "external_review_required"


def test_office_render_marks_partial_write_as_incomplete(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output_dir = "/mnt/user-data/workspace/partial-visual-qa"

    class FailingSandbox(_Sandbox):
        def update_file(self, path: str, content: bytes) -> None:
            if path.endswith("render-manifest.json"):
                raise OSError(errno.EIO, "manifest write failed", path)
            super().update_file(path, content)

    sandbox = FailingSandbox({source: _docx("hello")})
    _patch_sandbox(monkeypatch, sandbox)
    monkeypatch.setattr(
        "vassilflow.community.office.tools.office_engine.render",
        lambda *_args, **_kwargs: OfficeRenderResult(
            format="docx",
            renderer="test-renderer",
            renderer_version="1.0",
            pdfium_version="5.11.0",
            source_sha256="a" * 64,
            pipeline_fingerprint="b" * 64,
            page_count=1,
            start_page=1,
            dpi=120,
            pages=(
                RenderedOfficePage(
                    page=1,
                    filename="page-001.png",
                    width=1,
                    height=1,
                    sha256="digest",
                    data=_PNG,
                ),
            ),
        ),
    )

    result = json.loads(office_render_tool.func(SimpleNamespace(tools=[]), source, output_dir))

    assert result["ok"] is False
    assert "incomplete" in result["error"]
    assert f"{output_dir}/page-001.png" in sandbox.files
    assert f"{output_dir}/render-manifest.json" not in sandbox.files


def test_office_render_refuses_nonempty_output_directory_before_render(monkeypatch) -> None:
    source = "/mnt/user-data/uploads/in.docx"
    output_dir = "/mnt/user-data/workspace/visual-qa"
    existing_path = f"{output_dir}/keep.png"
    sandbox = _Sandbox({source: _docx("hello"), existing_path: _PNG})
    _patch_sandbox(monkeypatch, sandbox)

    result = json.loads(
        office_render_tool.func(
            SimpleNamespace(),
            source,
            output_dir,
        )
    )

    assert result["ok"] is False
    assert sandbox.files[existing_path] == _PNG
    assert f"{output_dir}/render-manifest.json" not in sandbox.files

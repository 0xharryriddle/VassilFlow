from __future__ import annotations

import hashlib
import io
import zipfile
from datetime import datetime

import pytest
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet.formula import ArrayFormula

from vassilflow.community.office import xlsx as xlsx_module
from vassilflow.community.office.engine import office_engine
from vassilflow.community.office.errors import OfficeOperationError, OfficePackageError
from vassilflow.community.office.models import (
    XlsxCellFormatOperation,
    XlsxCellFormatting,
    XlsxCellSelector,
)
from vassilflow.community.office.xlsx import edit_xlsx, inspect_xlsx, validate_xlsx

_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _replace_parts(
    data: bytes,
    replacements: dict[str, bytes],
    additions: dict[str, bytes] | None = None,
    removals: set[str] | None = None,
) -> bytes:
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(
            output,
            mode="w",
            compression=zipfile.ZIP_DEFLATED,
        ) as destination,
    ):
        for info in source.infolist():
            if info.filename in (removals or set()):
                continue
            destination.writestr(info, replacements.get(info.filename, source.read(info)))
        for name, payload in (additions or {}).items():
            destination.writestr(name, payload)
    return output.getvalue()


def _relocate_first_worksheet(data: bytes) -> bytes:
    source_part = "xl/worksheets/sheet1.xml"
    target_part = "xl/custom/sheet.xml"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        worksheet = archive.read(source_part)
        relationships = archive.read("xl/_rels/workbook.xml.rels").replace(
            b"worksheets/sheet1.xml",
            b"custom/sheet.xml",
        )
        content_types = archive.read("[Content_Types].xml").replace(
            b"/xl/worksheets/sheet1.xml",
            b"/xl/custom/sheet.xml",
        )
    return _replace_parts(
        data,
        {
            "xl/_rels/workbook.xml.rels": relationships,
            "[Content_Types].xml": content_types,
        },
        additions={target_part: worksheet},
        removals={source_part},
    )


def _convert_first_sheet_to_macro_sheet(
    data: bytes,
    *,
    relationship_type: str,
    content_type: str,
) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        content_types = etree.fromstring(archive.read("[Content_Types].xml"))
    worksheet_relationship = next(relationship for relationship in relationships if relationship.get("Target") == "/xl/worksheets/sheet1.xml" or relationship.get("Target") == "worksheets/sheet1.xml")
    worksheet_relationship.set("Type", relationship_type)
    worksheet_override = next(override for override in content_types if override.get("PartName") == "/xl/worksheets/sheet1.xml")
    worksheet_override.set("ContentType", content_type)
    return _replace_parts(
        data,
        {
            "xl/_rels/workbook.xml.rels": etree.tostring(
                relationships,
                encoding="UTF-8",
                xml_declaration=True,
            ),
            "[Content_Types].xml": etree.tostring(
                content_types,
                encoding="UTF-8",
                xml_declaration=True,
            ),
        },
    )


def _workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Metric"
    sheet["B1"] = "Value"
    sheet["A2"] = "Revenue"
    sheet["B2"] = 1234.5
    sheet["B2"].number_format = "#,##0.00"
    sheet["A3"] = "Projected"
    sheet["B3"] = "=B2*2"
    sheet["A4"] = "Approved"
    sheet["B4"] = True
    sheet["A5"] = "As of"
    sheet["B5"] = datetime(2026, 7, 14, 9, 30)
    sheet["C3"] = "#DIV/0!"
    sheet["C3"].data_type = "e"
    sheet["C2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    sheet["D1"] = "Merged heading"
    sheet.merge_cells("D1:E1")
    sheet.freeze_panes = "A2"
    sheet["A6"] = "one"
    sheet["A7"] = "two"
    sheet["A8"] = "three"
    for cell in sheet["A6:A8"]:
        cell[0].font = Font(italic=True, name="Calibri", size=11)
        cell[0].border = Border(bottom=Side(style="thin", color="FF112233"))
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "untouched"

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    data = output.getvalue()

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        sheet_payload = archive.read("xl/worksheets/sheet1.xml")
    root = etree.fromstring(sheet_payload)
    formula_cell = root.xpath(".//x:c[@r='B3']", namespaces={"x": _SHEET_NS})[0]
    value = formula_cell.find(f"{{{_SHEET_NS}}}v")
    assert value is not None
    value.text = "2469"
    return _replace_parts(
        data,
        {
            "xl/worksheets/sheet1.xml": etree.tostring(
                root,
                encoding="UTF-8",
                xml_declaration=True,
            )
        },
    )


def _cell_node(data: bytes, coordinate: str) -> etree._Element:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = etree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    return root.xpath(f".//x:c[@r='{coordinate}']", namespaces={"x": _SHEET_NS})[0]


def _cell_xf_count(data: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = etree.fromstring(archive.read("xl/styles.xml"))
    cell_xfs = root.find(f"{{{_SHEET_NS}}}cellXfs")
    assert cell_xfs is not None
    return len(cell_xfs)


def test_xlsx_inspection_exposes_bounded_values_formula_cache_and_styles() -> None:
    data = _workbook()

    summary = inspect_xlsx(data)
    inspected = inspect_xlsx(
        data,
        sheet_name="data",
        cell_range="A1:E5",
        include_cell_styles=True,
    )

    assert summary["format"] == "xlsx"
    assert [sheet["name"] for sheet in summary["sheets"]] == ["Data", "Hidden"]
    assert inspected["selected_sheet"] == "Data"
    assert inspected["range"] == "A1:E5"
    cells = {cell["coordinate"]: cell for cell in inspected["cells"]}
    assert cells["B2"]["value"] == 1234.5
    assert cells["B2"]["data_type"] == "number"
    assert cells["B3"]["formula"] == "=B2*2"
    assert cells["B3"]["cached_value"] == 2469
    assert cells["B4"]["data_type"] == "boolean"
    assert cells["B5"]["data_type"] == "date"
    assert cells["C3"]["data_type"] == "error"
    assert cells["C2"]["data_type"] == "blank"
    assert cells["C2"]["formatting"]["fill"]["foreground"]["value"] == "#FFFF00"
    assert inspected["merged_ranges"] == ["D1:E1"]


def test_xlsx_inspection_applies_a_shared_character_budget(monkeypatch) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "abcdefghij"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    monkeypatch.setattr(xlsx_module, "_MAX_INSPECT_OUTPUT_CHARS", 6)
    monkeypatch.setattr(xlsx_module, "_MAX_INSPECT_CELL_VALUE_CHARS", 100)

    inspected = inspect_xlsx(
        output.getvalue(),
        sheet_name="Sheet",
        cell_range="A1",
    )

    cell = inspected["cells"][0]
    assert cell["value"] == "abcdef"
    assert cell["value_truncated"] is True
    assert inspected["content_characters_returned"] == 6
    assert inspected["content_truncated"] is True


def test_xlsx_counts_relationship_resolved_worksheet_parts(monkeypatch) -> None:
    data = _relocate_first_worksheet(_workbook())
    expected = inspect_xlsx(_workbook())["declared_cell_count"]

    assert inspect_xlsx(data)["declared_cell_count"] == expected
    monkeypatch.setattr(xlsx_module, "_MAX_DECLARED_CELLS_PER_SHEET", 1)
    with pytest.raises(OfficePackageError, match="declares more than 1 cells"):
        inspect_xlsx(data)


def test_xlsx_rejects_unrecognized_relationship_type_for_declared_sheet() -> None:
    data = _workbook()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    first_sheet = next(relationship for relationship in relationships if (relationship.get("Type") or "").endswith("/worksheet"))
    first_sheet.set("Type", "urn:vassilflow:test:unsupported-sheet")
    malformed = _replace_parts(
        data,
        {
            "xl/_rels/workbook.xml.rels": etree.tostring(
                relationships,
                encoding="UTF-8",
                xml_declaration=True,
            )
        },
    )

    with pytest.raises(
        OfficePackageError,
        match="unsupported sheet relationship type",
    ):
        validate_xlsx(malformed)


def test_xlsx_selector_preserves_significant_sheet_name_whitespace() -> None:
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active["A1"] = "plain"
    spaced = workbook.create_sheet(" Data ")
    spaced["A1"] = "spaced"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    result, reports = edit_xlsx(
        output.getvalue(),
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(sheet_name=" Data ", ranges=["A1"]),
                formatting=XlsxCellFormatting(bold=True),
            )
        ],
    )

    assert reports[0]["sheet_name"] == " Data "
    edited = load_workbook(io.BytesIO(result))
    try:
        assert edited["Data"]["A1"].font.bold is not True
        assert edited[" Data "]["A1"].font.bold is True
    finally:
        edited.close()


def test_xlsx_inspects_and_preserves_multi_cell_array_formula_cache() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    for row, value in enumerate((2, 3, 4), start=1):
        worksheet.cell(row=row, column=1, value=value)
    worksheet["B1"] = ArrayFormula(ref="B1:B3", text="=A1:A3*2")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
        sheet = etree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    for row, cached_value in enumerate((4, 6, 8), start=1):
        row_node = sheet.xpath(f".//x:row[@r='{row}']", namespaces={"x": _SHEET_NS})[0]
        matches = row_node.xpath(f"./x:c[@r='B{row}']", namespaces={"x": _SHEET_NS})
        cell = matches[0] if matches else etree.SubElement(row_node, f"{{{_SHEET_NS}}}c", r=f"B{row}")
        value = cell.find(f"{{{_SHEET_NS}}}v")
        if value is None:
            value = etree.SubElement(cell, f"{{{_SHEET_NS}}}v")
        value.text = str(cached_value)
    data = _replace_parts(
        output.getvalue(),
        {
            "xl/worksheets/sheet1.xml": etree.tostring(
                sheet,
                encoding="UTF-8",
                xml_declaration=True,
            )
        },
    )

    inspected = inspect_xlsx(
        data,
        sheet_name="Sheet",
        cell_range="B1:B3",
    )

    assert inspected["cells"][0]["formula"] == "=A1:A3*2"
    assert inspected["cells"][0]["formula_type"] == "array"
    assert inspected["cells"][0]["formula_range"] == "B1:B3"
    assert inspected["cells"][0]["cached_value"] == 4
    assert "object at" not in str(inspected["cells"][0])

    edited, reports = edit_xlsx(
        data,
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(sheet_name="Sheet", ranges=["B1:B3"]),
                formatting=XlsxCellFormatting(bold=True),
            )
        ],
    )

    assert reports[0]["match_count"] == 3
    for coordinate in ("B1", "B2", "B3"):
        before = _cell_node(data, coordinate)
        after = _cell_node(edited, coordinate)
        before_content = [(etree.QName(child).localname, child.text, dict(child.attrib)) for child in before]
        after_content = [(etree.QName(child).localname, child.text, dict(child.attrib)) for child in after]
        assert after_content == before_content


def test_xlsx_formatting_preserves_formula_cache_structure_and_unrelated_parts() -> None:
    data = _workbook()
    formula_before = _cell_node(data, "B3")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        hidden_before = archive.read("xl/worksheets/sheet2.xml")
        core_before = archive.read("docProps/core.xml")

    result, reports = edit_xlsx(
        data,
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(
                    sheet_name="Data",
                    ranges=["B2:B3"],
                ),
                formatting=XlsxCellFormatting(
                    bold=True,
                    font_color="#008844",
                    fill_color="#FFF2CC",
                    horizontal_alignment="center",
                    vertical_alignment="center",
                    wrap_text=True,
                    number_format="$#,##0.00",
                ),
            )
        ],
    )

    assert reports[0]["match_count"] == 2
    assert reports[0]["sample_coordinates"] == ["B2", "B3"]
    assert validate_xlsx(result)["valid"] is True
    formula_after = _cell_node(result, "B3")
    before_formula = formula_before.find(f"{{{_SHEET_NS}}}f")
    after_formula = formula_after.find(f"{{{_SHEET_NS}}}f")
    before_value = formula_before.find(f"{{{_SHEET_NS}}}v")
    after_value = formula_after.find(f"{{{_SHEET_NS}}}v")
    assert before_formula is not None and after_formula is not None
    assert before_value is not None and after_value is not None
    assert (after_formula.text, after_value.text) == (before_formula.text, before_value.text)
    with zipfile.ZipFile(io.BytesIO(result)) as archive:
        assert archive.read("xl/worksheets/sheet2.xml") == hidden_before
        assert archive.read("docProps/core.xml") == core_before

    styled = load_workbook(io.BytesIO(result), data_only=False)
    cached = load_workbook(io.BytesIO(result), data_only=True)
    try:
        cell = styled["Data"]["B2"]
        assert cell.font.bold is True
        assert cell.font.color.rgb == "FF008844"
        assert cell.fill.fgColor.rgb == "FFFFF2CC"
        assert cell.alignment.horizontal == "center"
        assert cell.alignment.wrap_text is True
        assert cell.number_format == "$#,##0.00"
        assert styled["Data"]["B3"].value == "=B2*2"
        assert cached["Data"]["B3"].value == 2469
    finally:
        styled.close()
        cached.close()


def test_xlsx_edit_receipt_uses_exact_cell_paths_and_style_deltas() -> None:
    source = _workbook()
    operation = XlsxCellFormatOperation(
        cells=XlsxCellSelector(sheet_name="Data", ranges=["B2"]),
        formatting=XlsxCellFormatting(bold=True),
    )

    edited, reports, receipt = office_engine.edit_with_receipt(
        source,
        suffix=".xlsx",
        operations=[operation],
    )

    path = '/workbook/sheet[@name="Data"]/cell[B2]'
    operation_id = reports[0]["operation_id"]
    assert receipt["source"]["sha256"] == hashlib.sha256(source).hexdigest()
    assert receipt["result"]["sha256"] == hashlib.sha256(edited).hexdigest()
    assert receipt["applied_operation_ids"] == [operation_id]
    assert receipt["operations"][0]["target_paths"] == [path]
    assert receipt["semantic_changes"]["changed_target_paths"] == [path]
    assert {
        (
            delta["path"],
            delta["property"],
            delta["before"],
            delta["after"],
        )
        for delta in receipt["semantic_changes"]["semantic_deltas"]
    } == {(path, "formatting.font.bold", False, True)}
    assert {change["part_name"] for change in receipt["package_changes"]["parts"]} == {"xl/styles.xml", "xl/worksheets/sheet1.xml"}
    assert receipt["package_changes"]["relationship_change_count"] == 0


def test_xlsx_formatting_merges_base_style_and_deduplicates_new_cell_xf() -> None:
    data = _workbook()
    before_count = _cell_xf_count(data)

    result, reports = edit_xlsx(
        data,
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(sheet_name="Data", ranges=["A6:A8"]),
                formatting=XlsxCellFormatting(bold=True, font_color="#C00000"),
            )
        ],
    )

    assert reports[0]["match_count"] == 3
    assert _cell_xf_count(result) == before_count + 1
    workbook = load_workbook(io.BytesIO(result))
    try:
        for coordinate in ("A6", "A7", "A8"):
            cell = workbook["Data"][coordinate]
            assert cell.font.bold is True
            assert cell.font.italic is True
            assert cell.border.bottom.style == "thin"
            assert cell.border.bottom.color.rgb == "FF112233"
    finally:
        workbook.close()


def test_xlsx_selector_filters_are_conjunctive_and_transactional() -> None:
    data = _workbook()
    operation = XlsxCellFormatOperation(
        cells=XlsxCellSelector(
            sheet_name="Data",
            ranges=["A1:B5"],
            contains_text="missing",
            has_formula=True,
        ),
        formatting=XlsxCellFormatting(bold=True),
    )

    with pytest.raises(OfficeOperationError, match="no document changes were written"):
        edit_xlsx(data, [operation])

    assert _cell_node(data, "B3").get("s") == _cell_node(_workbook(), "B3").get("s")


def test_xlsx_no_match_returns_the_exact_source_package_when_allowed() -> None:
    data = _workbook()

    edited, reports = edit_xlsx(
        data,
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(
                    sheet_name="Data",
                    ranges=["A1:B5"],
                    contains_text="missing",
                ),
                formatting=XlsxCellFormatting(bold=True),
                require_match=False,
            )
        ],
    )

    assert edited == data
    assert reports[0]["match_count"] == 0


def test_xlsx_rgb_selectors_match_six_digit_openpyxl_colors() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    for row, color in enumerate(("00FF0000", "80FF0000", "FFFF0000"), start=1):
        worksheet.cell(row=row, column=1, value=f"font-{row}")
        worksheet.cell(row=row, column=1).font = Font(color=color)
    worksheet["A4"] = "fill"
    worksheet["A4"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    result, reports = edit_xlsx(
        output.getvalue(),
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(
                    sheet_name="Sheet",
                    ranges=["A1:A3"],
                    font_color="#FF0000",
                ),
                formatting=XlsxCellFormatting(bold=True),
            ),
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(
                    sheet_name="Sheet",
                    ranges=["A4"],
                    fill_color="#FFFF00",
                ),
                formatting=XlsxCellFormatting(italic=True),
            ),
        ],
    )

    assert [report["match_count"] for report in reports] == [3, 1]
    edited = load_workbook(io.BytesIO(result))
    try:
        assert all(edited.active.cell(row=row, column=1).font.bold is True for row in range(1, 4))
        assert edited.active["A4"].font.italic is True
    finally:
        edited.close()


@pytest.mark.parametrize(
    "part_name",
    [
        "xl/externalLinks/externalLink1.xml",
        "xl/ExternalLinks/ExternalLink1.xml",
    ],
)
def test_xlsx_edit_rejects_external_workbook_links_before_mutation(part_name: str) -> None:
    data = _replace_parts(
        _workbook(),
        {},
        additions={part_name: (b'<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>')},
    )
    assert validate_xlsx(data)["risky_features"] == ["external workbook links"]

    with pytest.raises(OfficeOperationError, match="external workbook links"):
        edit_xlsx(
            data,
            [
                XlsxCellFormatOperation(
                    cells=XlsxCellSelector(sheet_name="Data", ranges=["B2"]),
                    formatting=XlsxCellFormatting(bold=True),
                )
            ],
        )


def test_xlsx_detects_relationship_addressed_external_links() -> None:
    data = _workbook()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationship_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    etree.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        Id="rIdExternalAudit",
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink",
        Target="custom/link.xml",
    )
    risky = _replace_parts(
        data,
        {
            "xl/_rels/workbook.xml.rels": etree.tostring(
                relationships,
                encoding="UTF-8",
                xml_declaration=True,
            )
        },
        additions={"xl/custom/link.xml": b"<externalLink/>"},
    )

    assert validate_xlsx(risky)["risky_features"] == ["external workbook links"]
    with pytest.raises(OfficeOperationError, match="external workbook links"):
        xlsx_module.validate_xlsx_renderable(risky)


def test_xlsx_edit_preserves_unknown_worksheet_extension_markup() -> None:
    data = _workbook()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        sheet_payload = archive.read("xl/worksheets/sheet1.xml")
    root = etree.fromstring(sheet_payload)
    extension_list = etree.SubElement(root, f"{{{_SHEET_NS}}}extLst")
    extension = etree.SubElement(extension_list, f"{{{_SHEET_NS}}}ext", uri="urn:vassilflow:test-extension")
    etree.SubElement(extension, "{urn:vassilflow:test-extension}payload", marker="keep")
    extended = _replace_parts(
        data,
        {"xl/worksheets/sheet1.xml": etree.tostring(root, encoding="UTF-8", xml_declaration=True)},
    )

    result, _ = edit_xlsx(
        extended,
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(sheet_name="Data", ranges=["A1"]),
                formatting=XlsxCellFormatting(bold=True),
            )
        ],
    )

    before = _cell_node(extended, "A1").getroottree().getroot()
    after = _cell_node(result, "A1").getroottree().getroot()
    before_extension = before.xpath("./x:extLst", namespaces={"x": _SHEET_NS})[0]
    after_extension = after.xpath("./x:extLst", namespaces={"x": _SHEET_NS})[0]
    assert etree.tostring(before_extension, method="c14n") == etree.tostring(
        after_extension,
        method="c14n",
    )


def test_xlsx_inspection_supports_shared_strings_and_1904_date_epoch() -> None:
    workbook = Workbook()
    workbook.epoch = CALENDAR_MAC_1904
    workbook.active["A1"] = datetime(2026, 7, 14)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    date_inspection = inspect_xlsx(output.getvalue(), sheet_name="Sheet", cell_range="A1")
    assert date_inspection["cells"][0]["value"] == "2026-07-14T00:00:00"

    data = _workbook()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        content_types = etree.fromstring(archive.read("[Content_Types].xml"))
        relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        sheet = etree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    content_type_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    relationship_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    etree.SubElement(
        content_types,
        f"{{{content_type_ns}}}Override",
        PartName="/xl/sharedStrings.xml",
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    )
    etree.SubElement(
        relationships,
        f"{{{relationship_ns}}}Relationship",
        Id="rIdSharedStrings",
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
        Target="sharedStrings.xml",
    )
    a1 = sheet.xpath(".//x:c[@r='A1']", namespaces={"x": _SHEET_NS})[0]
    a1.set("t", "s")
    for child in list(a1):
        a1.remove(child)
    etree.SubElement(a1, f"{{{_SHEET_NS}}}v").text = "0"
    shared_strings = b'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1"><si><t>Metric</t></si></sst>'
    shared = _replace_parts(
        data,
        {
            "[Content_Types].xml": etree.tostring(content_types, encoding="UTF-8", xml_declaration=True),
            "xl/_rels/workbook.xml.rels": etree.tostring(
                relationships,
                encoding="UTF-8",
                xml_declaration=True,
            ),
            "xl/worksheets/sheet1.xml": etree.tostring(sheet, encoding="UTF-8", xml_declaration=True),
        },
        additions={"xl/sharedStrings.xml": shared_strings},
    )

    inspected = inspect_xlsx(shared, sheet_name="Data", cell_range="A1")
    assert inspected["cells"][0]["value"] == "Metric"
    with zipfile.ZipFile(io.BytesIO(shared)) as archive:
        shared_before = archive.read("xl/sharedStrings.xml")
    edited, _ = edit_xlsx(
        shared,
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(sheet_name="Data", ranges=["A1"]),
                formatting=XlsxCellFormatting(italic=True),
            )
        ],
    )
    with zipfile.ZipFile(io.BytesIO(edited)) as archive:
        assert archive.read("xl/sharedStrings.xml") == shared_before


@pytest.mark.parametrize(
    "cell_range",
    ["XFE1", "A1048577", "B2:A1", "Sheet1!A1", "A0"],
)
def test_xlsx_selector_rejects_invalid_or_out_of_bounds_ranges(cell_range: str) -> None:
    with pytest.raises(ValueError):
        XlsxCellSelector(sheet_name="Data", ranges=[cell_range])


def test_xlsx_selector_accepts_last_excel_cell_and_rejects_large_area() -> None:
    assert XlsxCellSelector(sheet_name="Data", ranges=["XFD1048576"]).ranges == ["XFD1048576"]
    assert XlsxCellSelector(
        sheet_name="Data",
        ranges=["A1:J1000", "A1:A1000"],
    ).ranges == ["A1:J1000", "A1:A1000"]
    with pytest.raises(ValueError, match="10,000"):
        XlsxCellSelector(sheet_name="Data", ranges=["A1:XFD1048576"])


def test_xlsx_validation_rejects_macro_or_wrong_workbook_content_type() -> None:
    data = _workbook()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        content_types = archive.read("[Content_Types].xml")
    invalid = content_types.replace(
        b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    )

    with pytest.raises(OfficePackageError, match="macro-free"):
        validate_xlsx(_replace_parts(data, {"[Content_Types].xml": invalid}))


@pytest.mark.parametrize(
    "relationship_type, content_type",
    [
        (
            "http://schemas.microsoft.com/office/2006/relationships/xlMacrosheet",
            "application/vnd.ms-excel.macrosheet+xml",
        ),
        (
            "http://schemas.microsoft.com/office/2006/relationships/xlIntlMacrosheet",
            "application/vnd.ms-excel.intlmacrosheet+xml",
        ),
    ],
)
def test_xlsx_rejects_excel_macro_sheets_before_inspect_or_render(
    relationship_type: str,
    content_type: str,
) -> None:
    data = _convert_first_sheet_to_macro_sheet(
        _workbook(),
        relationship_type=relationship_type,
        content_type=content_type,
    )

    with pytest.raises(OfficePackageError, match="macro content"):
        inspect_xlsx(data)
    with pytest.raises(OfficePackageError, match="macro content"):
        office_engine.render(data, suffix=".xlsx")


@pytest.mark.parametrize(
    "number_format, message",
    [
        ('"unterminated', "balanced double quotes"),
        ("0;0;0;0;0", "at most four"),
        ("[Red", "balanced square brackets"),
    ],
)
def test_xlsx_formatting_rejects_malformed_number_formats(
    number_format: str,
    message: str,
) -> None:
    with pytest.raises(ValueError, match=message):
        XlsxCellFormatting(number_format=number_format)


def test_xlsx_validation_rejects_malformed_existing_number_format() -> None:
    result, _ = edit_xlsx(
        _workbook(),
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(sheet_name="Data", ranges=["B2"]),
                formatting=XlsxCellFormatting(number_format='"USD" 0.00'),
            )
        ],
    )
    with zipfile.ZipFile(io.BytesIO(result)) as archive:
        styles = archive.read("xl/styles.xml").replace(
            b'formatCode="&quot;USD&quot; 0.00"',
            b'formatCode="&quot;unterminated"',
        )
    malformed = _replace_parts(result, {"xl/styles.xml": styles})

    with pytest.raises(OfficePackageError, match="balanced double quotes"):
        validate_xlsx(malformed)


def test_xlsx_validation_rejects_existing_number_format_over_255_characters() -> None:
    result, _ = edit_xlsx(
        _workbook(),
        [
            XlsxCellFormatOperation(
                cells=XlsxCellSelector(sheet_name="Data", ranges=["B2"]),
                formatting=XlsxCellFormatting(number_format='"USD" 0.00'),
            )
        ],
    )
    with zipfile.ZipFile(io.BytesIO(result)) as archive:
        styles = etree.fromstring(archive.read("xl/styles.xml"))
    custom_format = styles.xpath(".//x:numFmt", namespaces={"x": _SHEET_NS})[-1]
    custom_format.set("formatCode", "0" * 256)
    malformed = _replace_parts(
        result,
        {
            "xl/styles.xml": etree.tostring(
                styles,
                encoding="UTF-8",
                xml_declaration=True,
            )
        },
    )

    with pytest.raises(OfficePackageError, match="255 characters"):
        validate_xlsx(malformed)

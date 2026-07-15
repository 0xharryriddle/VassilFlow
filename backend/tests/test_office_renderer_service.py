from __future__ import annotations

import asyncio
import importlib.util
import io
import json
import sys
import warnings
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

REPO_ROOT = Path(__file__).resolve().parents[2]
MODULE_PATH = REPO_ROOT / "docker" / "office-renderer" / "app.py"
_CONTENT_TYPES = '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
_DEFAULT_MAIN_CONTENT_TYPES = '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
_ROOT_RELATIONSHIPS = (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="word/document.xml"/></Relationships>'
)


def _load_renderer_module():
    spec = importlib.util.spec_from_file_location("vassilflow_office_renderer_test", MODULE_PATH)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _docx() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("_rels/.rels", _ROOT_RELATIONSHIPS)
        archive.writestr("word/document.xml", "<document/>")
    return output.getvalue()


def _xlsx() -> bytes:
    workbook = Workbook()
    workbook.active["A1"] = "hello"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pptx(
    *,
    external_relationship_type: str | None = None,
    external_target: str = "https://example.test",
    extra_parts: dict[str, str] | None = None,
    slide_action: str | None = None,
) -> bytes:
    content_types = """<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
      <Default Extension="xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
      <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
      <Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>
    </Types>"""
    root_relationships = """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
    </Relationships>"""
    presentation = """<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst><p:sldSz cx="12192000" cy="6858000"/>
    </p:presentation>"""
    presentation_relationships = """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
    </Relationships>"""
    action = f'<a:hlinkClick r:id="" action="{slide_action}"/>' if slide_action is not None else ""
    slide = f"""<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:cSld><p:spTree>
        <p:nvGrpSpPr/><p:grpSpPr/>
        <p:sp>
          <p:nvSpPr><p:cNvPr id="1" name="Baseline">{action}</p:cNvPr></p:nvSpPr>
          <p:spPr/>
          <p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>Baseline</a:t></a:r></a:p></p:txBody>
        </p:sp>
      </p:spTree></p:cSld>
    </p:sld>"""
    slide_relationships = f"""<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rIdExternal" Type="{external_relationship_type}" Target="{external_target}" TargetMode="External"/>
    </Relationships>"""

    output = io.BytesIO()
    with zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_relationships)
        archive.writestr("ppt/presentation.xml", presentation)
        archive.writestr("ppt/_rels/presentation.xml.rels", presentation_relationships)
        archive.writestr("ppt/slides/slide1.xml", slide)
        if external_relationship_type is not None:
            archive.writestr("ppt/slides/_rels/slide1.xml.rels", slide_relationships)
        for part_name, payload in (extra_parts or {}).items():
            archive.writestr(part_name, payload)
    return output.getvalue()


def _relocate_first_xlsx_worksheet(data: bytes) -> bytes:
    source_part = "xl/worksheets/sheet1.xml"
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as destination,
    ):
        worksheet = source.read(source_part)
        for info in source.infolist():
            if info.filename == source_part:
                continue
            payload = source.read(info)
            if info.filename == "xl/_rels/workbook.xml.rels":
                payload = payload.replace(
                    b"worksheets/sheet1.xml",
                    b"custom/sheet.xml",
                )
            elif info.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"/xl/worksheets/sheet1.xml",
                    b"/xl/custom/sheet.xml",
                )
            destination.writestr(info, payload)
        destination.writestr("xl/custom/sheet.xml", worksheet)
    return output.getvalue()


def _xlsx_with_external_link_relationship(data: bytes) -> bytes:
    output = io.BytesIO()
    marker = b"</Relationships>"
    relationship = b'<Relationship Id="rIdExternalAudit" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" Target="custom/link.xml"/>'
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as destination,
    ):
        for info in source.infolist():
            payload = source.read(info)
            if info.filename == "xl/_rels/workbook.xml.rels":
                assert marker in payload
                payload = payload.replace(marker, relationship + marker)
            destination.writestr(info, payload)
        destination.writestr("xl/custom/link.xml", b"<externalLink/>")
    return output.getvalue()


def _xlsx_with_macro_sheet(
    data: bytes,
    *,
    relationship_type: str,
    content_type: str,
) -> bytes:
    output = io.BytesIO()
    worksheet_relationship = b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
    worksheet_content_type = b"application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as destination,
    ):
        for info in source.infolist():
            payload = source.read(info)
            if info.filename == "xl/_rels/workbook.xml.rels":
                assert worksheet_relationship in payload
                payload = payload.replace(
                    worksheet_relationship,
                    relationship_type.encode(),
                    1,
                )
            elif info.filename == "[Content_Types].xml":
                assert worksheet_content_type in payload
                payload = payload.replace(
                    worksheet_content_type,
                    content_type.encode(),
                    1,
                )
            destination.writestr(info, payload)
    return output.getvalue()


def _two_page_pdf(path: Path) -> None:
    first = Image.new("RGB", (72, 72), "white")
    second = Image.new("RGB", (72, 72), "black")
    try:
        first.save(path, format="PDF", save_all=True, append_images=[second], resolution=72)
    finally:
        first.close()
        second.close()


def test_renderer_service_rasterizes_bounded_pdf_page_window(tmp_path: Path) -> None:
    module = _load_renderer_module()
    pdf_path = tmp_path / "document.pdf"
    _two_page_pdf(pdf_path)

    page_count, pages = module._render_pages(
        pdf_path,
        start_page=2,
        max_pages=1,
        dpi=96,
    )

    assert page_count == 2
    assert len(pages) == 1
    assert pages[0].page == 2
    assert pages[0].filename == "page-002.png"
    assert pages[0].data.startswith(b"\x89PNG\r\n\x1a\n")


def test_renderer_service_rejects_oversized_page_before_bitmap_allocation(monkeypatch, tmp_path: Path) -> None:
    module = _load_renderer_module()

    class FakePage:
        def get_size(self):
            return 10_000, 10_000

        def render(self, **_kwargs):
            raise AssertionError("bitmap allocation must not run")

        def close(self):
            return None

    class FakePdf:
        def __len__(self):
            return 1

        def __getitem__(self, _index):
            return FakePage()

        def close(self):
            return None

    monkeypatch.setattr(module.pdfium, "PdfDocument", lambda _path: FakePdf())

    try:
        module._render_pages(tmp_path / "unused.pdf", start_page=1, max_pages=1, dpi=120)
    except ValueError as exc:
        assert "raster dimensions" in str(exc)
    else:
        raise AssertionError("oversized page must be rejected before rendering")


def test_renderer_service_kills_raster_worker_at_hard_timeout(monkeypatch, tmp_path: Path) -> None:
    module = _load_renderer_module()

    def timeout(*_args, **_kwargs):
        raise module.subprocess.TimeoutExpired(cmd="raster-worker", timeout=45)

    monkeypatch.setattr(module.subprocess, "run", timeout)

    try:
        module._render_pages_isolated(
            tmp_path / "document.pdf",
            tmp_path,
            start_page=1,
            max_pages=1,
            dpi=120,
        )
    except TimeoutError as exc:
        assert "45 second" in str(exc)
    else:
        raise AssertionError("raster worker timeout must be enforced by the parent process")


def test_renderer_service_endpoint_returns_archive(monkeypatch) -> None:
    module = _load_renderer_module()
    expected = b"PK\x03\x04render"
    captured: dict[str, object] = {}

    def fake_render(document: bytes, *, start_page: int, max_pages: int, dpi: int) -> bytes:
        captured.update(
            document=document,
            start_page=start_page,
            max_pages=max_pages,
            dpi=dpi,
        )
        return expected

    monkeypatch.setattr(module, "_render_docx", fake_render)
    client = TestClient(module.app)
    document = _docx()

    response = client.post(
        "/v1/render/docx?start_page=2&max_pages=3&dpi=144",
        content=document,
        headers={"Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert response.content == expected
    assert captured == {
        "document": document,
        "start_page": 2,
        "max_pages": 3,
        "dpi": 144,
    }


def test_renderer_service_rejects_non_docx_media_type() -> None:
    module = _load_renderer_module()
    client = TestClient(module.app)

    response = client.post(
        "/v1/render/docx",
        content=b"not a docx",
        headers={"Content-Type": "text/plain"},
    )

    assert response.status_code == 415


def test_renderer_service_xlsx_endpoint_returns_archive(monkeypatch) -> None:
    module = _load_renderer_module()
    expected = b"PK\x03\x04xlsx-render"
    captured: dict[str, object] = {}

    def fake_render(document: bytes, *, start_page: int, max_pages: int, dpi: int) -> bytes:
        captured.update(
            document=document,
            start_page=start_page,
            max_pages=max_pages,
            dpi=dpi,
        )
        return expected

    monkeypatch.setattr(module, "_render_xlsx", fake_render)
    client = TestClient(module.app)
    document = _xlsx()

    response = client.post(
        "/v1/render/xlsx?start_page=1&max_pages=2&dpi=120",
        content=document,
        headers={"Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    )

    assert response.status_code == 200
    assert response.content == expected
    assert captured == {
        "document": document,
        "start_page": 1,
        "max_pages": 2,
        "dpi": 120,
    }


def test_renderer_service_pptx_endpoint_returns_archive(monkeypatch) -> None:
    module = _load_renderer_module()
    expected = b"PK\x03\x04pptx-render"
    captured: dict[str, object] = {}

    def fake_render(document: bytes, *, start_page: int, max_pages: int, dpi: int) -> bytes:
        captured.update(
            document=document,
            start_page=start_page,
            max_pages=max_pages,
            dpi=dpi,
        )
        return expected

    monkeypatch.setattr(module, "_render_pptx", fake_render)
    client = TestClient(module.app)
    document = _pptx()

    response = client.post(
        "/v1/render/pptx?start_page=1&max_pages=2&dpi=120",
        content=document,
        headers={"Content-Type": "application/vnd.openxmlformats-officedocument.presentationml.presentation"},
    )

    assert response.status_code == 200
    assert response.content == expected
    assert captured == {
        "document": document,
        "start_page": 1,
        "max_pages": 2,
        "dpi": 120,
    }


def test_renderer_service_validates_pptx_and_allows_safe_external_hyperlinks() -> None:
    module = _load_renderer_module()
    hyperlink_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"

    assert module._validate_pptx(_pptx()) == 1
    assert module._validate_pptx(_pptx(external_relationship_type=hyperlink_type)) == 1


def test_renderer_service_rejects_external_pptx_linked_resources() -> None:
    module = _load_renderer_module()
    image_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

    with pytest.raises(ValueError, match="external relationships"):
        module._validate_pptx(
            _pptx(
                external_relationship_type=image_type,
                external_target="https://example.test/image.png",
            )
        )


def test_renderer_service_rejects_unsafe_external_pptx_hyperlinks() -> None:
    module = _load_renderer_module()
    hyperlink_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"

    with pytest.raises(ValueError, match="external relationships"):
        module._validate_pptx(
            _pptx(
                external_relationship_type=hyperlink_type,
                external_target="file:///tmp/local.png",
            )
        )


def test_renderer_service_rejects_unsafe_pptx_slide_actions() -> None:
    module = _load_renderer_module()

    with pytest.raises(ValueError, match="unsafe actions"):
        module._validate_pptx(_pptx(slide_action="javascript:alert(1)"))


def test_renderer_service_validates_every_internal_pptx_relationship_target() -> None:
    module = _load_renderer_module()
    chart_relationships = """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rIdData" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" Target="../embeddings/missing.xlsx"/>
    </Relationships>"""

    with pytest.raises(ValueError, match="targets a missing part"):
        module._validate_pptx(
            _pptx(
                extra_parts={
                    "ppt/charts/chart1.xml": "<chart/>",
                    "ppt/charts/_rels/chart1.xml.rels": chart_relationships,
                }
            )
        )


def test_renderer_service_uses_impress_pdf_export_for_pptx(monkeypatch, tmp_path: Path) -> None:
    module = _load_renderer_module()
    captured: dict[str, object] = {}

    def fake_run(command, **kwargs):
        captured["command"] = command
        captured["env"] = kwargs["env"]
        output_dir = Path(command[command.index("--outdir") + 1])
        (output_dir / "document.pdf").write_bytes(b"%PDF-baseline")
        return SimpleNamespace(returncode=0, stdout=b"", stderr=b"")

    monkeypatch.setattr(module, "_SOFFICE", "/usr/bin/soffice")
    monkeypatch.setattr(module.subprocess, "run", fake_run)
    job_dir = tmp_path / "job"
    job_dir.mkdir()

    output = module._convert_to_pdf(_pptx(), job_dir, format_name="pptx")

    assert output.read_bytes() == b"%PDF-baseline"
    assert module._CONVERSION_FILTERS["pptx"] in captured["command"]
    assert "ExportHiddenSlides" in module._CONVERSION_FILTERS["pptx"]


def test_renderer_service_pptx_archive_maps_pages_to_source_slides() -> None:
    module = _load_renderer_module()
    page = module._RenderedPage(
        page=2,
        filename="page-002.png",
        width=1,
        height=1,
        sha256="a" * 64,
        data=b"png",
    )

    payload = module._build_archive(
        format_name="pptx",
        page_count=3,
        start_page=2,
        max_pages=1,
        dpi=120,
        pages=(page,),
        source_sha256="b" * 64,
    )

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        manifest = json.loads(archive.read("manifest.json"))
    assert manifest["pipeline_contract_version"] == module._PIPELINE_CONTRACT_VERSION
    assert manifest["pages"][0]["source_slide"] == 2


def test_renderer_service_rejects_pptx_slide_to_page_count_drift(
    monkeypatch,
    tmp_path: Path,
) -> None:
    module = _load_renderer_module()
    page = module._RenderedPage(
        page=1,
        filename="page-001.png",
        width=1,
        height=1,
        sha256="a" * 64,
        data=b"png",
    )
    monkeypatch.setattr(module, "_validate_pptx", lambda _document: 2)
    monkeypatch.setattr(
        module,
        "_convert_to_pdf",
        lambda *_args, **_kwargs: tmp_path / "document.pdf",
    )
    monkeypatch.setattr(
        module,
        "_render_pages_isolated",
        lambda *_args, **_kwargs: (1, (page,)),
    )

    with pytest.raises(ValueError, match="slide-to-page mapping"):
        module._render_office(
            b"pptx",
            format_name="pptx",
            start_page=1,
            max_pages=1,
            dpi=120,
        )


def test_renderer_service_validates_xlsx_opc_controls() -> None:
    module = _load_renderer_module()

    module._validate_xlsx(_xlsx())


def test_renderer_service_bounds_declared_xlsx_cells(monkeypatch) -> None:
    module = _load_renderer_module()
    workbook = Workbook()
    workbook.active["A1"] = "one"
    workbook.active["A2"] = "two"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    document = _relocate_first_xlsx_worksheet(output.getvalue())
    monkeypatch.setattr(module, "_MAX_DECLARED_CELLS_PER_SHEET", 1)

    try:
        module._validate_xlsx(document)
    except ValueError as exc:
        assert "declares more than 1 cells" in str(exc)
    else:
        raise AssertionError("oversized XLSX worksheets must be rejected before conversion")


def test_renderer_service_rejects_relationship_addressed_external_links() -> None:
    module = _load_renderer_module()

    try:
        module._validate_xlsx(_xlsx_with_external_link_relationship(_xlsx()))
    except ValueError as exc:
        assert "external workbook links" in str(exc)
    else:
        raise AssertionError("external workbook links must be rejected before conversion")


@pytest.mark.parametrize(
    "relationship_type, content_type",
    [
        (
            "http://schemas.microsoft.com/office/2006/relationships/xlMacrosheet",
            "application/vnd.ms-excel.macrosheet+xml",
        ),
        (
            "http://schemas.microsoft.com/office/2006/relationships/xlIntlMacrosheet",
            "application/vnd.ms-excel.intlmacrosheet+xml",
        ),
    ],
)
def test_renderer_service_rejects_excel_macro_sheets(
    relationship_type: str,
    content_type: str,
) -> None:
    module = _load_renderer_module()
    document = _xlsx_with_macro_sheet(
        _xlsx(),
        relationship_type=relationship_type,
        content_type=content_type,
    )

    try:
        module._validate_xlsx(document)
    except ValueError as exc:
        assert "macros" in str(exc)
    else:
        raise AssertionError("Excel macro sheets must be rejected before conversion")


def test_renderer_service_rejects_malformed_xlsx_number_format() -> None:
    workbook = Workbook()
    workbook.active["A1"] = 42
    workbook.active["A1"].number_format = '"USD" 0.00'
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    malformed = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(output.getvalue())) as source,
        zipfile.ZipFile(malformed, mode="w", compression=zipfile.ZIP_DEFLATED) as destination,
    ):
        for info in source.infolist():
            payload = source.read(info)
            if info.filename == "xl/styles.xml":
                payload = payload.replace(
                    b'formatCode="&quot;USD&quot; 0.00"',
                    b'formatCode="&quot;unterminated"',
                )
            destination.writestr(info, payload)

    module = _load_renderer_module()
    try:
        module._validate_xlsx(malformed.getvalue())
    except ValueError as exc:
        assert "unbalanced double quotes" in str(exc)
    else:
        raise AssertionError("malformed XLSX number formats must be rejected before conversion")


def test_renderer_service_rejects_xlsx_number_format_over_255_characters() -> None:
    workbook = Workbook()
    workbook.active["A1"] = 42
    workbook.active["A1"].number_format = '"USD" 0.00'
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    oversized = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(output.getvalue())) as source,
        zipfile.ZipFile(oversized, mode="w", compression=zipfile.ZIP_DEFLATED) as destination,
    ):
        for info in source.infolist():
            payload = source.read(info)
            if info.filename == "xl/styles.xml":
                payload = payload.replace(
                    b'formatCode="&quot;USD&quot; 0.00"',
                    f'formatCode="{"0" * 256}"'.encode(),
                )
            destination.writestr(info, payload)

    module = _load_renderer_module()
    try:
        module._validate_xlsx(oversized.getvalue())
    except ValueError as exc:
        assert "number format exceeds" in str(exc)
    else:
        raise AssertionError("oversized XLSX number formats must be rejected before conversion")


def test_renderer_service_rejects_busy_request_before_reading_body(monkeypatch) -> None:
    module = _load_renderer_module()
    module._RENDER_SLOTS = asyncio.Semaphore(0)
    module._QUEUE_TIMEOUT_SECONDS = 0.01

    async def fail_if_read(_request):
        raise AssertionError("body must not be read before admission")

    monkeypatch.setattr(module, "_read_request_body", fail_if_read)
    client = TestClient(module.app)

    response = client.post(
        "/v1/render/docx",
        content=_docx(),
        headers={"Content-Type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    )

    assert response.status_code == 429
    assert response.headers["retry-after"] == "5"


def test_renderer_service_rejects_duplicate_package_entries() -> None:
    module = _load_renderer_module()
    output = io.BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        with zipfile.ZipFile(output, mode="w") as archive:
            archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
            archive.writestr("_rels/.rels", _ROOT_RELATIONSHIPS)
            archive.writestr("word/document.xml", "<document/>")
            archive.writestr("word/document.xml", "<document/>")

    try:
        module._validate_docx(output.getvalue())
    except ValueError as exc:
        assert "duplicate" in str(exc)
    else:
        raise AssertionError("duplicate DOCX package entries must be rejected")


def test_renderer_service_rejects_invalid_opc_controls() -> None:
    module = _load_renderer_module()
    output = io.BytesIO()
    with zipfile.ZipFile(output, mode="w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("_rels/.rels", _ROOT_RELATIONSHIPS)
        archive.writestr("word/document.xml", "<document/>")

    try:
        module._validate_docx(output.getvalue())
    except ValueError as exc:
        assert "Content_Types" in str(exc)
    else:
        raise AssertionError("invalid OPC controls must be rejected")


def test_renderer_service_accepts_main_content_type_from_default_extension() -> None:
    module = _load_renderer_module()
    output = io.BytesIO()
    with zipfile.ZipFile(output, mode="w") as archive:
        archive.writestr("[Content_Types].xml", _DEFAULT_MAIN_CONTENT_TYPES)
        archive.writestr("_rels/.rels", _ROOT_RELATIONSHIPS)
        archive.writestr("word/document.xml", "<document/>")

    module._validate_docx(output.getvalue())

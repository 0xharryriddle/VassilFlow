"""Build deterministic native PPTX geometry corpus artifacts."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import struct
import zlib
from datetime import datetime
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from vassilflow.community.office.engine import office_engine
from vassilflow.community.office.generation import _canonicalize_pptx_package
from vassilflow.community.office.pptx import validate_pptx

_EMU_PER_INCH = 914_400
_FIXED_TIMESTAMP = datetime(2026, 7, 17, 0, 0, 0)


def _png_chunk(chunk_type: bytes, payload: bytes) -> bytes:
    checksum = zlib.crc32(chunk_type + payload) & 0xFFFFFFFF
    return struct.pack(">I", len(payload)) + chunk_type + payload + struct.pack(">I", checksum)


def build_geometry_image() -> bytes:
    """Return the deterministic 200x100 visual embedded by every corpus picture."""

    width = 200
    height = 100
    scanlines = bytearray()
    for y in range(height):
        scanlines.append(0)
        for x in range(width):
            if x < 5 or x >= width - 5 or y < 5 or y >= height - 5:
                color = (23, 33, 43, 255)
            elif x < width // 2 and y < height // 2:
                color = (31, 111, 235, 255)
            elif x >= width // 2 and y < height // 2:
                color = (15, 118, 110, 255)
            elif x < width // 2:
                color = (255, 183, 77, 255)
            else:
                color = (220, 38, 38, 255)
            if abs((x * height) - (y * width)) < width * 2:
                color = (255, 255, 255, 255)
            scanlines.extend(color)
    header = struct.pack(">IIBBBBB", width, height, 8, 6, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + _png_chunk(b"IHDR", header) + _png_chunk(b"IDAT", zlib.compress(bytes(scanlines), level=9)) + _png_chunk(b"IEND", b"")


def _set_name(shape: Any, name: str) -> None:
    shape.name = name


def _set_picture_alt_text(picture: Any, alt_text: str) -> None:
    properties = picture._element.nvPicPr.cNvPr  # noqa: SLF001 - fixture OOXML only
    properties.set("descr", alt_text)
    properties.set("title", alt_text)


def _set_group_transform(
    group: Any,
    *,
    x: int,
    y: int,
    width: int,
    height: int,
    child_x: int,
    child_y: int,
    child_width: int,
    child_height: int,
    rotation_degrees: int = 0,
    flip_horizontal: bool = False,
    flip_vertical: bool = False,
) -> None:
    transform = group._element.grpSpPr.xfrm  # noqa: SLF001 - fixture OOXML only
    transform.off.x = int(x)
    transform.off.y = int(y)
    transform.ext.cx = int(width)
    transform.ext.cy = int(height)
    transform.chOff.x = int(child_x)
    transform.chOff.y = int(child_y)
    transform.chExt.cx = int(child_width)
    transform.chExt.cy = int(child_height)
    if rotation_degrees:
        transform.set("rot", str(rotation_degrees * 60_000))
    else:
        transform.attrib.pop("rot", None)
    if flip_horizontal:
        transform.set("flipH", "1")
    else:
        transform.attrib.pop("flipH", None)
    if flip_vertical:
        transform.set("flipV", "1")
    else:
        transform.attrib.pop("flipV", None)


def _style_title(title: Any, text: str, *, left: float = 0.55, width: float = 12.2) -> None:
    title.text = text
    title.left = Inches(left)
    title.top = Inches(0.25)
    title.width = Inches(width)
    title.height = Inches(0.6)
    _set_name(title, f"VFGEOM:TITLE:{text.upper().replace(' ', '_')}")
    paragraph = title.text_frame.paragraphs[0]
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(24)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(23, 33, 43)


def _add_label(slide: Any, text: str, *, left: float, top: float, width: float) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.55))
    _set_name(box, f"VFGEOM:LABEL:{text.upper().replace(' ', '_')}")
    box.text = text
    paragraph = box.text_frame.paragraphs[0]
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(14)
    paragraph.font.color.rgb = RGBColor(86, 97, 111)


def _style_slide(slide: Any) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(247, 248, 250)


def _add_picture(
    shapes: Any,
    image: bytes,
    *,
    name: str,
    alt_text: str,
    left: float,
    top: float,
    width: float,
    height: float,
) -> Any:
    picture = shapes.add_picture(
        io.BytesIO(image),
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    _set_name(picture, name)
    _set_picture_alt_text(picture, alt_text)
    return picture


def build_geometry_seed() -> tuple[bytes, bytes, dict[str, Any]]:
    """Build the seed, embedded asset, and source-bound preflight evidence."""

    image = build_geometry_image()
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    properties = presentation.core_properties
    properties.title = "PPTX Geometry Corpus"
    properties.subject = "Native grouped geometry and clipping evidence"
    properties.author = "VassilFlow"
    properties.last_modified_by = "VassilFlow"
    properties.comments = "vassilflow-pptx-geometry-corpus-v1"
    properties.created = _FIXED_TIMESTAMP
    properties.modified = _FIXED_TIMESTAMP
    properties.revision = 1
    title_layout = presentation.slide_layouts[5]

    nested_slide = presentation.slides.add_slide(title_layout)
    _style_slide(nested_slide)
    if nested_slide.shapes.title is None:
        raise RuntimeError("Geometry corpus title layout is missing its title placeholder")
    _style_title(nested_slide.shapes.title, "Nested group transform", left=8.25, width=4.5)
    _add_label(
        nested_slide,
        "8 x 4 inch picture frame; 12.5% clipped",
        left=8.3,
        top=1.05,
        width=4.3,
    )
    outer = nested_slide.shapes.add_group_shape()
    _set_name(outer, "VF_GEOM_NESTED_OUTER")
    inner = outer.shapes.add_group_shape()
    _set_name(inner, "VF_GEOM_NESTED_INNER")
    _add_picture(
        inner.shapes,
        image,
        name="VF_GEOM_NESTED_PICTURE",
        alt_text="Nested picture transformed by two groups and clipped at the top slide edge",
        left=0,
        top=0,
        width=2,
        height=1,
    )
    _set_group_transform(
        inner,
        x=0,
        y=0,
        width=4 * _EMU_PER_INCH,
        height=2 * _EMU_PER_INCH,
        child_x=0,
        child_y=0,
        child_width=2 * _EMU_PER_INCH,
        child_height=_EMU_PER_INCH,
        flip_vertical=True,
    )
    _set_group_transform(
        outer,
        x=2 * _EMU_PER_INCH,
        y=_EMU_PER_INCH,
        width=8 * _EMU_PER_INCH,
        height=4 * _EMU_PER_INCH,
        child_x=0,
        child_y=0,
        child_width=4 * _EMU_PER_INCH,
        child_height=2 * _EMU_PER_INCH,
        rotation_degrees=90,
        flip_horizontal=True,
    )

    clipping_slide = presentation.slides.add_slide(title_layout)
    _style_slide(clipping_slide)
    if clipping_slide.shapes.title is None:
        raise RuntimeError("Geometry corpus title layout is missing its title placeholder")
    _style_title(clipping_slide.shapes.title, "Slide boundary evidence")
    _add_label(
        clipping_slide,
        "One frame is fully outside; one is 25% clipped",
        left=0.55,
        top=1.0,
        width=6.2,
    )
    _add_picture(
        clipping_slide.shapes,
        image,
        name="VF_GEOM_OUTSIDE_PICTURE",
        alt_text="Picture frame positioned fully beyond the left slide edge",
        left=-2,
        top=2,
        width=1,
        height=1,
    )
    _add_picture(
        clipping_slide.shapes,
        image,
        name="VF_GEOM_CLIPPED_PICTURE",
        alt_text="Picture frame positioned one quarter beyond the left slide edge",
        left=-0.5,
        top=3.25,
        width=2,
        height=1,
    )

    offset_slide = presentation.slides.add_slide(title_layout)
    _style_slide(offset_slide)
    if offset_slide.shapes.title is None:
        raise RuntimeError("Geometry corpus title layout is missing its title placeholder")
    _style_title(offset_slide.shapes.title, "Child offset and rotation")
    _add_label(
        offset_slide,
        "Non-zero chOff maps a 6 x 4 child box into a rotated 3 x 2 frame",
        left=0.55,
        top=1.0,
        width=6.3,
    )
    offset_group = offset_slide.shapes.add_group_shape()
    _set_name(offset_group, "VF_GEOM_OFFSET_GROUP")
    _add_picture(
        offset_group.shapes,
        image,
        name="VF_GEOM_OFFSET_PICTURE",
        alt_text="Picture filling a non-zero child coordinate box inside a rotated group",
        left=1,
        top=2,
        width=6,
        height=4,
    )
    _set_group_transform(
        offset_group,
        x=7 * _EMU_PER_INCH,
        y=3.5 * _EMU_PER_INCH,
        width=3 * _EMU_PER_INCH,
        height=2 * _EMU_PER_INCH,
        child_x=_EMU_PER_INCH,
        child_y=2 * _EMU_PER_INCH,
        child_width=6 * _EMU_PER_INCH,
        child_height=4 * _EMU_PER_INCH,
        rotation_degrees=30,
    )

    package = io.BytesIO()
    presentation.save(package)
    artifact = _canonicalize_pptx_package(package.getvalue())
    validation = validate_pptx(artifact)
    if validation["valid"] is not True or validation["risky_features"]:
        raise RuntimeError("Generated geometry corpus is not a safe, valid PPTX package")
    preflight = office_engine.preflight(artifact, suffix=".pptx")
    if preflight["source_sha256"] != hashlib.sha256(artifact).hexdigest():
        raise RuntimeError("Geometry preflight evidence is not bound to the generated seed")
    return artifact, image, preflight


def _json_bytes(value: Any) -> bytes:
    return (json.dumps(value, indent=2, sort_keys=True) + "\n").encode("utf-8")


def _write(path: Path, payload: bytes, *, overwrite: bool) -> None:
    if path.exists() and not overwrite:
        raise FileExistsError(f"Destination already exists: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    artifact, image, preflight = build_geometry_seed()
    outputs = {
        "asset-v1.png": image,
        "seed-v1.pptx": artifact,
        "preflight-v1.json": _json_bytes(preflight),
    }
    for name, payload in outputs.items():
        _write(args.output_dir / name, payload, overwrite=args.overwrite)

    summary = {
        name: {
            "sha256": hashlib.sha256(payload).hexdigest(),
            "size_bytes": len(payload),
        }
        for name, payload in outputs.items()
    }
    print(json.dumps(summary, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()

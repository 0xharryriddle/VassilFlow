"""Build the deterministic PPTX seed used by native round-trip fixtures."""

from __future__ import annotations

import argparse
import copy
import hashlib
import io
import struct
import zipfile
import zlib
from datetime import datetime
from pathlib import Path, PurePosixPath

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from vassilflow.community.office.models import PptxObjectSelector, PptxTextReplacement
from vassilflow.community.office.pptx import edit_pptx, inspect_pptx, validate_pptx

_FIXED_ZIP_TIME = (2026, 7, 14, 0, 0, 0)
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_PACKAGE_RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_IMAGE_RELATIONSHIP_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
_PNG_IMAGE_PART = "ppt/media/image1.png"
_JPEG_IMAGE_PART = "ppt/media/image2.jpg"
_TILE_IMAGE_PART = "ppt/media/image3.png"
_CENTER_IMAGE_PART = "ppt/media/image4.png"
_BACKGROUND_STRETCH_IMAGE_PART = "ppt/media/image5.jpg"
_BACKGROUND_TILE_IMAGE_PART = "ppt/media/image6.png"
_BACKGROUND_CENTER_IMAGE_PART = "ppt/media/image7.png"
_PICTURE_SHARED_IMAGE_PART = "ppt/media/image8.png"
_PICTURE_JPEG_IMAGE_PART = "ppt/media/image9.jpg"
_IMAGE_RELATIONSHIP_ID = "rId2"
_DIRECT_FILL_TAGS = {
    qn("a:noFill"),
    qn("a:solidFill"),
    qn("a:gradFill"),
    qn("a:blipFill"),
    qn("a:pattFill"),
    qn("a:grpFill"),
}


def _set_linear_gradient(
    parent,
    *,
    stops: tuple[tuple[float, str, float | None], ...],
    angle_degrees: float,
    scaled: bool,
    rotate_with_shape: bool,
) -> None:
    fills = [child for child in parent if child.tag in _DIRECT_FILL_TAGS]
    if len(fills) != 1:
        raise RuntimeError("Round-trip gradient target requires one direct fill")

    gradient = etree.Element(qn("a:gradFill"))
    gradient.set("rotWithShape", "1" if rotate_with_shape else "0")
    stop_list = etree.SubElement(gradient, qn("a:gsLst"))
    for position_percent, color, opacity_percent in stops:
        stop = etree.SubElement(stop_list, qn("a:gs"))
        stop.set("pos", str(round(position_percent * 1_000)))
        rgb = etree.SubElement(stop, qn("a:srgbClr"))
        rgb.set("val", color.removeprefix("#").upper())
        if opacity_percent is not None:
            alpha = etree.SubElement(rgb, qn("a:alpha"))
            alpha.set("val", str(round(opacity_percent * 1_000)))
    linear = etree.SubElement(gradient, qn("a:lin"))
    linear.set("ang", str(round(angle_degrees * 60_000)))
    linear.set("scaled", "1" if scaled else "0")
    parent.replace(fills[0], gradient)


def _set_path_gradient(
    parent,
    *,
    stops: tuple[tuple[float, str, float | None], ...],
    path: str,
    fill_to_rectangle: tuple[float, float, float, float],
    rotate_with_shape: bool,
) -> None:
    fills = [child for child in parent if child.tag in _DIRECT_FILL_TAGS]
    if len(fills) != 1:
        raise RuntimeError("Round-trip gradient target requires one direct fill")

    gradient = etree.Element(qn("a:gradFill"))
    gradient.set("rotWithShape", "1" if rotate_with_shape else "0")
    stop_list = etree.SubElement(gradient, qn("a:gsLst"))
    for position_percent, color, opacity_percent in stops:
        stop = etree.SubElement(stop_list, qn("a:gs"))
        stop.set("pos", str(round(position_percent * 1_000)))
        rgb = etree.SubElement(stop, qn("a:srgbClr"))
        rgb.set("val", color.removeprefix("#").upper())
        if opacity_percent is not None:
            alpha = etree.SubElement(rgb, qn("a:alpha"))
            alpha.set("val", str(round(opacity_percent * 1_000)))
    path_geometry = etree.SubElement(gradient, qn("a:path"))
    path_geometry.set("path", path)
    fill_rectangle = etree.SubElement(path_geometry, qn("a:fillToRect"))
    for attribute, percentage in zip(
        ("l", "t", "r", "b"),
        fill_to_rectangle,
        strict=True,
    ):
        fill_rectangle.set(attribute, str(round(percentage * 1_000)))
    parent.replace(fills[0], gradient)


def _set_pattern_fill(
    parent,
    *,
    preset: str,
    foreground_color: str,
    background_color: str,
) -> None:
    fills = [child for child in parent if child.tag in _DIRECT_FILL_TAGS]
    if len(fills) != 1:
        raise RuntimeError("Round-trip pattern target requires one direct fill")

    pattern = etree.Element(qn("a:pattFill"))
    pattern.set("prst", preset)
    for child_name, color in (
        ("fgClr", foreground_color),
        ("bgClr", background_color),
    ):
        color_slot = etree.SubElement(pattern, qn(f"a:{child_name}"))
        rgb = etree.SubElement(color_slot, qn("a:srgbClr"))
        rgb.set("val", color.removeprefix("#").upper())
    parent.replace(fills[0], pattern)


def _png_chunk(chunk_type: bytes, payload: bytes) -> bytes:
    return struct.pack(">I", len(payload)) + chunk_type + payload + struct.pack(">I", zlib.crc32(chunk_type + payload) & 0xFFFFFFFF)


def _build_image_fill_png(
    *,
    colors: tuple[tuple[int, int, int], tuple[int, int, int], tuple[int, int, int]] = (
        (68, 114, 196),
        (112, 173, 71),
        (192, 43, 58),
    ),
    border_color: tuple[int, int, int] = (24, 36, 43),
    diagonal_color: tuple[int, int, int] = (255, 255, 255),
    width: int = 640,
    height: int = 360,
) -> bytes:
    scanlines = bytearray()
    for y in range(height):
        scanlines.append(0)
        for x in range(width):
            if x < 18 or x >= width - 18 or y < 18 or y >= height - 18:
                red, green, blue = border_color
            elif abs(x - y * 16 // 9) < 7 or abs((width - 1 - x) - y * 16 // 9) < 7:
                red, green, blue = diagonal_color
            elif x < width // 3:
                red, green, blue = colors[0]
            elif x < width * 2 // 3:
                red, green, blue = colors[1]
            else:
                red, green, blue = colors[2]
            scanlines.extend((red, green, blue, 255))
    header = struct.pack(">IIBBBBB", width, height, 8, 6, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + _png_chunk(b"IHDR", header) + _png_chunk(b"IDAT", zlib.compress(bytes(scanlines), level=9)) + _png_chunk(b"IEND", b"")


def build_image_fill_replacement_png() -> bytes:
    """Return the deterministic PNG used by image-fill writer acceptance."""
    return _build_image_fill_png(
        colors=(
            (112, 48, 160),
            (255, 192, 0),
            (0, 150, 136),
        ),
        border_color=(54, 37, 17),
        diagonal_color=(255, 242, 204),
    )


def _build_image_fill_jpeg(
    *,
    colors: tuple[tuple[int, int, int], tuple[int, int, int], tuple[int, int, int]] = (
        (0, 120, 167),
        (255, 192, 0),
        (192, 43, 58),
    ),
    border_color: tuple[int, int, int] = (24, 36, 43),
    diagonal_color: tuple[int, int, int] = (255, 255, 255),
) -> bytes:
    width = 640
    height = 360
    pixels = bytearray()
    for y in range(height):
        for x in range(width):
            if x < 18 or x >= width - 18 or y < 18 or y >= height - 18:
                color = border_color
            elif abs(x - y * 16 // 9) < 7 or abs((width - 1 - x) - y * 16 // 9) < 7:
                color = diagonal_color
            elif x < width // 3:
                color = colors[0]
            elif x < width * 2 // 3:
                color = colors[1]
            else:
                color = colors[2]
            pixels.extend(color)

    image = Image.frombytes("RGB", (width, height), bytes(pixels))
    output = io.BytesIO()
    image.save(
        output,
        format="JPEG",
        quality=92,
        subsampling=0,
        progressive=False,
        optimize=False,
        dpi=(96, 96),
        exif=b"",
    )
    return output.getvalue()


def build_image_fill_replacement_jpeg() -> bytes:
    """Return the deterministic baseline JPEG used by writer acceptance."""
    return _build_image_fill_jpeg(
        colors=(
            (126, 96, 191),
            (0, 176, 240),
            (146, 208, 80),
        ),
        border_color=(49, 31, 18),
        diagonal_color=(255, 242, 204),
    )


def _set_image_fill(
    parent,
    *,
    relationship_id: str,
    crop: tuple[float, float, float, float] | None,
    rotate_with_shape: bool | None,
    mode: str = "stretch",
    tile_attributes: dict[str, str] | None = None,
) -> None:
    fills = [child for child in parent if child.tag in _DIRECT_FILL_TAGS]
    if len(fills) != 1:
        raise RuntimeError("Round-trip image target requires one direct fill")

    image_fill = etree.Element(qn("a:blipFill"))
    if rotate_with_shape is not None:
        image_fill.set("rotWithShape", "1" if rotate_with_shape else "0")
    blip = etree.SubElement(image_fill, qn("a:blip"))
    blip.set(qn("r:embed"), relationship_id)
    if crop is not None:
        source_rectangle = etree.SubElement(image_fill, qn("a:srcRect"))
        for attribute, percentage in zip(("l", "t", "r", "b"), crop, strict=True):
            source_rectangle.set(attribute, str(round(percentage * 1_000)))
    if mode == "stretch":
        if tile_attributes is not None:
            raise RuntimeError("Stretch image fill does not accept tile attributes")
        stretch = etree.SubElement(image_fill, qn("a:stretch"))
        etree.SubElement(stretch, qn("a:fillRect"))
    elif mode in {"tile", "center"}:
        if crop is not None or tile_attributes is None:
            raise RuntimeError("Tiled image fill requires explicit attributes and no crop")
        tile = etree.SubElement(image_fill, qn("a:tile"))
        for attribute, value in tile_attributes.items():
            tile.set(attribute, value)
    else:
        raise RuntimeError(f"Unsupported round-trip image fill mode: {mode}")
    parent.replace(fills[0], image_fill)


def _embed_image_fill(
    data: bytes,
    *,
    slide_number: int,
    shape_name: str,
    image_part: str,
    image_payload: bytes,
    extension: str,
    content_type: str,
    crop: tuple[float, float, float, float] | None,
    rotate_with_shape: bool,
    relationship_id: str = _IMAGE_RELATIONSHIP_ID,
    mode: str = "stretch",
    tile_attributes: dict[str, str] | None = None,
) -> bytes:
    replacements: dict[str, bytes] = {}
    slide_part = f"ppt/slides/slide{slide_number}.xml"
    relationship_part = f"ppt/slides/_rels/slide{slide_number}.xml.rels"

    with zipfile.ZipFile(io.BytesIO(data)) as source:
        content_types = etree.fromstring(source.read("[Content_Types].xml"))
        matching_defaults = content_types.xpath(
            "ct:Default[translate(@Extension, 'ABCDEFGHIJKLMNOPQRSTUVWXYZ', 'abcdefghijklmnopqrstuvwxyz')=$extension]",
            namespaces={"ct": _CONTENT_TYPES_NS},
            extension=extension,
        )
        if not matching_defaults:
            default = etree.Element(f"{{{_CONTENT_TYPES_NS}}}Default")
            default.set("Extension", extension)
            default.set("ContentType", content_type)
            first_override = next(
                (index for index, child in enumerate(content_types) if child.tag == f"{{{_CONTENT_TYPES_NS}}}Override"),
                len(content_types),
            )
            content_types.insert(first_override, default)
        replacements["[Content_Types].xml"] = etree.tostring(
            content_types,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        relationships = etree.fromstring(source.read(relationship_part))
        if relationships.xpath(
            "rel:Relationship[@Id=$relationship_id]",
            namespaces={"rel": _PACKAGE_RELATIONSHIPS_NS},
            relationship_id=relationship_id,
        ):
            raise RuntimeError("Round-trip image relationship id is already in use")
        relationship = etree.SubElement(
            relationships,
            f"{{{_PACKAGE_RELATIONSHIPS_NS}}}Relationship",
        )
        relationship.set("Id", relationship_id)
        relationship.set("Type", _IMAGE_RELATIONSHIP_TYPE)
        relationship.set("Target", f"../media/{PurePosixPath(image_part).name}")
        replacements[relationship_part] = etree.tostring(
            relationships,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        slide = etree.fromstring(source.read(slide_part))
        matching_names = slide.xpath(
            ".//p:cNvPr[@name=$shape_name]",
            namespaces={"p": "http://schemas.openxmlformats.org/presentationml/2006/main"},
            shape_name=shape_name,
        )
        if len(matching_names) != 1:
            raise RuntimeError("Round-trip seed did not produce one image fill target")
        shape = matching_names[0].getparent().getparent()
        shape_properties = shape.find(qn("p:spPr"))
        if shape_properties is None:
            raise RuntimeError("Round-trip image target has no shape properties")
        _set_image_fill(
            shape_properties,
            relationship_id=relationship_id,
            crop=crop,
            rotate_with_shape=rotate_with_shape,
            mode=mode,
            tile_attributes=tile_attributes,
        )
        replacements[slide_part] = etree.tostring(
            slide,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as destination:
            destination.comment = source.comment
            for source_info in source.infolist():
                destination.writestr(
                    copy.copy(source_info),
                    replacements.get(source_info.filename, source.read(source_info)),
                )
            image_info = zipfile.ZipInfo(image_part, _FIXED_ZIP_TIME)
            image_info.compress_type = zipfile.ZIP_DEFLATED
            image_info.external_attr = 0o600 << 16
            destination.writestr(image_info, image_payload)
    return output.getvalue()


def _direct_background(fill) -> etree._Element:
    background = etree.Element(qn("p:bg"))
    properties = etree.SubElement(background, qn("p:bgPr"))
    properties.set("shadeToTitle", "0")
    properties.append(fill)
    return background


def _solid_background(color: str) -> etree._Element:
    fill = etree.Element(qn("a:solidFill"))
    rgb = etree.SubElement(fill, qn("a:srgbClr"))
    rgb.set("val", color.removeprefix("#").upper())
    return _direct_background(fill)


def _linear_gradient_background() -> etree._Element:
    background = _direct_background(etree.Element(qn("a:solidFill")))
    properties = background.find(qn("p:bgPr"))
    if properties is None:
        raise RuntimeError("Round-trip background has no properties")
    _set_linear_gradient(
        properties,
        stops=(
            (0, "#FFF2CC", None),
            (42.5, "#5B9BD5", 90),
            (100, "#203864", None),
        ),
        angle_degrees=132.5,
        scaled=True,
        rotate_with_shape=False,
    )
    return background


def _image_background(
    *,
    relationship_id: str,
    mode: str,
    tile_attributes: dict[str, str] | None = None,
) -> etree._Element:
    background = _direct_background(etree.Element(qn("a:solidFill")))
    properties = background.find(qn("p:bgPr"))
    if properties is None:
        raise RuntimeError("Round-trip background has no properties")
    _set_image_fill(
        properties,
        relationship_id=relationship_id,
        crop=None,
        rotate_with_shape=None,
        mode=mode,
        tile_attributes=tile_attributes,
    )
    return background


def _embed_slide_backgrounds(data: bytes) -> bytes:
    image_specs = {
        10: (
            _BACKGROUND_STRETCH_IMAGE_PART,
            _build_image_fill_jpeg(
                colors=((31, 78, 121), (112, 173, 71), (255, 192, 0)),
                border_color=(24, 36, 43),
                diagonal_color=(255, 255, 255),
            ),
        ),
        11: (
            _BACKGROUND_TILE_IMAGE_PART,
            _build_image_fill_png(
                colors=((255, 192, 0), (0, 176, 240), (192, 43, 58)),
                border_color=(56, 87, 35),
                diagonal_color=(255, 255, 255),
                width=320,
                height=180,
            ),
        ),
        12: (
            _BACKGROUND_CENTER_IMAGE_PART,
            _build_image_fill_png(
                colors=((112, 48, 160), (0, 150, 136), (237, 125, 49)),
                border_color=(54, 37, 17),
                diagonal_color=(255, 242, 204),
                width=1280,
                height=720,
            ),
        ),
    }
    backgrounds = {
        8: _solid_background("#D9EAF7"),
        9: _linear_gradient_background(),
        10: _image_background(relationship_id="rId2", mode="stretch"),
        11: _image_background(
            relationship_id="rId2",
            mode="tile",
            tile_attributes={
                "tx": "127000",
                "ty": "-63500",
                "sx": "40000",
                "sy": "55000",
                "algn": "tl",
                "flip": "none",
            },
        ),
        12: _image_background(
            relationship_id="rId2",
            mode="center",
            tile_attributes={
                "sx": "100000",
                "sy": "100000",
                "algn": "ctr",
                "flip": "none",
            },
        ),
    }
    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        for slide_number, background in backgrounds.items():
            slide_part = f"ppt/slides/slide{slide_number}.xml"
            slide = etree.fromstring(source.read(slide_part))
            common_slide_data = slide.find(qn("p:cSld"))
            if common_slide_data is None:
                raise RuntimeError("Round-trip slide has no common slide data")
            existing = common_slide_data.find(qn("p:bg"))
            if existing is not None:
                common_slide_data.remove(existing)
            shape_tree = common_slide_data.find(qn("p:spTree"))
            if shape_tree is None:
                raise RuntimeError("Round-trip slide has no shape tree")
            common_slide_data.insert(common_slide_data.index(shape_tree), background)
            replacements[slide_part] = etree.tostring(
                slide,
                xml_declaration=True,
                encoding="UTF-8",
                standalone=True,
            )

        for slide_number, (image_part, _payload) in image_specs.items():
            relationship_part = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
            relationships = etree.fromstring(source.read(relationship_part))
            if relationships.xpath(
                "rel:Relationship[@Id='rId2']",
                namespaces={"rel": _PACKAGE_RELATIONSHIPS_NS},
            ):
                raise RuntimeError("Round-trip background relationship id is already in use")
            relationship = etree.SubElement(
                relationships,
                f"{{{_PACKAGE_RELATIONSHIPS_NS}}}Relationship",
            )
            relationship.set("Id", "rId2")
            relationship.set("Type", _IMAGE_RELATIONSHIP_TYPE)
            relationship.set("Target", f"../media/{PurePosixPath(image_part).name}")
            replacements[relationship_part] = etree.tostring(
                relationships,
                xml_declaration=True,
                encoding="UTF-8",
                standalone=True,
            )

        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as destination:
            destination.comment = source.comment
            for source_info in source.infolist():
                destination.writestr(
                    copy.copy(source_info),
                    replacements.get(source_info.filename, source.read(source_info)),
                )
            for image_part, image_payload in image_specs.values():
                image_info = zipfile.ZipInfo(image_part, _FIXED_ZIP_TIME)
                image_info.compress_type = zipfile.ZIP_DEFLATED
                image_info.external_attr = 0o600 << 16
                destination.writestr(image_info, image_payload)
    return output.getvalue()


def _picture_element(
    *,
    object_id: int,
    name: str,
    description: str,
    relationship_id: str,
    x_emu: int,
    y_emu: int,
    width_emu: int,
    height_emu: int,
    mode: str,
    crop: tuple[float, float, float, float] | None = None,
    fill_rectangle: tuple[float, float, float, float] | None = None,
    tile_attributes: dict[str, str] | None = None,
    dpi: int | None = None,
    rotate_with_shape: bool | None = None,
    compression_state: str | None = None,
    alpha_percent: float | None = None,
    brightness_percent: float | None = None,
    contrast_percent: float | None = None,
    grayscale: bool = False,
) -> etree._Element:
    picture = etree.Element(qn("p:pic"))
    non_visual = etree.SubElement(picture, qn("p:nvPicPr"))
    drawing_properties = etree.SubElement(non_visual, qn("p:cNvPr"))
    drawing_properties.set("id", str(object_id))
    drawing_properties.set("name", name)
    drawing_properties.set("descr", description)
    picture_properties = etree.SubElement(non_visual, qn("p:cNvPicPr"))
    locks = etree.SubElement(picture_properties, qn("a:picLocks"))
    locks.set("noChangeAspect", "1")
    etree.SubElement(non_visual, qn("p:nvPr"))

    blip_fill = etree.SubElement(picture, qn("p:blipFill"))
    if dpi is not None:
        blip_fill.set("dpi", str(dpi))
    if rotate_with_shape is not None:
        blip_fill.set("rotWithShape", "1" if rotate_with_shape else "0")
    blip = etree.SubElement(blip_fill, qn("a:blip"))
    blip.set(qn("r:embed"), relationship_id)
    if compression_state is not None:
        blip.set("cstate", compression_state)
    if alpha_percent is not None:
        alpha = etree.SubElement(blip, qn("a:alphaModFix"))
        alpha.set("amt", str(round(alpha_percent * 1_000)))
    if grayscale:
        etree.SubElement(blip, qn("a:grayscl"))
    if brightness_percent is not None or contrast_percent is not None:
        luminance = etree.SubElement(blip, qn("a:lum"))
        if brightness_percent is not None:
            luminance.set("bright", str(round(brightness_percent * 1_000)))
        if contrast_percent is not None:
            luminance.set("contrast", str(round(contrast_percent * 1_000)))
    if crop is not None:
        source_rectangle = etree.SubElement(blip_fill, qn("a:srcRect"))
        for attribute, percentage in zip(("l", "t", "r", "b"), crop, strict=True):
            source_rectangle.set(attribute, str(round(percentage * 1_000)))
    if mode == "stretch":
        if tile_attributes is not None:
            raise RuntimeError("Stretch picture does not accept tile attributes")
        stretch = etree.SubElement(blip_fill, qn("a:stretch"))
        fill = etree.SubElement(stretch, qn("a:fillRect"))
        if fill_rectangle is not None:
            for attribute, percentage in zip(
                ("l", "t", "r", "b"),
                fill_rectangle,
                strict=True,
            ):
                fill.set(attribute, str(round(percentage * 1_000)))
    elif mode in {"tile", "center"}:
        if crop is not None or fill_rectangle is not None or tile_attributes is None:
            raise RuntimeError("Tiled picture requires explicit attributes and no rectangles")
        tile = etree.SubElement(blip_fill, qn("a:tile"))
        for attribute, value in tile_attributes.items():
            tile.set(attribute, value)
    else:
        raise RuntimeError(f"Unsupported round-trip picture mode: {mode}")

    shape_properties = etree.SubElement(picture, qn("p:spPr"))
    transform = etree.SubElement(shape_properties, qn("a:xfrm"))
    offset = etree.SubElement(transform, qn("a:off"))
    offset.set("x", str(x_emu))
    offset.set("y", str(y_emu))
    extents = etree.SubElement(transform, qn("a:ext"))
    extents.set("cx", str(width_emu))
    extents.set("cy", str(height_emu))
    geometry = etree.SubElement(shape_properties, qn("a:prstGeom"))
    geometry.set("prst", "rect")
    etree.SubElement(geometry, qn("a:avLst"))
    line = etree.SubElement(shape_properties, qn("a:ln"))
    line.set("w", "19050")
    line_fill = etree.SubElement(line, qn("a:solidFill"))
    line_color = etree.SubElement(line_fill, qn("a:srgbClr"))
    line_color.set("val", "18242B")
    return picture


def _embed_picture_corpus(data: bytes) -> bytes:
    slide_number = 13
    slide_part = f"ppt/slides/slide{slide_number}.xml"
    relationship_part = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
    shared_payload = _build_image_fill_png(
        colors=((192, 43, 58), (68, 114, 196), (112, 173, 71)),
        border_color=(24, 36, 43),
        diagonal_color=(255, 242, 204),
    )
    center_payload = _build_image_fill_jpeg(
        colors=((0, 150, 136), (237, 125, 49), (112, 48, 160)),
        border_color=(49, 31, 18),
        diagonal_color=(255, 255, 255),
    )
    pictures = (
        _picture_element(
            object_id=3,
            name="VF_RT_PICTURE_EFFECT",
            description="Source replacement preserves crop and blip effects",
            relationship_id="rId2",
            x_emu=685_800,
            y_emu=1_371_600,
            width_emu=3_429_000,
            height_emu=4_191_000,
            mode="stretch",
            crop=(10, 5, 15, 12),
            fill_rectangle=(-5, 0, 3, 0),
            dpi=150,
            rotate_with_shape=False,
            compression_state="screen",
            alpha_percent=85,
            brightness_percent=12,
            contrast_percent=-8,
        ),
        _picture_element(
            object_id=4,
            name="VF_RT_PICTURE_SHARED",
            description="Shared source remains unchanged when its peer is replaced",
            relationship_id="rId2",
            x_emu=4_343_400,
            y_emu=1_371_600,
            width_emu=3_429_000,
            height_emu=4_191_000,
            mode="stretch",
            rotate_with_shape=True,
            grayscale=True,
        ),
        _picture_element(
            object_id=5,
            name="VF_RT_PICTURE_JPEG",
            description="Baseline JPEG picture source",
            relationship_id="rId3",
            x_emu=8_001_000,
            y_emu=1_371_600,
            width_emu=3_429_000,
            height_emu=4_191_000,
            mode="stretch",
            dpi=96,
        ),
    )

    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        slide = etree.fromstring(source.read(slide_part))
        shape_tree = slide.find(f"{qn('p:cSld')}/{qn('p:spTree')}")
        if shape_tree is None:
            raise RuntimeError("Round-trip picture slide has no shape tree")
        shape_tree.extend(pictures)
        replacements[slide_part] = etree.tostring(
            slide,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        relationships = etree.fromstring(source.read(relationship_part))
        for relationship_id, image_part in (
            ("rId2", _PICTURE_SHARED_IMAGE_PART),
            ("rId3", _PICTURE_JPEG_IMAGE_PART),
        ):
            if relationships.xpath(
                "rel:Relationship[@Id=$relationship_id]",
                namespaces={"rel": _PACKAGE_RELATIONSHIPS_NS},
                relationship_id=relationship_id,
            ):
                raise RuntimeError("Round-trip picture relationship id is already in use")
            relationship = etree.SubElement(
                relationships,
                f"{{{_PACKAGE_RELATIONSHIPS_NS}}}Relationship",
            )
            relationship.set("Id", relationship_id)
            relationship.set("Type", _IMAGE_RELATIONSHIP_TYPE)
            relationship.set("Target", f"../media/{PurePosixPath(image_part).name}")
        replacements[relationship_part] = etree.tostring(
            relationships,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as destination:
            destination.comment = source.comment
            for source_info in source.infolist():
                destination.writestr(
                    copy.copy(source_info),
                    replacements.get(source_info.filename, source.read(source_info)),
                )
            for image_part, image_payload in (
                (_PICTURE_SHARED_IMAGE_PART, shared_payload),
                (_PICTURE_JPEG_IMAGE_PART, center_payload),
            ):
                image_info = zipfile.ZipInfo(image_part, _FIXED_ZIP_TIME)
                image_info.compress_type = zipfile.ZIP_DEFLATED
                image_info.external_attr = 0o600 << 16
                destination.writestr(image_info, image_payload)
    return output.getvalue()


def _add_title(slide, text: str, *, foreground_label: bool = False) -> None:
    shape = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.7))
    shape.name = "VF_RT_TITLE"
    if foreground_label:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shape.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos Display"
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x18, 0x24, 0x2B)


def _build_source() -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "VassilFlow PPTX native round-trip corpus"
    presentation.core_properties.subject = "Interop fixture for bounded PPTX editing"
    presentation.core_properties.author = "VassilFlow QA"
    presentation.core_properties.last_modified_by = "VassilFlow QA"
    presentation.core_properties.created = datetime(2026, 7, 14, 0, 0, 0)
    presentation.core_properties.modified = datetime(2026, 7, 14, 0, 0, 0)

    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    _add_title(slide, "Run and paragraph preservation")

    split = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(5.7), Inches(1.1))
    split.name = "VF_RT_SPLIT_RUN"
    split.text_frame.clear()
    paragraph = split.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    first = paragraph.add_run()
    first.text = "Alpha"
    first.font.name = "Aptos"
    first.font.size = Pt(28)
    first.font.bold = True
    first.font.color.rgb = RGBColor(0xC0, 0x2B, 0x3A)
    second = paragraph.add_run()
    second.text = "Beta"
    second.font.name = "Georgia"
    second.font.size = Pt(20)
    second.font.italic = True
    second.font.color.rgb = RGBColor(0x10, 0x7C, 0x71)

    paragraphs = slide.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(5.7), Inches(2.6))
    paragraphs.name = "VF_RT_PARAGRAPHS"
    paragraphs.text_frame.clear()
    paragraphs.text_frame.word_wrap = True
    for index, (text, alignment, size) in enumerate(
        (
            ("Left paragraph", PP_ALIGN.LEFT, 18),
            ("Centered paragraph", PP_ALIGN.CENTER, 20),
            ("Right paragraph", PP_ALIGN.RIGHT, 16),
        )
    ):
        paragraph = paragraphs.text_frame.paragraphs[0] if index == 0 else paragraphs.text_frame.add_paragraph()
        paragraph.alignment = alignment
        paragraph.space_after = Pt(8)
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)

    sentinel = slide.shapes.add_textbox(Inches(7.0), Inches(1.45), Inches(5.2), Inches(1.2))
    sentinel.name = "VF_RT_EDIT_SENTINEL"
    sentinel.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    sentinel.text_frame.paragraphs[0].text = "Draft native round-trip sentinel"
    sentinel.text_frame.paragraphs[0].runs[0].font.name = "Aptos"
    sentinel.text_frame.paragraphs[0].runs[0].font.size = Pt(22)

    styled = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(3.0), Inches(4.5), Inches(2.0))
    styled.name = "VF_RT_STYLE"
    styled.fill.solid()
    styled.fill.fore_color.rgb = RGBColor(0x33, 0x66, 0x99)
    styled.line.color.rgb = RGBColor(0xE0, 0x4B, 0x42)
    styled.line.width = Pt(2.5)
    styled.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    styled.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    styled.text_frame.paragraphs[0].text = "Direct fill and line"
    styled.text_frame.paragraphs[0].runs[0].font.name = "Aptos"
    styled.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    styled.text_frame.paragraphs[0].runs[0].font.bold = True
    styled.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    gradient_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.0),
        Inches(5.35),
        Inches(4.5),
        Inches(1.25),
    )
    gradient_shape.name = "VF_RT_GRADIENT"
    gradient_shape.fill.solid()
    gradient_shape.fill.fore_color.rgb = RGBColor(0x20, 0x38, 0x64)
    gradient_shape.line.color.rgb = RGBColor(0x18, 0x24, 0x2B)
    gradient_shape.line.width = Pt(1.5)
    gradient_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    gradient_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    gradient_shape.text_frame.paragraphs[0].text = "Authored linear gradient"
    gradient_run = gradient_shape.text_frame.paragraphs[0].runs[0]
    gradient_run.font.name = "Aptos"
    gradient_run.font.size = Pt(17)
    gradient_run.font.bold = True
    gradient_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _set_linear_gradient(
        gradient_shape._element.spPr,
        stops=(
            (0, "#203864", None),
            (37.5, "#4472C4", 85),
            (100, "#70AD47", None),
        ),
        angle_degrees=37.5,
        scaled=True,
        rotate_with_shape=True,
    )

    slide = presentation.slides.add_slide(blank)
    _add_title(slide, "Unicode and layout preservation")
    unicode_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.7), Inches(1.5))
    unicode_box.name = "VF_RT_UNICODE"
    unicode_box.text_frame.word_wrap = True
    unicode_box.text_frame.paragraphs[0].text = "Hello, 世界! Καλημέρα κόσμε. مرحبا بالعالم."
    unicode_box.text_frame.paragraphs[0].runs[0].font.name = "Noto Sans"
    unicode_box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)

    margins = slide.shapes.add_textbox(Inches(1.1), Inches(3.4), Inches(5.2), Inches(2.0))
    margins.name = "VF_RT_MARGINS"
    margins.text_frame.margin_left = Inches(0.35)
    margins.text_frame.margin_right = Inches(0.2)
    margins.text_frame.margin_top = Inches(0.25)
    margins.text_frame.margin_bottom = Inches(0.15)
    margins.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    margins.text_frame.paragraphs[0].text = "Bottom anchored with authored margins"
    margins.text_frame.paragraphs[0].runs[0].font.name = "Aptos"
    margins.text_frame.paragraphs[0].runs[0].font.size = Pt(18)

    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(7.15),
        Inches(4.1),
        Inches(11.7),
        Inches(5.45),
    )
    connector.name = "VF_RT_CONNECTOR"
    connector.line.color.rgb = RGBColor(0x44, 0x72, 0xC4)
    connector.line.width = Pt(3)
    line = connector.line._ln
    if line is None:
        raise RuntimeError("Round-trip connector did not produce direct line properties")
    line.set("cmpd", "dbl")
    line.set("algn", "ctr")
    head_end = etree.Element(qn("a:headEnd"))
    head_end.set("type", "diamond")
    head_end.set("w", "sm")
    head_end.set("len", "med")
    line.append(head_end)
    tail_end = etree.Element(qn("a:tailEnd"))
    tail_end.set("type", "triangle")
    tail_end.set("w", "lg")
    tail_end.set("len", "lg")
    line.append(tail_end)

    gradient_connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(7.15),
        Inches(6.15),
        Inches(11.7),
        Inches(6.15),
    )
    gradient_connector.name = "VF_RT_GRADIENT_CONNECTOR"
    gradient_connector.line.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    gradient_connector.line.width = Pt(5)
    gradient_line = gradient_connector.line._ln
    if gradient_line is None:
        raise RuntimeError("Round-trip gradient connector did not produce direct line properties")
    _set_linear_gradient(
        gradient_line,
        stops=(
            (0, "#C00000", None),
            (50, "#FFC000", None),
            (100, "#70AD47", None),
        ),
        angle_degrees=0,
        scaled=True,
        rotate_with_shape=True,
    )

    slide = presentation.slides.add_slide(blank)
    _add_title(slide, "Path gradient and preset geometry preservation")

    path_gradient_shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(0.9),
        Inches(1.55),
        Inches(5.3),
        Inches(4.65),
    )
    path_gradient_shape.name = "VF_RT_PATH_GRADIENT"
    path_gradient_shape.fill.solid()
    path_gradient_shape.fill.fore_color.rgb = RGBColor(0x20, 0x38, 0x64)
    path_gradient_shape.line.color.rgb = RGBColor(0x18, 0x24, 0x2B)
    path_gradient_shape.line.width = Pt(1.5)
    path_gradient_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    path_gradient_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    path_gradient_shape.text_frame.paragraphs[0].text = "Authored path gradient"
    path_gradient_run = path_gradient_shape.text_frame.paragraphs[0].runs[0]
    path_gradient_run.font.name = "Aptos"
    path_gradient_run.font.size = Pt(20)
    path_gradient_run.font.bold = True
    path_gradient_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _set_path_gradient(
        path_gradient_shape._element.spPr,
        stops=(
            (0, "#FFFFFF", None),
            (35, "#5B9BD5", 90),
            (100, "#203864", None),
        ),
        path="circle",
        fill_to_rectangle=(35, 25, 65, 75),
        rotate_with_shape=False,
    )

    preset_geometry_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.0),
        Inches(2.25),
        Inches(4.7),
        Inches(2.7),
    )
    preset_geometry_shape.name = "VF_RT_PRESET_GEOMETRY"
    preset_geometry_shape.fill.solid()
    preset_geometry_shape.fill.fore_color.rgb = RGBColor(0x70, 0xAD, 0x47)
    preset_geometry_shape.line.color.rgb = RGBColor(0x38, 0x59, 0x23)
    preset_geometry_shape.line.width = Pt(2)
    preset_geometry_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    preset_geometry_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    preset_geometry_shape.text_frame.paragraphs[0].text = "Preset geometry"
    preset_geometry_run = preset_geometry_shape.text_frame.paragraphs[0].runs[0]
    preset_geometry_run.font.name = "Aptos"
    preset_geometry_run.font.size = Pt(20)
    preset_geometry_run.font.bold = True
    preset_geometry_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    slide = presentation.slides.add_slide(blank)
    _add_title(slide, "Pattern fill preservation")

    pattern_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(2.15),
        Inches(1.55),
        Inches(9.0),
        Inches(4.85),
    )
    pattern_shape.name = "VF_RT_PATTERN_FILL"
    pattern_shape.fill.solid()
    pattern_shape.fill.fore_color.rgb = RGBColor(0xD9, 0xEA, 0xF7)
    pattern_shape.line.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    pattern_shape.line.width = Pt(2)
    pattern_shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    pattern_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    pattern_shape.text_frame.paragraphs[0].text = "Authored direct RGB pattern"
    pattern_run = pattern_shape.text_frame.paragraphs[0].runs[0]
    pattern_run.font.name = "Aptos"
    pattern_run.font.size = Pt(24)
    pattern_run.font.bold = True
    pattern_run.font.color.rgb = RGBColor(0x18, 0x24, 0x2B)
    _set_pattern_fill(
        pattern_shape._element.spPr,
        preset="diagCross",
        foreground_color="#4472C4",
        background_color="#D9EAF7",
    )

    slide = presentation.slides.add_slide(blank)
    _add_title(slide, "Image fill preservation")

    image_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(2.15),
        Inches(1.55),
        Inches(9.0),
        Inches(4.85),
    )
    image_shape.name = "VF_RT_IMAGE_FILL"
    image_shape.fill.solid()
    image_shape.fill.fore_color.rgb = RGBColor(0x44, 0x72, 0xC4)
    image_shape.line.color.rgb = RGBColor(0x18, 0x24, 0x2B)
    image_shape.line.width = Pt(2)

    slide = presentation.slides.add_slide(blank)
    _add_title(slide, "Baseline JPEG fill preservation")

    jpeg_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(2.15),
        Inches(1.55),
        Inches(9.0),
        Inches(4.85),
    )
    jpeg_shape.name = "VF_RT_JPEG_FILL"
    jpeg_shape.fill.solid()
    jpeg_shape.fill.fore_color.rgb = RGBColor(0x00, 0x78, 0xA7)
    jpeg_shape.line.color.rgb = RGBColor(0x59, 0x55, 0x95)
    jpeg_shape.line.width = Pt(2)

    slide = presentation.slides.add_slide(blank)
    _add_title(slide, "Tile and centered image framing")

    tile_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.75),
        Inches(1.55),
        Inches(5.8),
        Inches(4.9),
    )
    tile_shape.name = "VF_RT_IMAGE_TILE"
    tile_shape.fill.solid()
    tile_shape.fill.fore_color.rgb = RGBColor(0x70, 0xAD, 0x47)
    tile_shape.line.color.rgb = RGBColor(0x38, 0x59, 0x23)
    tile_shape.line.width = Pt(3)

    center_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.78),
        Inches(1.55),
        Inches(5.8),
        Inches(4.9),
    )
    center_shape.name = "VF_RT_IMAGE_CENTER"
    center_shape.fill.solid()
    center_shape.fill.fore_color.rgb = RGBColor(0x70, 0x35, 0xA0)
    center_shape.line.color.rgb = RGBColor(0x59, 0x55, 0x95)
    center_shape.line.width = Pt(3)

    for title in (
        "Direct solid slide background",
        "Direct linear-gradient slide background",
        "Stretched JPEG slide background",
        "Tiled PNG slide background",
        "Centered PNG slide background",
    ):
        slide = presentation.slides.add_slide(blank)
        _add_title(slide, title, foreground_label=True)

    slide = presentation.slides.add_slide(blank)
    _add_title(slide, "Embedded picture source preservation")

    output = io.BytesIO()
    presentation.save(output)
    with_png = _embed_image_fill(
        output.getvalue(),
        slide_number=5,
        shape_name="VF_RT_IMAGE_FILL",
        image_part=_PNG_IMAGE_PART,
        image_payload=_build_image_fill_png(),
        extension="png",
        content_type="image/png",
        crop=(12.5, 5.0, 8.0, 10.0),
        rotate_with_shape=False,
    )
    with_jpeg = _embed_image_fill(
        with_png,
        slide_number=6,
        shape_name="VF_RT_JPEG_FILL",
        image_part=_JPEG_IMAGE_PART,
        image_payload=_build_image_fill_jpeg(),
        extension="jpg",
        content_type="image/jpeg",
        crop=(6.0, 9.0, 14.0, 4.0),
        rotate_with_shape=True,
    )
    with_tile = _embed_image_fill(
        with_jpeg,
        slide_number=7,
        shape_name="VF_RT_IMAGE_TILE",
        image_part=_TILE_IMAGE_PART,
        image_payload=_build_image_fill_png(
            colors=((255, 192, 0), (0, 176, 240), (192, 43, 58)),
            border_color=(56, 87, 35),
            diagonal_color=(255, 255, 255),
        ),
        extension="png",
        content_type="image/png",
        crop=None,
        rotate_with_shape=False,
        relationship_id="rId2",
        mode="tile",
        tile_attributes={
            "tx": "127000",
            "ty": "-63500",
            "sx": "45000",
            "sy": "65000",
            "algn": "br",
            "flip": "x",
        },
    )
    with_center = _embed_image_fill(
        with_tile,
        slide_number=7,
        shape_name="VF_RT_IMAGE_CENTER",
        image_part=_CENTER_IMAGE_PART,
        image_payload=_build_image_fill_png(
            colors=((112, 48, 160), (0, 150, 136), (237, 125, 49)),
            border_color=(54, 37, 17),
            diagonal_color=(255, 242, 204),
        ),
        extension="png",
        content_type="image/png",
        crop=None,
        rotate_with_shape=True,
        relationship_id="rId3",
        mode="center",
        tile_attributes={
            "sx": "100000",
            "sy": "100000",
            "algn": "ctr",
            "flip": "none",
        },
    )
    return _embed_picture_corpus(_embed_slide_backgrounds(with_center))


def _normalize_zip_timestamps(data: bytes) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as destination:
        destination.comment = source.comment
        for source_info in source.infolist():
            info = copy.copy(source_info)
            info.date_time = _FIXED_ZIP_TIME
            destination.writestr(info, source.read(source_info))
    return output.getvalue()


def build_seed() -> bytes:
    source = _normalize_zip_timestamps(_build_source())
    inspected = inspect_pptx(
        source,
        max_slides=20,
        include_formatting=True,
        selector=PptxObjectSelector(name_equals="VF_RT_EDIT_SENTINEL"),
    )
    targets = [item for slide in inspected["slides"] for item in slide["objects"]]
    if len(targets) != 1 or targets[0]["identity_source"] != "cNvPr.id":
        raise RuntimeError("Round-trip seed did not produce one stable edit sentinel")
    edited, reports = edit_pptx(
        source,
        [
            PptxTextReplacement(
                paths=[targets[0]["path"]],
                find="Draft",
                replace="Final",
            )
        ],
    )
    if reports[0]["match_count"] != 1 or not validate_pptx(edited)["valid"]:
        raise RuntimeError("Round-trip seed edit failed validation")
    return edited


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    output = args.output.resolve()
    if output.exists() and not args.overwrite:
        raise FileExistsError(f"Destination already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = build_seed()
    output.write_bytes(payload)
    print(f"wrote {len(payload)} bytes sha256={hashlib.sha256(payload).hexdigest()} -> {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

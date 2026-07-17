import sys
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from skill_loader import load  # noqa: E402

ppt = load("ppt-generation")


def test_generate_ppt_reads_bom_plan_with_explicit_flattened_acknowledgement(tmp_path):
    plan = tmp_path / "plan.json"
    plan.write_text(
        '{"aspect_ratio":"16:9","slides":[{"title":"Smoke","key_points":["bom"]}]}',
        encoding="utf-8-sig",
    )
    slide = tmp_path / "slide.png"
    Image.new("RGB", (1280, 720), (38, 99, 235)).save(slide)
    output = tmp_path / "deck.pptx"

    result = ppt.generate_ppt(
        str(plan),
        [str(slide)],
        str(output),
        acknowledge_flattened_output=True,
    )

    assert result == {
        "ok": True,
        "output_mode": "raster_composite",
        "editable_objects": False,
        "slide_count": 1,
        "output_path": str(output),
        "output_sha256": result["output_sha256"],
    }
    assert len(result["output_sha256"]) == 64
    deck = Presentation(output)
    assert len(deck.slides) == 1
    assert len(deck.slides[0].shapes) == 1
    assert deck.slides[0].shapes[0].shape_type is not None


def test_generate_ppt_refuses_unacknowledged_flattening(tmp_path):
    plan = tmp_path / "plan.json"
    plan.write_text('{"slides":[{}]}', encoding="utf-8")
    slide = tmp_path / "slide.png"
    Image.new("RGB", (1280, 720), (38, 99, 235)).save(slide)

    with pytest.raises(ValueError, match="acknowledged explicitly"):
        ppt.generate_ppt(str(plan), [str(slide)], str(tmp_path / "deck.pptx"))


def test_generate_ppt_requires_one_image_per_plan_slide(tmp_path):
    plan = tmp_path / "plan.json"
    plan.write_text('{"slides":[{},{}]}', encoding="utf-8")
    slide = tmp_path / "slide.png"
    Image.new("RGB", (1280, 720), (38, 99, 235)).save(slide)

    with pytest.raises(ValueError, match="must match plan slide count"):
        ppt.generate_ppt(
            str(plan),
            [str(slide)],
            str(tmp_path / "deck.pptx"),
            acknowledge_flattened_output=True,
        )

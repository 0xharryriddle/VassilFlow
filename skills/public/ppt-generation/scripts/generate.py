"""Explicit fallback for composing flattened slide images into a PPTX.

The resulting slides contain one raster image each. Use the Office generation
tool for editable native presentations.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import tempfile
from pathlib import Path
from typing import Any

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.util import Inches

_SLIDE_SIZES = {
    "16:9": (Inches(13.333), Inches(7.5)),
    "4:3": (Inches(10), Inches(7.5)),
}
_SUPPORTED_IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png"}


def _load_plan(plan_file: str) -> dict[str, Any]:
    path = Path(plan_file)
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except FileNotFoundError as exc:
        raise ValueError(f"Plan file does not exist: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"Plan file is not valid JSON: {exc.msg}") from exc
    if not isinstance(value, dict):
        raise ValueError("Plan must be a JSON object")
    slides = value.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("Plan must contain a non-empty slides array")
    if not all(isinstance(slide, dict) for slide in slides):
        raise ValueError("Every plan slide must be a JSON object")
    return value


def _validated_image(path_value: str) -> tuple[Path, int, int]:
    path = Path(path_value)
    if path.suffix.lower() not in _SUPPORTED_IMAGE_SUFFIXES:
        raise ValueError(f"Slide image must be PNG or JPEG: {path}")
    if not path.is_file():
        raise ValueError(f"Slide image does not exist: {path}")
    try:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            width, height = image.size
    except (OSError, UnidentifiedImageError) as exc:
        raise ValueError(f"Slide image is invalid: {path}") from exc
    if width <= 0 or height <= 0:
        raise ValueError(f"Slide image has invalid dimensions: {path}")
    return path, width, height


def _add_notes(slide: Any, slide_info: dict[str, Any]) -> None:
    notes: list[str] = []
    for label, field in (("Title", "title"), ("Subtitle", "subtitle")):
        value = slide_info.get(field)
        if isinstance(value, str) and value.strip():
            notes.append(f"{label}: {value.strip()}")
    key_points = slide_info.get("key_points")
    if isinstance(key_points, list) and key_points:
        notes.append("Key Points:")
        notes.extend(f"  - {point}" for point in key_points if isinstance(point, str))
    if notes:
        text_frame = slide.notes_slide.notes_text_frame
        if text_frame is not None:
            text_frame.text = "\n".join(notes)


def generate_ppt(
    plan_file: str,
    slide_images: list[str],
    output_file: str,
    *,
    acknowledge_flattened_output: bool = False,
    overwrite: bool = False,
) -> dict[str, Any]:
    """Compose a flattened, image-only PPTX after explicit acknowledgement."""

    if not acknowledge_flattened_output:
        raise ValueError(
            "Flattened output must be acknowledged explicitly; use office_generate "
            "for editable native PPTX output"
        )

    plan = _load_plan(plan_file)
    aspect_ratio = plan.get("aspect_ratio", "16:9")
    if aspect_ratio not in _SLIDE_SIZES:
        raise ValueError("aspect_ratio must be 16:9 or 4:3")
    slides_info = plan["slides"]
    if len(slide_images) != len(slides_info):
        raise ValueError(
            f"Slide image count ({len(slide_images)}) must match plan slide count "
            f"({len(slides_info)})"
        )

    output = Path(output_file)
    if output.suffix.lower() != ".pptx":
        raise ValueError("Output file must end with .pptx")
    if output.exists() and not overwrite:
        raise FileExistsError(f"Output already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)

    images = [_validated_image(value) for value in slide_images]
    slide_width, slide_height = _SLIDE_SIZES[aspect_ratio]
    presentation = Presentation()
    presentation.slide_width = slide_width
    presentation.slide_height = slide_height
    blank_layout = presentation.slide_layouts[6]

    for index, (image_path, image_width, image_height) in enumerate(images):
        slide = presentation.slides.add_slide(blank_layout)
        image_aspect = image_width / image_height
        slide_aspect = int(slide_width) / int(slide_height)
        if image_aspect > slide_aspect:
            width = int(slide_width)
            height = round(width / image_aspect)
            left = 0
            top = (int(slide_height) - height) // 2
        else:
            height = int(slide_height)
            width = round(height * image_aspect)
            left = (int(slide_width) - width) // 2
            top = 0
        slide.shapes.add_picture(str(image_path), left, top, width, height)
        _add_notes(slide, slides_info[index])

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output.stem}-",
            suffix=".pptx",
            dir=output.parent,
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        presentation.save(temporary_path)
        os.replace(temporary_path, output)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)

    payload = output.read_bytes()
    return {
        "ok": True,
        "output_mode": "raster_composite",
        "editable_objects": False,
        "slide_count": len(images),
        "output_path": str(output),
        "output_sha256": hashlib.sha256(payload).hexdigest(),
    }


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Compose an explicitly flattened image-only PowerPoint presentation"
    )
    parser.add_argument("--plan-file", required=True, help="Path to the JSON presentation plan")
    parser.add_argument(
        "--slide-images",
        nargs="+",
        required=True,
        help="PNG or JPEG slide images in plan order",
    )
    parser.add_argument("--output-file", required=True, help="Output .pptx path")
    parser.add_argument(
        "--acknowledge-flattened-output",
        action="store_true",
        help="Acknowledge that slide elements will not be independently editable",
    )
    parser.add_argument("--overwrite", action="store_true", help="Replace an existing output")
    return parser


if __name__ == "__main__":
    argument_parser = _parser()
    arguments = argument_parser.parse_args()
    try:
        result = generate_ppt(
            arguments.plan_file,
            arguments.slide_images,
            arguments.output_file,
            acknowledge_flattened_output=arguments.acknowledge_flattened_output,
            overwrite=arguments.overwrite,
        )
    except Exception as exc:
        argument_parser.exit(1, f"Error: {exc}\n")
    print(json.dumps(result, indent=2, sort_keys=True))
